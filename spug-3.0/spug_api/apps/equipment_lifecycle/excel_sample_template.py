# Copyright: (c) OpenSpug Organization. https://github.com/openspug
# Released under the AGPL-3.0 License.
"""设备全生命周期——Excel 上传版维护台账验收样例模板（2026-09-28）。

生成任务书要求的验收样本 .xlsx：信道 × 主备电台 ×
PS/TX/RX/BB/显示模块/备注，固定行列、固定填写位置；包含合并单元格、
行高列宽、边框字体、打印区域、纸张方向、页边距、页眉页脚与第二个
工作表，作为版式验收（Excel/PDF 导出对比）的标准输入。
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

MODULE_HEADERS = ['信道', '电台', 'PS模块', 'TX模块', 'RX模块', 'BB模块', '显示模块', '备注']
CHANNEL_ROWS = [('主电台', '备电台'), ('主电台', '备电台')]

COLUMN_WIDTHS = {'A': 12, 'B': 12, 'C': 12, 'D': 12, 'E': 12, 'F': 12, 'G': 14, 'H': 20}
HEADER_ROW_HEIGHT = 26
DATA_ROW_HEIGHT = 24


def build_acceptance_sample_xlsx():
    """生成验收样例：工作表1=信道维护记录（打印版式），工作表2=维护结论。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '信道维护记录'
    _fill_channel_sheet(ws)

    summary = wb.create_sheet('维护结论')
    _fill_summary_sheet(summary)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _sheet_styles():
    header_font = Font(name='宋体', size=11, bold=True)
    header_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2',
                              fill_type='solid')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='000000')
    return header_font, header_fill, center, Border(left=thin, right=thin,
                                                    top=thin, bottom=thin)


def _fill_channel_sheet(ws):
    header_font, header_fill, center, border = _sheet_styles()
    last_col = len(MODULE_HEADERS)

    title = ws.cell(row=1, column=1, value='信道维护记录表')
    title.font = Font(name='宋体', size=14, bold=True)
    title.alignment = center
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.row_dimensions[1].height = 30

    for col_idx, header in enumerate(MODULE_HEADERS, 1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[2].height = HEADER_ROW_HEIGHT

    row_idx = 3
    for channel_no, (main_radio, backup_radio) in enumerate(CHANNEL_ROWS, 1):
        ws.merge_cells(start_row=row_idx, start_column=1,
                       end_row=row_idx + 1, end_column=1)
        anchor = ws.cell(row=row_idx, column=1)
        anchor.value = '第%d信道' % channel_no
        anchor.alignment = center
        anchor.border = border
        for offset, radio in ((0, main_radio), (1, backup_radio)):
            radio_cell = ws.cell(row=row_idx + offset, column=2, value=radio)
            radio_cell.alignment = center
            for col_idx in range(1, last_col + 1):
                ws.cell(row=row_idx + offset, column=col_idx).border = border
            ws.row_dimensions[row_idx + offset].height = DATA_ROW_HEIGHT
        row_idx += 2

    for letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[letter].width = width

    # 打印版式：A4 纵向、按宽度缩放、水平居中、自定义页边距与页眉页脚
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.print_area = 'A1:H%d' % (row_idx - 1)
    ws.page_margins.left = 0.6
    ws.page_margins.right = 0.6
    ws.page_margins.top = 0.8
    ws.page_margins.bottom = 0.8
    ws.oddHeader.center.text = '信道维护记录（验收样例）'
    ws.oddFooter.center.text = '第 &P 页 / 共 &N 页'


def _fill_summary_sheet(ws):
    header_font, header_fill, center, border = _sheet_styles()
    ws.cell(row=1, column=1, value='维护结论')
    ws.cell(row=1, column=2, value='内容（填写）')
    for col_idx in (1, 2):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    for row_idx, label in ((2, '维护日期'), (3, '维护人员'), (4, '总体结论')):
        label_cell = ws.cell(row=row_idx, column=1, value=label)
        label_cell.border = border
        label_cell.alignment = Alignment(horizontal='left', vertical='center')
        ws.cell(row=row_idx, column=2).border = border
        ws.row_dimensions[row_idx].height = 22
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 24
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = 9
    ws.print_area = 'A1:B4'
