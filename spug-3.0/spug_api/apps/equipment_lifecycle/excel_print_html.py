# Copyright: (c) OpenSpug Organization. https://github.com/openspug
# Released under the AGPL-3.0 License.
"""设备全生命周期——Excel 台账统一打印 HTML 渲染（2026-10-21）。

把版本快照（布局 meta）+ 填写值 + build_print_plan 的打印计划渲染成
一份确定性 HTML，交由无头 Chromium（excel_browser_pdf_service）打印为
PDF。渲染规则与网页 LayoutTable 完全同源（同一换算常量、同一字号/
内边距/边框/合并语义），这是"网页翻阅 = 导出 PDF = 实际打印"的基础：

- 纸张：每页一个 .page div，尺寸 = 计划中的真实纸张 mm，@page margin 0，
  Chromium 按页原样输出（单页尺寸固定，内容不撑页）；
- 单元格：padding 2px 4px、字号 pt（缺省 11pt 与网页一致）、边框 1px
  solid、white-space pre/pre-wrap（与网页 cellCss 相同语义）、垂直居中
  （与网页 td 默认一致，非 Excel 的底对齐——以网页视觉基准为准）；
- 合并：按页窗口裁剪 colSpan/rowSpan，锚点不在本页窗口时渲染空合并格
  （保边框），跨页合并由打印计划保证只在块高于整页时发生；
- 值：绑定字段以填写值覆盖模板文本（与 Excel 导出同一 _prepare_writes
  归一化口径）；文本保留空格与换行，HTML 转义防注入；
- 打印区域参考线（网页虚线）不属于本渲染器，绝不进入 PDF。
"""
import html as html_lib

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from apps.equipment_lifecycle.excel_print_layout import (
    DEFAULT_FONT_SIZE_PT,
    column_width_px,
    row_height_px,
)

# 网页与打印共用的字体栈（两端同一策略，确保中文正常显示与近似字宽）：
# 容器安装 fonts-noto-cjk 后落到 Noto Sans CJK SC。
BASE_FONT_STACK = (
    '"Segoe UI", "Microsoft YaHei", "PingFang SC", "Noto Sans CJK SC", '
    '"WenQuanYi Micro Hei", "Helvetica Neue", Arial, sans-serif')
SERIF_FONT_NAMES = ('宋体', 'SimSun', 'NSimSun', 'Times New Roman', '楷体',
                    'KaiTi', '仿宋', 'FangSong')
_TD_PADDING = '2px 4px'


def _css_color(rgb):
    if not rgb or not isinstance(rgb, str):
        return None
    if len(rgb) == 8:
        return '#' + rgb[2:]
    if len(rgb) == 6:
        return '#' + rgb
    return None


def _font_family(name):
    """单元格字体名 → CSS font-family（保留原名 + 同类中文字体兜底）。"""
    if not name or not isinstance(name, str):
        return BASE_FONT_STACK
    safe = name.replace('"', '')
    if any(marker in safe for marker in SERIF_FONT_NAMES):
        return '"%s", "SimSun", "Noto Serif CJK SC", "Noto Serif SC", serif' % safe
    return '"%s", "Microsoft YaHei", "Noto Sans CJK SC", "WenQuanYi Micro Hei", sans-serif' % safe


def _font_css(css, font):
    """字体相关样式（字号/字体族/加粗斜体/下划线删除线/字色）。"""
    if font.get('size'):
        css['font-size'] = '%gpt' % float(font['size'])
    if font.get('name'):
        css['font-family'] = _font_family(font['name'])
    if font.get('bold'):
        css['font-weight'] = 'bold'
    if font.get('italic'):
        css['font-style'] = 'italic'
    decorations = []
    if font.get('underline'):
        decorations.append('underline')
    if font.get('strike'):
        decorations.append('line-through')
    if decorations:
        css['text-decoration'] = ' '.join(decorations)
    color = _css_color(font.get('color'))
    if color:
        css['color'] = color


def _align_border_css(css, style):
    """对齐与边框样式（h/v/wrap 语义与网页 cellCss 相同）。"""
    align = style.get('align') or {}
    if align.get('h'):
        css['text-align'] = str(align['h'])
    if align.get('v'):
        css['vertical-align'] = str(align['v'])
    if align.get('wrap'):
        css['white-space'] = 'pre-wrap'
        css['word-break'] = 'break-word'
    border = style.get('border') or {}
    for side in ('top', 'bottom', 'left', 'right'):
        part = border.get(side)
        if part and part.get('style'):
            color = _css_color(part.get('color')) or '#000000'
            css['border-%s' % side] = '1px solid %s' % color


def _cell_css(cell_style, value):
    """单元格内联样式（与网页 LayoutTable.cellCss 同一映射语义）。"""
    css = {
        'padding': _TD_PADDING,
        'overflow': 'hidden',
        'vertical-align': 'middle',
        'font-size': '%gpt' % DEFAULT_FONT_SIZE_PT,
        # 行高 1：单行文本高度 = 字号，行高不被字体行框撑大（Excel 行高
        # 是精确值；两端同一规则，几何与模板行高逐像素一致）
        'line-height': '1',
        'white-space': 'pre',
    }
    style = cell_style or {}
    _font_css(css, font=style.get('font') or {})
    fill = _css_color(style.get('fill'))
    if fill:
        css['background'] = fill
    _align_border_css(css, style)
    return css


def _style_attr(css):
    return ';'.join('%s:%s' % (key, _escape(str(value)))
                    for key, value in css.items())


def _escape(text):
    return html_lib.escape(str(text), quote=True)


def _cell_text(cell, value):
    """展示文本：绑定字段以填写值覆盖模板文本（值归一化在导出层完成）。"""
    if value is not None and value != '':
        return str(value)
    if cell is None:
        return ''
    raw = cell.get('v')
    return '' if raw is None else str(raw)


def _row_occupancy(sheet, row, values_map, col_start, col_end):
    """行内单元格占用表（Excel 溢出语义：有内容才占用，仅边框不算）。"""
    occupied = {}
    for col in range(col_start, col_end + 1):
        coordinate = '%s%d' % (get_column_letter(col), row)
        cell = (sheet.get('cells') or {}).get(coordinate) or {}
        raw = cell.get('v')
        if raw is not None and raw != '':
            occupied[col] = True
        elif values_map.get(coordinate) not in (None, ''):
            occupied[col] = True
    return occupied


def _overflow_allowed(sheet, row, col, col_end, values_map, style):
    """非换行且有文本的单元格，右侧直到窗口尽头均为空格时允许向右溢出
    （与 Excel 显示语义一致，避免长文本被列宽静默裁切）。"""
    align = (style or {}).get('align') or {}
    if align.get('wrap'):
        return False
    occupied = _row_occupancy(sheet, row, values_map, col + 1, col_end)
    return not any(occupied.values())


def _page_merge_windows(sheet, rows, col_start, col_end):
    """合并区域与页窗口的交集：渲染格 = 交集左上，跨度 = 交集大小。

    返回 {(row, col): {'row_span', 'col_span', 'is_anchor'}}（is_anchor
    表示交集左上即模板锚点，承载文本；否则渲染空合并格保边框）；
    covered：交集内被覆盖的坐标集合。
    """
    row_set = set(rows)
    spans = {}
    covered = set()
    for merged in sheet.get('merged_ranges') or []:
        try:
            c1, r1, c2, r2 = range_boundaries(str(merged))
        except Exception:
            continue
        rs = [r for r in range(r1, r2 + 1) if r in row_set]
        cs = [c for c in range(c1, c2 + 1) if col_start <= c <= col_end]
        if not rs or not cs:
            continue
        top, left = rs[0], cs[0]
        spans[(top, left)] = {
            'row_span': len(rs), 'col_span': len(cs),
            'is_anchor': (top, left) == (r1, c1),
        }
        for row in rs:
            for col in cs:
                if (row, col) != (top, left):
                    covered.add((row, col))
    return spans, covered


def _render_row(sheet, row, widths, spans, covered, values_map, heights):
    """渲染一行：td 定高（行高精确值，合并锚点 = 跨行高度和），
    配合 overflow hidden 保证几何不被内容撑大（与 Excel 行高语义一致）。"""
    cells = []
    for col, width in widths:
        coordinate = '%s%d' % (get_column_letter(col), row)
        if (row, col) in covered:
            continue
        span = spans.get((row, col))
        cell = (sheet.get('cells') or {}).get(coordinate)
        # values_map 已按工作表过滤、以单元格坐标为键（render_print_html）
        value = values_map.get(coordinate)
        css = _cell_css(cell.get('s') if cell else None, value)
        row_height = heights.get(row, 0.0)
        if span and span['row_span'] > 1:
            row_height = round(sum(
                heights.get(r, 0.0)
                for r in range(row, row + span['row_span'])), 2)
        css['height'] = '%gpx' % row_height
        text = _cell_text(cell, value)
        # Excel 溢出语义：右侧为连续空格时长文本向右溢出显示（z-index
        # 提升绘制层级，盖过右侧空格的边框；页面边缘由内容盒裁剪）
        if text and not span and _overflow_allowed(
                sheet, row, col, widths[-1][0], values_map,
                cell.get('s') if cell else None):
            css['overflow'] = 'visible'
            css['position'] = 'relative'
            css['z-index'] = '1'
        attrs = ['style="%s"' % _style_attr(css)]
        if span:
            attrs.append('colspan="%d"' % span['col_span'])
            attrs.append('rowspan="%d"' % span['row_span'])
        cells.append('<td %s>%s</td>' % (' '.join(attrs), _escape(text)))
    row_height = heights.get(row, 0.0)
    return '<tr style="height:%gpx">%s</tr>' % (row_height, ''.join(cells))


def _render_page_table(sheet, page, scale):
    """渲染单页窗口内的表格（列宽/行高/字号按缩放等比，数值与网页同源）。"""
    col_start, col_end = page['col_start'], page['col_end']
    hidden_cols = set(sheet.get('hidden_columns') or [])
    col_widths = sheet.get('column_widths') or {}
    widths = []
    table_width = 0.0
    for col in range(col_start, col_end + 1):
        letter = get_column_letter(col)
        if letter in hidden_cols:
            continue
        width = round(column_width_px(col_widths.get(letter)) * scale, 2)
        widths.append((col, width))
        table_width += width
    heights_map = sheet.get('row_heights') or {}
    hidden_rows = set(sheet.get('hidden_rows') or [])

    def row_h(row):
        if str(row) in hidden_rows:
            return 0.0
        return round(row_height_px(heights_map.get(str(row))) * scale, 2)

    repeat_rows = page.get('repeat_rows') or []
    content_rows = list(range(page['row_start'], page['row_end'] + 1))
    all_rows = repeat_rows + content_rows
    table_height = round(sum(row_h(row) for row in all_rows), 2)
    heights = {row: row_h(row) for row in all_rows}
    spans, covered = _page_merge_windows(sheet, all_rows, col_start, col_end)
    rows_html = []
    for row in all_rows:
        rows_html.append(_render_row(sheet, row, widths, spans, covered,
                                     page.get('values_map') or {}, heights))
    cols_html = ''.join(
        '<col style="width:%gpx"/>' % width for _col, width in widths)
    return table_width, table_height, (
        '<table class="el-print-table" style="width:%gpx">'
        '<colgroup>%s</colgroup><tbody>%s</tbody></table>'
        % (table_width, cols_html, ''.join(rows_html)))


def _plan_geometry(plan):
    """计划 → 页面几何（页宽高/边距 mm 与可打印区 px、居中标志）。"""
    paper = plan.get('paper') or {}
    margins = plan.get('margins_mm') or {}
    page_w = float(paper.get('width_mm') or 210)
    page_h = float(paper.get('height_mm') or 297)
    margin_top = float(margins.get('top') or 0)
    margin_bottom = float(margins.get('bottom') or 0)
    margin_left = float(margins.get('left') or 0)
    margin_right = float(margins.get('right') or 0)
    return {
        'page_w': page_w, 'page_h': page_h,
        'margin_top': margin_top, 'margin_left': margin_left,
        'printable_w_px': max(
            2.0, (page_w - margin_left - margin_right) * 96 / 25.4),
        'printable_h_px': max(
            2.0, (page_h - margin_top - margin_bottom) * 96 / 25.4),
        'center_h': bool(plan.get('horizontal_centered')),
        'center_v': bool(plan.get('vertical_centered')),
    }


def _sheet_values(values_by_cell, sheet_name):
    """按工作表过滤坐标值映射（键收敛为坐标）。"""
    return {
        key[1]: value for key, value in (values_by_cell or {}).items()
        if key[0] == sheet_name}


def _sheet_page_divs(sheet, sheet_plan, values_map, geometry):
    """渲染单个工作表的全部页 div（缩放/居中偏移在此应用）。"""
    divs = []
    scale = float(sheet_plan.get('scale') or 1.0)
    for page in sheet_plan.get('pages') or []:
        page = dict(page, values_map=values_map)
        table_width, table_height, table_html = _render_page_table(
            sheet, page, scale)
        offset_x = max(0.0, (geometry['printable_w_px'] - table_width) / 2) \
            if geometry['center_h'] else 0.0
        offset_y = max(0.0, (geometry['printable_h_px'] - table_height) / 2) \
            if geometry['center_v'] else 0.0
        divs.append(_render_page_div(
            geometry['page_w'], geometry['page_h'], geometry['margin_top'],
            geometry['margin_left'], offset_x, offset_y, table_html))
    return divs


def render_print_html(meta, values_by_cell, plan):
    """渲染统一打印 HTML。

    meta：版本快照形态（layout.sheets 提供单元格与样式）；values_by_cell：
    {(sheet_name, 'A1'): 值}；plan：excel_print_layout.build_print_plan
    的打印计划（纸张/边距/缩放/分页唯一事实来源）。
    """
    sheets_by_name = {
        (sheet.get('name') or ''): sheet
        for sheet in ((meta or {}).get('layout') or {}).get('sheets') or []
    }
    geometry = _plan_geometry(plan)
    body_pages = []
    for sheet_name, sheet_plan in (plan.get('sheets') or {}).items():
        sheet = sheets_by_name.get(sheet_name)
        if sheet is None:
            continue
        body_pages.extend(_sheet_page_divs(
            sheet, sheet_plan, _sheet_values(values_by_cell, sheet_name),
            geometry))
    if not body_pages:
        body_pages.append('<div class="el-print-page"></div>')
    return _document(
        geometry['page_w'], geometry['page_h'],
        geometry['printable_w_px'] / 96 * 25.4,
        geometry['printable_h_px'] / 96 * 25.4, body_pages)


def _render_page_div(page_w, page_h, margin_top, margin_left, offset_x,
                     offset_y, table_html):
    """单页 div：绝对定位于页边距原点，偏移实现水平/垂直居中。"""
    return (
        '<div class="el-print-page" style="width:%gmm;height:%gmm">'
        '<div class="el-print-content" style="top:%gmm;left:%gmm">'
        '<div class="el-print-table-wrap" style="margin:%gpx 0 0 %gpx">%s</div>'
        '</div></div>'
        % (page_w, page_h, margin_top, margin_left, offset_y, offset_x,
           table_html))


def _document(page_w, page_h, printable_w_mm, printable_h_mm, body_pages):
    css = (
        '@page { size: %gmm %gmm; margin: 0; }\n'
        'html, body { margin: 0; padding: 0; }\n'
        'body { font-family: %s; -webkit-print-color-adjust: exact; '
        'print-color-adjust: exact; }\n'
        '.el-print-page { position: relative; overflow: hidden; '
        'page-break-after: always; background: #fff; }\n'
        '.el-print-page:last-child { page-break-after: auto; }\n'
        '.el-print-content { position: absolute; width: %gmm; height: %gmm; '
        'overflow: hidden; }\n'
        '.el-print-table { border-collapse: collapse; table-layout: fixed; }\n'
        '.el-print-table td { padding: %s; overflow: hidden; box-sizing: '
        'border-box; }\n'
        % (page_w, page_h, BASE_FONT_STACK, printable_w_mm, printable_h_mm,
           _TD_PADDING))
    return (
        '<!DOCTYPE html><html><head><meta charset="utf-8">'
        '<title>维护台账打印</title><style>%s</style></head><body>%s</body>'
        '</html>' % (css, ''.join(body_pages)))
