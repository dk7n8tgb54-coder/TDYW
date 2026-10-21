# Copyright: (c) OpenSpug Organization. https://github.com/openspug
# Released under the AGPL-3.0 License.
"""设备全生命周期——Excel 上传版维护台账模板核心服务（2026-09-28）。

由验证原型 apps/ledger_prototype/prototype_services.py 提取的稳定逻辑
（正式路线唯一实现，原型经兼容外壳委托到本模块）：

- parse_workbook：openpyxl 解析 .xlsx 布局——合并单元格、行高、列宽、
  单元格值与基本样式（字体/填充/对齐/边框/数字格式）；
- validate_bindings / validate_values：字段绑定与填写值校验（规则见下）；
- export_filled_excel / export_pdf：以服务端保存的 meta + 源文件为唯一事实，
  导出前重新读取源文件并校验合并关系，不信任前端提交的布局；PDF 转换
  前净化打印区域（清除不覆盖内容的残留打印区域，防 LibreOffice 输出
  整份空白），导出 Excel 与源文件不做任何改写；
- preview_blank_pdf：模板草稿空白预览 PDF（不写入填写值，未绑定字段
  保持空白），同样经布局校验与打印区域净化后转换，只读源文件；
- preview_pdf：预览通用入口，可携带临时样例值（validate_values 校验，
  只写预览副本）与未保存打印参数（显式校验后仅本次生效），供打印
  设置"应用当前设置并预览"使用，不写库、不改草稿与版本；
- assess_compatibility：对草稿源文件做单次解析的兼容性评估（不支持
  特性 / 打印区域覆盖），供模板详情返回 status/warnings/unsupported；
- collect_unsupported_features：上传前显式检测平台不支持的文件特性
  （公式/图表/图片/外部链接/透视表），任一存在即拒绝导入，绝不静默丢失。

绑定规则：
- 一个单元格最多绑定一个字段；一个字段只能绑定一个单元格（双向唯一）；
- 合并区域只能绑定左上角单元格；
- 已有内容（标题/表头）的单元格禁止绑定为填写字段；
- field_id 必须是 UUID，不由坐标充当、后端不派生不改写（批量绑定由
  前端为每个单元格生成独立 UUID，服务端逐项校验唯一性）；
- 字段绑定协议只含 field_id/sheet_name/cell_coordinate/field_name/
  input_type/is_required/unit/options 八个键；default_value/
  operation_guide/abnormal_rule/maintenance_cycles 已于 2026-09-29
  退役：新保存的绑定不写入、校验不读取，历史草稿/版本快照中的旧属性
  原样保留（读取端兼容，填写与导出不受影响）；
- 仅支持固定行列、固定填写位置的 .xlsx；不支持动态增行、公式计算、
  宏、图表、图片、外部链接与 Word 模板。
"""
import hashlib
import io
import os
import tempfile
import uuid
from datetime import datetime, date
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string, range_boundaries

from apps.equipment_lifecycle import excel_pdf_service

PROTOCOL_VERSION = 1
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 500
MAX_COLUMNS = 100
MAX_BINDINGS = 300

INPUT_TYPE_TEXT = 'text'
INPUT_TYPE_NUMBER = 'number'
INPUT_TYPE_CHOICE = 'choice'
INPUT_TYPE_STATUS = 'status'
INPUT_TYPE_DATE = 'date'
INPUT_TYPES = {
    INPUT_TYPE_TEXT: '文本',
    INPUT_TYPE_NUMBER: '数值',
    INPUT_TYPE_CHOICE: '单选',
    INPUT_TYPE_STATUS: '正常/异常',
    INPUT_TYPE_DATE: '日期',
}
STATUS_OPTIONS = ['正常', '异常']

XLSX_SUFFIX = '.xlsx'

UNSUPPORTED_FEATURE_LABELS = {
    'formula': '公式计算',
    'chart': '图表',
    'image': '图片',
    'external_link': '外部链接',
    'pivot': '透视表',
}


class ExcelLayoutError(Exception):
    """业务校验错误，View 层捕获后以 HTTP 200 + error 返回。"""


# ==================== 不支持特性检测 ====================

def collect_unsupported_features(data):
    """解析前检测平台不支持的文件特性，返回特性键列表（供明确拒绝提示）。

    openpyxl 重写 .xlsx 时公式缓存值、图表、图片、透视表会丢失或失真，
    因此不允许静默导入；检测通过私有属性兼容 openpyxl 3.x。
    """
    try:
        workbook = load_workbook(io.BytesIO(data))
    except Exception:
        return []
    features = set()
    for worksheet in workbook.worksheets:
        features |= _sheet_unsupported_features(worksheet)
    if getattr(workbook, '_external_links', None):
        features.add('external_link')
    return sorted(features, key=UNSUPPORTED_FEATURE_LABELS.get)


def _sheet_unsupported_features(worksheet):
    """单个工作表内不支持特性检测（公式/图表/图片/透视表）。"""
    features = set()
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.data_type == 'f':
                features.add('formula')
    if getattr(worksheet, '_charts', None):
        features.add('chart')
    if getattr(worksheet, '_images', None):
        features.add('image')
    if getattr(worksheet, '_pivots', None):
        features.add('pivot')
    return features


def describe_unsupported_features(features):
    return '、'.join(UNSUPPORTED_FEATURE_LABELS[key] for key in features)


# ==================== 布局解析 ====================

def parse_workbook(filename, data):
    """解析上传的 .xlsx，返回模板 meta（不含租户/创建人等登记信息）。"""
    from apps.equipment_lifecycle.excel_pdf_params import default_pdf_params
    if not filename or not filename.lower().endswith(XLSX_SUFFIX):
        raise ExcelLayoutError('仅支持 .xlsx 文件（暂不支持 .xls 与 WPS 私有格式）')
    try:
        wb = load_workbook(io.BytesIO(data))
    except Exception:
        raise ExcelLayoutError('文件解析失败，请确认上传的是有效的 .xlsx 文件')
    if not wb.worksheets:
        raise ExcelLayoutError('工作簿中没有工作表')
    sheets = [_parse_sheet(ws) for ws in wb.worksheets]
    active = _active_sheet_name(wb, sheets)
    return {
        'source_filename': filename,
        'source_file_hash': hashlib.sha256(data).hexdigest(),
        'sheet_name': active,
        'layout': {'active_sheet': active, 'sheets': sheets},
        'field_bindings': [],
        'sample_values': {},
        'pdf_params': default_pdf_params(wb),
        'protocol_version': PROTOCOL_VERSION,
    }


def _active_sheet_name(workbook, sheets):
    active = workbook.active
    if active is not None and any(s['name'] == active.title for s in sheets):
        return active.title
    return sheets[0]['name']


def _parse_sheet(ws):
    if ws.max_row > MAX_ROWS or ws.max_column > MAX_COLUMNS:
        raise ExcelLayoutError(
            '工作表「%s」超出支持范围（最多 %d 行 × %d 列）'
            % (ws.title, MAX_ROWS, MAX_COLUMNS))
    merged_ranges, covered = _merged_layout(ws)
    column_widths, hidden_columns = column_layout(ws)
    return {
        'name': ws.title,
        'max_row': ws.max_row,
        'max_column': ws.max_column,
        'merged_ranges': merged_ranges,
        'row_heights': _row_heights(ws),
        'column_widths': column_widths,
        'hidden_rows': _hidden_rows(ws),
        'hidden_columns': hidden_columns,
        'cells': _sheet_cells(ws, covered),
        'print': _print_setup(ws),
    }


def _merged_layout(ws):
    """返回 (合并区域字符串列表, 被合并覆盖的空格坐标集合)。"""
    merged_ranges = []
    covered = set()
    for rng in ws.merged_cells.ranges:
        merged_ranges.append(str(rng))
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                if (row, col) != (rng.min_row, rng.min_col):
                    covered.add((row, col))
    return merged_ranges, covered


def _row_heights(ws):
    heights = {}
    for idx in range(1, ws.max_row + 1):
        height = ws.row_dimensions[idx].height
        if height is not None:
            heights[str(idx)] = round(float(height), 2)
    return heights


def _default_column_width(ws):
    """工作表显式声明的缺省列宽（sheetFormatPr.defaultColWidth；未声明为 None）。"""
    try:
        width = ws.sheet_format.defaultColWidth
    except AttributeError:
        return None
    if width is None:
        return None
    try:
        return round(float(width), 2)
    except (TypeError, ValueError):
        return None


def _dimension_column_range(dim, max_column):
    """列维度覆盖的 (起始列, 结束列)，钳制到内容区（1..max_column）。

    读取文件时 openpyxl 把 <col min="1" max="3"> 的范围记在维度的
    min/max 属性上（只按起始列落键）；程序构造的维度（min/max 为空）
    按 index 字母回退为单列。无法定位返回 (None, None)。
    """
    min_col = dim.min
    if not min_col:
        try:
            min_col = column_index_from_string(dim.index)
        except (ValueError, TypeError, AttributeError):
            return None, None
    max_col = dim.max or min_col
    if max_col < min_col:
        max_col = min_col
    return max(1, min_col), min(max_col, max_column)


def column_layout(ws):
    """列宽与隐藏列（一次遍历；展开 <col> 的 min/max 覆盖范围）。

    - openpyxl 读取 <col min="1" max="3" width="36"/> 时只按起始列落键
      （范围记在 min/max 属性），B/C 在 column_dimensions 中不存在；
      按字母下标访问缺失列会用 openpyxl 内置缺省宽度凭空造出 width=13
      的假维度——因此只遍历已存在的维度并展开范围，单列（min=max）
      自然覆盖；
    - 未设显式宽度的列使用 sheetFormatPr.defaultColWidth；Excel 未声明
      时的内置缺省 8.43 由前端兜底，不在此伪造；
    - 重叠范围按 XML 文档顺序后者覆盖前者（与 Excel 行为一致）；
    - 范围超出内容区（max_column）的部分忽略（渲染不出无内容列）。
    返回 (column_widths dict, hidden_columns list)。公开供
    excel_pdf_params 内容宽度估算复用（同一 min/max 展开语义）。
    """
    max_column = ws.max_column
    covered = {}
    for dim in list(ws.column_dimensions.values()):
        min_col, max_col = _dimension_column_range(dim, max_column)
        if min_col is None:
            continue
        for idx in range(min_col, max_col + 1):
            covered[idx] = dim
    widths = {idx: round(float(dim.width), 2)
              for idx, dim in covered.items() if dim.width is not None}
    default_width = _default_column_width(ws)
    if default_width is not None:
        for idx in range(1, max_column + 1):
            widths.setdefault(idx, default_width)
    return ({get_column_letter(idx): widths[idx] for idx in sorted(widths)},
            [get_column_letter(idx) for idx in sorted(covered)
             if covered[idx].hidden])


def _hidden_rows(ws):
    return [str(idx) for idx in range(1, ws.max_row + 1)
            if ws.row_dimensions[idx].hidden]


def refresh_column_layout(layout, source_path):
    """以源文件为准修正布局中的列宽与隐藏列（读取端兼容，快照不可变）。

    历史草稿/版本快照由旧解析逻辑生成：<col min max> 范围只落在起始列，
    其余列被误记为 openpyxl 缺省宽度 13、隐藏范围漏报。这里用仍留存的
    源文件重新推导 column_widths/hidden_columns，返回修正后的布局；
    布局非法或源文件缺失/解析失败时原样返回（详情不因修正失败阻断，
    导出路径另有源文件校验与报错）。修正幂等：新解析逻辑产出的快照
    修正前后一致，已发布版本无需重建。
    """
    if not isinstance(layout, dict):
        return layout
    try:
        data = _read_source(source_path)
        workbook = load_workbook(io.BytesIO(data))
    except Exception:
        return layout
    worksheets = {ws.title: ws for ws in workbook.worksheets}
    patched = []
    changed = False
    for sheet_meta in layout.get('sheets') or []:
        worksheet = sheet_meta.get('name') if isinstance(sheet_meta, dict) else None
        if worksheet not in worksheets:
            patched.append(sheet_meta)
            continue
        widths, hidden = column_layout(worksheets[worksheet])
        updated = dict(sheet_meta, column_widths=widths, hidden_columns=hidden)
        changed = changed or updated != sheet_meta
        patched.append(updated)
    if not changed:
        return layout
    return {**layout, 'sheets': patched}


def refresh_print_layout(layout, source_path):
    """以源文件为准补全布局快照的打印设置（读取端修正，快照不可变）。

    历史版本快照的 print 段只有 orientation/paper_size/fit（2026-09-28
    解析口径），缺 scale/page_margins/print_area——浏览器打印引擎需要
    物理边距与打印区域才能复现源文件版式。这里对缺失的键用仍留存的
    源文件重新推导补齐（只补缺，不覆盖已有键，幂等：新解析逻辑产出的
    快照补齐前后一致）；布局非法或源文件缺失/解析失败时原样返回，
    不影响详情与导出。
    """
    if not isinstance(layout, dict):
        return layout
    try:
        data = _read_source(source_path)
        workbook = load_workbook(io.BytesIO(data))
    except Exception:
        return layout
    worksheets = {ws.title: ws for ws in workbook.worksheets}
    patched = []
    changed = False
    for sheet_meta in layout.get('sheets') or []:
        worksheet = sheet_meta.get('name') if isinstance(sheet_meta, dict) else None
        if worksheet not in worksheets:
            patched.append(sheet_meta)
            continue
        current = dict(sheet_meta.get('print') or {})
        updated = _print_setup(worksheets[worksheet])
        merged = dict(current)
        for key, value in updated.items():
            if key not in merged or merged.get(key) in (None, '', {}):
                merged[key] = value
        if merged != current:
            changed = True
            patched.append(dict(sheet_meta, print=merged))
        else:
            patched.append(sheet_meta)
    if not changed:
        return layout
    return {**layout, 'sheets': patched}


def _sheet_cells(ws, covered):
    """收集有值或有样式的单元格（被合并覆盖的空格无独立展示意义）。"""
    cells = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            value = _cell_value(cell.value)
            has_value = value is not None and value != ''
            if (cell.row, cell.column) in covered and not has_value:
                continue
            style = _cell_style(cell)
            if not has_value and style is None:
                continue
            cells[cell.coordinate] = {'v': value, 's': style}
    return cells


def _print_setup(ws):
    """工作表打印设置快照（统一打印布局的源文件口径）。

    orientation/paper_size/fit_to_width/fit_to_height 为 2026-09-28 既有
    键；2026-10-21 补充 scale（百分比缩放）、page_margins（物理边距，
    inch→cm）与 print_area（净化 $ 与表名前缀），供浏览器打印引擎在
    不经 LibreOffice 的情况下复现源文件打印版式。
    """
    page_setup = ws.page_setup
    setup = {
        'orientation': page_setup.orientation or 'default',
        'paper_size': page_setup.paperSize,
        'fit_to_width': page_setup.fitToWidth,
        'fit_to_height': page_setup.fitToHeight,
        'scale': page_setup.scale,
    }
    margins = ws.page_margins
    if margins is not None:
        setup['page_margins'] = {
            key: round(float(getattr(margins, key)) * 2.54, 2)
            for key in ('top', 'bottom', 'left', 'right')
            if getattr(margins, key) is not None
        }
    parts = []
    for part in (ws.print_area or []):
        text = str(part).strip().strip("'")
        if '!' in text:
            text = text.rsplit('!', 1)[-1]
        if text:
            parts.append(text.replace('$', ''))
    setup['print_area'] = ','.join(parts)
    return setup


def _cell_value(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d %H:%M:%S') if isinstance(value, datetime) \
            else value.strftime('%Y-%m-%d')
    if isinstance(value, Decimal):
        return float(value)
    return str(value)


def _cell_style(cell):
    """提取基本样式；完全默认则返回 None（压缩 payload）。"""
    border = _border_part(cell.border)
    fill = _fill_part(cell.fill)
    font = _font_part(cell.font)
    align = _align_part(cell.alignment)
    number_format = cell.number_format if cell.number_format != 'General' else None
    if not border and fill is None and font is None and align is None \
            and number_format is None:
        return None
    return {
        'font': font,
        'fill': fill,
        'align': align,
        'border': border,
        'number_format': number_format,
    }


def _border_part(border):
    part = {}
    for side_name in ('top', 'bottom', 'left', 'right'):
        side = getattr(border, side_name, None)
        if side is not None and side.style:
            part[side_name] = {'style': side.style, 'color': _color_str(side.color)}
    return part


def _fill_part(fill):
    if fill is None or fill.fill_type in (None, 'none'):
        return None
    return _color_str(fill.fgColor) or _color_str(fill.start_color)


def _font_part(font):
    if font is None:
        return None
    decorated = (
        font.bold or font.italic or font.underline or font.strike
        or (font.name and font.name != 'Calibri')
        or (font.size and float(font.size) != 11.0))
    if not decorated:
        return None
    return {
        'name': font.name,
        'size': float(font.size) if font.size else None,
        'bold': bool(font.bold),
        'italic': bool(font.italic),
        'underline': font.underline or None,
        'strike': bool(font.strike),
        'color': _color_str(font.color),
    }


def _align_part(align):
    if align is None or not (align.horizontal or align.vertical or align.wrap_text):
        return None
    return {
        'h': align.horizontal,
        'v': align.vertical,
        'wrap': bool(align.wrap_text),
    }


def _color_str(color):
    if color is None:
        return None
    try:
        rgb = color.rgb
        if isinstance(rgb, str) and rgb:
            return rgb
    except Exception:
        pass
    return None


# ==================== 字段绑定校验 ====================

def validate_bindings(meta, bindings):
    """校验并归一化字段绑定列表；非法直接抛 ExcelLayoutError。"""
    if bindings is None:
        bindings = []
    if not isinstance(bindings, list):
        raise ExcelLayoutError('field_bindings 必须是列表')
    if len(bindings) > MAX_BINDINGS:
        raise ExcelLayoutError('字段绑定数量超过上限（%d）' % MAX_BINDINGS)
    sheets_by_name = {s['name']: s for s in meta['layout']['sheets']}
    seen_fields = {}
    seen_cells = set()
    return [_validate_binding(index, raw, sheets_by_name, seen_fields, seen_cells)
            for index, raw in enumerate(bindings, 1)]


def _validate_binding(index, raw, sheets_by_name, seen_fields, seen_cells):
    label = '第%d个字段' % index
    if not isinstance(raw, dict):
        raise ExcelLayoutError('%s 格式不正确' % label)
    field_id = _binding_field_id(raw, label)
    sheet, coordinate, col_idx, row_idx = _binding_cell(raw, sheets_by_name, label)
    attributes = _binding_attributes(raw, label)
    input_type = attributes['input_type']
    _check_uniqueness(field_id, attributes['field_name'], sheet['name'],
                      coordinate, seen_fields, seen_cells, label)
    _check_merged_anchor(sheet, coordinate, col_idx, row_idx, label)
    _check_blank_cell(sheet, coordinate, label)
    seen_fields[field_id] = attributes['field_name']
    seen_cells.add((sheet['name'], coordinate))
    binding = {
        'field_id': field_id,
        'sheet_name': sheet['name'],
        'cell_coordinate': coordinate,
    }
    binding.update(attributes)
    return binding


def _binding_field_id(raw, label):
    field_id = str(raw.get('field_id') or '').strip()
    try:
        uuid.UUID(field_id)
    except (ValueError, AttributeError, TypeError):
        raise ExcelLayoutError('%s 的 field_id 必须是 UUID' % label)
    return field_id


def _binding_cell(raw, sheets_by_name, label):
    """校验工作表与单元格坐标（存在性、格式、范围），返回定位信息。"""
    sheet_name = str(raw.get('sheet_name') or '').strip()
    if sheet_name not in sheets_by_name:
        raise ExcelLayoutError('%s 的工作表「%s」不存在' % (label, sheet_name))
    sheet = sheets_by_name[sheet_name]
    coordinate = str(raw.get('cell_coordinate') or '').strip().upper()
    col_idx, row_idx = _parse_coordinate(coordinate, label)
    if row_idx > sheet['max_row'] or col_idx > sheet['max_column']:
        raise ExcelLayoutError(
            '%s 的单元格 %s 超出工作表范围（%d 行 × %d 列）'
            % (label, coordinate, sheet['max_row'], sheet['max_column']))
    return sheet, coordinate, col_idx, row_idx


def _binding_attributes(raw, label):
    """校验并归一化字段属性（仅新协议八键中的属性四键）。

    请求中携带的 default_value/operation_guide/abnormal_rule/
    maintenance_cycles 一律忽略不落库（历史数据兼容由读取端保证）。
    """
    field_name = str(raw.get('field_name') or '').strip()
    if not field_name:
        raise ExcelLayoutError('%s 缺少字段名称' % label)
    if len(field_name) > 50:
        raise ExcelLayoutError('%s 的字段名称过长（≤50 字）' % label)
    input_type = str(raw.get('input_type') or '').strip()
    if input_type not in INPUT_TYPES:
        raise ExcelLayoutError('%s 的输入类型必须是：%s'
                               % (label, '、'.join(INPUT_TYPES.values())))
    return {
        'field_name': field_name,
        'input_type': input_type,
        'is_required': bool(raw.get('is_required', False)),
        'unit': str(raw.get('unit') or '')[:20],
        'options': _validate_options(input_type, raw.get('options'), label),
    }


def _parse_coordinate(coordinate, label):
    if not coordinate:
        raise ExcelLayoutError('%s 缺少单元格坐标' % label)
    try:
        column_part, row_part = coordinate_from_string(coordinate)
        return column_index_from_string(column_part), int(row_part)
    except Exception:
        raise ExcelLayoutError('%s 的单元格坐标「%s」不合法' % (label, coordinate))


def _validate_options(input_type, options, label):
    if input_type == INPUT_TYPE_STATUS:
        return list(STATUS_OPTIONS)
    if input_type != INPUT_TYPE_CHOICE:
        return []
    if not isinstance(options, list) or not options:
        raise ExcelLayoutError('%s 为单选类型，必须提供候选项' % label)
    cleaned = []
    for option in options:
        text = str(option or '').strip()
        if not text:
            raise ExcelLayoutError('%s 的候选项不能为空' % label)
        if len(text) > 50:
            raise ExcelLayoutError('%s 的候选项过长（≤50 字）' % label)
        if text not in cleaned:
            cleaned.append(text)
    if len(cleaned) > 20:
        raise ExcelLayoutError('%s 的候选项过多（≤20 个）' % label)
    return cleaned


def _check_uniqueness(field_id, field_name, sheet_name, coordinate,
                      seen_fields, seen_cells, label):
    if field_id in seen_fields:
        raise ExcelLayoutError('字段 %s 重复绑定（一个字段只能绑定一个单元格）'
                               % seen_fields[field_id])
    if (sheet_name, coordinate) in seen_cells:
        raise ExcelLayoutError('单元格 %s 已绑定字段（一个单元格最多绑定一个字段）'
                               % coordinate)


def _find_merged_range(sheet, col_idx, row_idx):
    """返回包含该坐标的合并区域边界，无则 None。"""
    for merged in sheet['merged_ranges']:
        min_col, min_row, max_col, max_row = range_boundaries(merged)
        if min_row <= row_idx <= max_row and min_col <= col_idx <= max_col:
            return (min_col, min_row, max_col, max_row, merged)
    return None


def _check_merged_anchor(sheet, coordinate, col_idx, row_idx, label):
    merged = _find_merged_range(sheet, col_idx, row_idx)
    if not merged:
        return
    min_col, min_row, _max_col, _max_row, merged_text = merged
    if (col_idx, row_idx) != (min_col, min_row):
        anchor = '%s%d' % (get_column_letter(min_col), min_row)
        raise ExcelLayoutError(
            '%s 位于合并区域 %s，只能绑定左上角单元格 %s'
            % (label, merged_text, anchor))


def _check_blank_cell(sheet, coordinate, label):
    """已有内容（标题/表头）的单元格不能绑定为填写字段。"""
    cell_data = sheet['cells'].get(coordinate)
    if cell_data is not None and cell_data.get('v') not in (None, ''):
        raise ExcelLayoutError(
            '单元格 %s 已有内容（标题/表头等），不能绑定为填写字段' % coordinate)


# ==================== 填写值校验与导出 ====================

def validate_values(bindings, values):
    """校验填写值（键必须为已绑定 field_id，值类型按字段定义），返回归一化 dict。"""
    if values is None:
        return {}
    if not isinstance(values, dict):
        raise ExcelLayoutError('填写值必须是对象')
    known = {b['field_id']: b for b in bindings}
    normalized = {}
    for field_id, value in values.items():
        binding = known.get(field_id)
        if binding is None:
            raise ExcelLayoutError('填写值包含未绑定的字段，已拒绝')
        if value is None or value == '':
            continue
        normalized[field_id] = _typed_value(binding, value)
    return normalized


def _typed_value(binding, value):
    """按字段类型归一化填写值；文本类型保留原始空格与换行（写入单元格）。

    只有需要精确匹配的受控类型（日期格式、候选项）对匹配副本做空白
    清理；文本值原样返回——单元格内容不因服务端归一化而丢失空格。
    """
    name = binding['field_name']
    input_type = binding['input_type']
    if input_type == INPUT_TYPE_NUMBER:
        return _number_value(value, name)
    text = str(value)
    if input_type == INPUT_TYPE_DATE:
        return _date_value(text.strip(), name)
    if input_type in (INPUT_TYPE_CHOICE, INPUT_TYPE_STATUS):
        stripped = text.strip()
        if stripped not in binding['options']:
            raise ExcelLayoutError('字段「%s」的值必须是候选项之一' % name)
        return stripped
    return text


def _number_value(value, name):
    if isinstance(value, bool):
        raise ExcelLayoutError('字段「%s」的值必须是数字' % name)
    try:
        return float(value)
    except (TypeError, ValueError):
        raise ExcelLayoutError('字段「%s」的值必须是数字' % name)


def _date_value(text, name):
    try:
        datetime.strptime(text, '%Y-%m-%d')
    except ValueError:
        raise ExcelLayoutError('字段「%s」的日期格式必须为 YYYY-MM-DD' % name)
    return text


def export_filled_excel(meta, source_file_path, values, check_required=True):
    """按服务端保存的绑定关系向源模板副本写入填写值，返回 xlsx 字节。

    check_required=False 时不做必填拦截（打印预览允许留空，类型校验保留）。
    """
    writes = _prepare_writes(meta, values, check_required=check_required)
    data = _read_source(source_file_path)
    try:
        workbook = load_workbook(io.BytesIO(data))
    except Exception:
        raise ExcelLayoutError('源模板文件读取失败')
    _verify_layout_matches(workbook, meta)
    for sheet_name, coordinate, value in writes:
        workbook[sheet_name][coordinate] = value
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_pdf(meta, source_file_path, values, converter=None, pdf_params=None,
               check_required=True):
    """填充后转换为 PDF，返回 pdf 字节。

    引擎分派（resolve_engine）：chromium（浏览器打印引擎）走统一打印布局
    —— excel_print_layout 计算分页计划（与网页 print_plan 同源），经
    excel_print_html 渲染统一打印 HTML 后由无头 Chromium 打印，网页与 PDF
    排版规则完全一致；libreoffice 走既有 openpyxl 填值 + headless 转换。

    converter 须提供 convert_xlsx_to_pdf；显式传入时固定走 LibreOffice
    路径（既有测试补丁目标），引擎分派仅在缺省转换器时生效。pdf_params
    显式传入时优先生效（导出高级设置覆盖），否则使用 meta 携带的模板级
    参数（旧版本快照无参数时按内置缺省）。LibreOffice 路径转换前对填写
    结果做打印区域净化（prepare_pdf_workbook）：WPS/Excel 误留的空打印
    区域会被 LibreOffice 严格执行，产生整份空白 PDF。两条路径共用同一
    必填与类型校验（_prepare_writes，check_required=False 供打印预览
    允许留空，与模板预览同口径），导出 Excel 与源文件不做任何改写。
    """
    effective = pdf_params if pdf_params is not None else meta.get('pdf_params')
    if converter is None and excel_pdf_service.resolve_engine() == \
            excel_pdf_service.ENGINE_CHROMIUM:
        return _export_pdf_browser(meta, values, effective,
                                   check_required=check_required)
    filled = prepare_pdf_workbook(
        export_filled_excel(meta, source_file_path, values,
                            check_required=check_required),
        effective, meta.get('sheet_name'))
    return _xlsx_bytes_to_pdf(filled, converter or excel_pdf_service)


def _export_pdf_browser(meta, values, effective, check_required=True):
    """浏览器打印引擎导出：统一打印布局 → 统一打印 HTML → Chromium PDF。"""
    from apps.equipment_lifecycle import (
        excel_browser_pdf_service,
        excel_print_html,
        excel_print_layout,
    )
    # _prepare_writes 返回 [(sheet, coordinate, value)]：与 Excel 导出同一
    # 必填+类型校验口径，这里转坐标映射供 HTML 渲染取值
    values_by_cell = {
        (sheet_name, coordinate): value
        for sheet_name, coordinate, value in _prepare_writes(
            meta, values, check_required=check_required)
    }
    plan = excel_print_layout.build_print_plan(meta, pdf_params=effective)
    html = excel_print_html.render_print_html(meta, values_by_cell, plan)
    return _render_pdf_via_browser(html)


def _render_pdf_via_browser(html):
    """HTML 落临时目录 → Chromium 打印 → 读取 pdf 字节（目录自动清理）。"""
    import tempfile
    from apps.equipment_lifecycle import excel_browser_pdf_service
    with tempfile.TemporaryDirectory(prefix='el_excel_browser_pdf_') as tmp_dir:
        pdf_path = excel_browser_pdf_service.render_html_to_pdf(html, tmp_dir)
        with open(pdf_path, 'rb') as f:
            return f.read()


def preview_blank_pdf(meta, source_file_path, converter=None):
    """模板草稿空白预览 PDF：不写入任何填写值（未填写字段保持空白）。

    以服务端保存的 meta 校验源文件（合并关系一致）并净化打印区域后
    转换；预览即发布后台账未填写时的最终版式基准。只读源文件，
    不改草稿、不创建版本、不产生任何写库副作用。
    """
    return preview_pdf(meta, source_file_path, converter=converter)


def main_sheet_size(meta):
    """主工作表 (行数, 列数)（打印区域/分页参数越界校验；缺失返回 None）。"""
    for sheet_meta in (meta.get('layout') or {}).get('sheets') or []:
        if sheet_meta.get('name') == meta.get('sheet_name'):
            return sheet_meta.get('max_row'), sheet_meta.get('max_column')
    return None, None


def preview_pdf(meta, source_file_path, values=None, pdf_params=None,
                converter=None):
    """模板草稿预览 PDF（可携带临时样例值与未保存参数，均只作用于本次）。

    - values（键为 field_id）经 validate_values 校验（未知 field_id 拒绝、
      按字段类型归一化）后写入预览副本；必填不做检查（预览允许留空），
      空值/无样例时生成空白版式；
    - pdf_params 显式传入时按主工作表尺寸整体校验后优先生效（未保存参数
      预览），缺省使用 meta 携带的模板级参数（发布快照同源）；两者都只
      作用于临时副本：不改草稿源文件、不写库、不创建版本；
    - 引擎分派与 export_pdf 一致：chromium 走统一打印布局（校验通过后
      直接从快照渲染，不经源文件）；libreoffice 走既有 prepare 路径
      （布局校验（合并关系一致）+ 打印区域净化后 LibreOffice 转换），
      两条路径都绝不返回伪造 PDF。
    """
    from apps.equipment_lifecycle.excel_pdf_params import (
        apply_pdf_params,
        validate_pdf_params,
    )
    if pdf_params is not None:
        max_row, max_column = main_sheet_size(meta)
        pdf_params = validate_pdf_params(pdf_params, max_row=max_row,
                                         max_column=max_column)
    else:
        pdf_params = meta.get('pdf_params')
    if excel_pdf_service.resolve_engine() == excel_pdf_service.ENGINE_CHROMIUM:
        return _preview_pdf_browser(meta, values, pdf_params)
    data = _read_source(source_file_path)
    try:
        workbook = load_workbook(io.BytesIO(data))
    except Exception:
        raise ExcelLayoutError('源模板文件读取失败')
    _verify_layout_matches(workbook, meta)
    _apply_preview_values(workbook, meta, values)
    _sanitize_print_areas(workbook)
    apply_pdf_params(workbook, pdf_params, meta.get('sheet_name'))
    output = io.BytesIO()
    workbook.save(output)
    return _xlsx_bytes_to_pdf(output.getvalue(), converter or excel_pdf_service)


def _preview_pdf_browser(meta, values, effective):
    """浏览器打印引擎预览：样例值经 validate_values 归一化后按坐标映射。"""
    from apps.equipment_lifecycle import (
        excel_print_html,
        excel_print_layout,
    )
    normalized = validate_values(meta.get('field_bindings') or [], values or {})
    values_by_cell = {}
    for binding in meta.get('field_bindings') or []:
        value = normalized.get(binding['field_id'])
        if value is None or value == '':
            continue
        values_by_cell[(binding['sheet_name'],
                        binding['cell_coordinate'])] = value
    plan = excel_print_layout.build_print_plan(meta, pdf_params=effective)
    html = excel_print_html.render_print_html(meta, values_by_cell, plan)
    return _render_pdf_via_browser(html)


def _apply_preview_values(workbook, meta, values):
    """把临时样例值写入预览副本（validate_values 校验；空值不写入）。"""
    if not values:
        return
    normalized = validate_values(meta.get('field_bindings') or [], values)
    for binding in meta.get('field_bindings') or []:
        value = normalized.get(binding['field_id'])
        if value is None or value == '':
            continue
        workbook[binding['sheet_name']][binding['cell_coordinate']] = value


def _xlsx_bytes_to_pdf(data, converter):
    """xlsx 字节落临时目录 → LibreOffice 转换 → 读取 pdf 字节（自动清理）。"""
    with tempfile.TemporaryDirectory(prefix='el_excel_pdf_') as tmp_dir:
        xlsx_path = os.path.join(tmp_dir, 'filled.xlsx')
        with open(xlsx_path, 'wb') as f:
            f.write(data)
        pdf_path = converter.convert_xlsx_to_pdf(xlsx_path, tmp_dir)
        with open(pdf_path, 'rb') as f:
            return f.read()


def _print_area_covers_content(worksheet):
    """打印区域是否覆盖至少一个有值单元格（未设打印区域返回 True）。"""
    for coordinate_range in worksheet.print_area or []:
        plain = str(coordinate_range).replace('$', '')
        try:
            min_col, min_row, max_col, max_row = range_boundaries(plain)
        except Exception:
            return True
        for row in worksheet.iter_rows(min_row=min_row, max_row=max_row,
                                       min_col=min_col, max_col=max_col):
            for cell in row:
                if cell.value not in (None, ''):
                    return True
    return False


def _sanitize_print_areas(workbook):
    """清除「不覆盖任何有值单元格」的打印区域（就地修改）。"""
    for worksheet in workbook.worksheets:
        if not _print_area_covers_content(worksheet):
            worksheet.print_area = []


def prepare_pdf_workbook(data, pdf_params=None, sheet_name=None):
    """净化 PDF 转换用的填写结果副本，返回 xlsx 字节。

    只清除「不覆盖任何有值单元格」的打印区域：这类区域是 WPS/Excel 里
    误留的（常见为指向单个空单元格），LibreOffice 严格按打印区域输出，
    会产生整份空白 PDF。覆盖内容的打印区域原样保留（尊重模板打印版式）。
    可选应用模板级 PDF 打印参数（excel_pdf_params.apply_pdf_params）：
    页面设置作用于全部工作表，重复表头行与手动分页行只作用于主工作表。
    仅作用于 PDF 转换副本：导出 Excel 与服务端源文件保持原始字节不动。
    """
    from apps.equipment_lifecycle.excel_pdf_params import apply_pdf_params
    try:
        workbook = load_workbook(io.BytesIO(data))
    except Exception:
        raise ExcelLayoutError('PDF 转换预处理失败：填写结果无法解析')
    _sanitize_print_areas(workbook)
    apply_pdf_params(workbook, pdf_params, sheet_name)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


# ==================== 草稿兼容性评估 ====================

def assess_compatibility(data):
    """对草稿源文件字节做单次解析的兼容性评估（文件维度）。

    返回 {'status', 'warnings', 'unsupported'}：
    - status: ok 可继续发布 | warning 可发布但可能有格式差异 | error
      阻止发布（对应发布/预览本身会拒绝的条件，不新增阻断口径）；
    - unsupported：不支持特性中文标签列表（历史草稿若含公式/图表等，
      上传拦截之前遗留，发布与 PDF 转换都不可用）；
    - warnings：打印区域不覆盖内容等可自动修正的差异提示；
    - 解析失败（文件损坏）视为 error。
    部署环境维度（LibreOffice 可用性）由调用方补充。
    """
    result = {'status': 'ok', 'warnings': [], 'unsupported': []}
    try:
        workbook = load_workbook(io.BytesIO(data))
    except Exception:
        result['status'] = 'error'
        result['warnings'].append('草稿源文件无法解析，请重新上传后再发布或预览')
        return result
    features = set()
    for worksheet in workbook.worksheets:
        features |= _sheet_unsupported_features(worksheet)
        _check_print_area_coverage(result, worksheet)
    if getattr(workbook, '_external_links', None):
        features.add('external_link')
    if features:
        result['unsupported'] = [
            UNSUPPORTED_FEATURE_LABELS[key]
            for key in sorted(features, key=UNSUPPORTED_FEATURE_LABELS.get)]
    result['status'] = _compatibility_status(features, result['warnings'])
    return result


def _check_print_area_coverage(result, worksheet):
    """打印区域不覆盖内容时追加 warning（PDF 转换会自动忽略该区域）。"""
    if worksheet.print_area and not _print_area_covers_content(worksheet):
        result['warnings'].append(
            '工作表「%s」的打印区域未覆盖任何内容，生成 PDF 时将忽略该区域'
            % worksheet.title)


def _compatibility_status(features, warnings):
    if features:
        return 'error'
    return 'warning' if warnings else 'ok'


def _prepare_writes(meta, values, check_required=True):
    """全量校验（必填 + 类型）后再写文件；空值不写入，保留原单元格。

    check_required=False 时不拦截必填（打印预览允许留空，与模板预览
    同一设计口径），字段类型校验保留。
    """
    writes = []
    missing = []
    for binding in meta.get('field_bindings') or []:
        raw = (values or {}).get(binding['field_id'])
        if raw is None or raw == '':
            if binding['is_required'] and check_required:
                missing.append(binding['field_name'])
            continue
        writes.append((binding['sheet_name'], binding['cell_coordinate'],
                       _typed_value(binding, raw)))
    if missing:
        raise ExcelLayoutError('必填字段未填写：%s' % '、'.join(missing))
    return writes


def _read_source(source_file_path):
    if not source_file_path or not os.path.isfile(source_file_path):
        raise ExcelLayoutError('模板源文件不存在')
    with open(source_file_path, 'rb') as f:
        return f.read()


def _verify_layout_matches(workbook, meta):
    """导出前重新读取源文件校验：工作表与合并关系必须与服务端保存一致。"""
    for sheet_meta in meta['layout']['sheets']:
        if sheet_meta['name'] not in workbook.sheetnames:
            raise ExcelLayoutError('源模板校验失败：工作表「%s」缺失'
                                   % sheet_meta['name'])
        current = sorted(str(rng) for rng in
                         workbook[sheet_meta['name']].merged_cells.ranges)
        if current != sorted(sheet_meta['merged_ranges']):
            raise ExcelLayoutError('源模板校验失败：工作表「%s」合并关系发生变化'
                                   % sheet_meta['name'])
