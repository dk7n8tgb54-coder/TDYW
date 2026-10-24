# -*- coding: utf-8 -*-
"""运行处置记录——故障等级（fault_level）契约测试（2026-10-23 新增字段）。

覆盖真实行为契约（真实 HTTP + 隔离测试数据库 + 数据库状态断言）：
1. 设备故障创建可保存四种固定等级（level1一级/level2二级/level3三级/
   level4四级），详情返回 fault_level 与 fault_level_text；
2. 取值校验：等级不在四个选项内拒绝（空串=未选择合法）；旧参数 `level`
   仍按退役参数拒绝（fault_level 不受退役守卫影响）；
3. 类型边界：非故障记录（一般记录/异常事件）携带 fault_level 提交不改变
   原有行为——服务层强制清空为空串，恒不产生带等级的非故障记录；
4. 编辑：故障可修改/清空等级（显式空串=清空），未提交保留原值；
   历史记录（字段默认空串）读取、编辑兼容；
5. 异常事件认定设备故障（identify-fault）可选填等级，非法值拒绝；
6. 导出 xlsx 含"故障等级"列：故障行展示等级文本，非故障行为空；
7. 列表/导出按等级筛选：后端查询实现（与 _filter_records 同一口径），
   只作用于设备故障记录，与类型/系统筛选、分页、租户隔离、权限正交组合。
"""
import io

from django.test import TestCase

from openpyxl import load_workbook

from apps.equipment_lifecycle.operation_models import SystemOperationRecord
from apps.equipment_lifecycle.tests.base import (
    ALL_OPERATION_PERMS, P_OP_ADD, P_OP_VIEW, RECORDS_URL, make_client,
    make_record, make_system, make_user,
)


def _make_fault_user():
    user = make_user('fault_level_user', perms=ALL_OPERATION_PERMS)
    return user, make_client(user), make_system(user)


class FaultLevelCreateTest(TestCase):
    """设备故障创建：四级等级保存 / 非法值拒绝 / 非故障类型行为不变。"""

    def setUp(self):
        self.user, self.client, self.system = _make_fault_user()

    def _payload(self, **overrides):
        payload = {
            'record_type': 'fault',
            'system': self.system.id,
            'occurred_at': '2026-10-23 08:30',
            'title': '发射机故障',
            'description': '面板告警灯亮，输出功率下降。',
            'registrar_name': '李四',
        }
        payload.update(overrides)
        return payload

    def test_create_fault_saves_each_of_four_levels(self):
        for idx, (level, text) in enumerate((
                ('level1', '一级'), ('level2', '二级'),
                ('level3', '三级'), ('level4', '四级'))):
            # 错开发生时间与描述：避开 30 秒窗口期业务查重
            res = self.client.post(
                RECORDS_URL + '/',
                self._payload(
                    fault_level=level,
                    occurred_at='2026-10-23 0%d:30' % (idx + 1),
                    description='面板告警灯亮，输出功率下降。（%s）' % text),
                content_type='application/json')
            self.assertEqual(res.status_code, 200)
            self.assertFalse(res.json().get('error'), res.json())
            data = res.json()['data']
            record = SystemOperationRecord.objects.get(pk=data['id'])
            self.assertEqual(record.record_type, 'fault')
            self.assertEqual(record.fault_level, level)
            self.assertEqual(data['fault_level'], level)
            self.assertEqual(data['fault_level_text'], text)

    def test_create_fault_without_level_defaults_empty(self):
        res = self.client.post(RECORDS_URL + '/', self._payload(),
                               content_type='application/json')
        self.assertFalse(res.json().get('error'), res.json())
        record = SystemOperationRecord.objects.get(pk=res.json()['data']['id'])
        self.assertEqual(record.fault_level, '')
        self.assertEqual(res.json()['data']['fault_level_text'], '')

    def test_create_fault_invalid_level_rejected(self):
        for bad in ('level5', '一级', 'high', '重大'):
            res = self.client.post(
                RECORDS_URL + '/', self._payload(fault_level=bad),
                content_type='application/json')
            self.assertEqual(res.status_code, 200)
            self.assertIn('故障等级取值不合法', res.json().get('error', ''))
        self.assertFalse(SystemOperationRecord.objects.exists())

    def test_non_fault_record_level_forced_empty(self):
        """一般记录/异常事件携带等级提交：行为与存量一致（服务层清空）。"""
        for record_type, title in (
                ('routine', 'UPS电池巡检'), ('abnormal', 'UPS切换异常')):
            res = self.client.post(
                RECORDS_URL + '/',
                self._payload(
                    record_type=record_type, title=title, fault_level='level2'),
                content_type='application/json')
            data = res.json()
            if not data.get('error'):
                record = SystemOperationRecord.objects.get(pk=data['data']['id'])
                self.assertEqual(record.record_type, record_type)
                self.assertEqual(record.fault_level, '')
                self.assertEqual(data['data']['fault_level_text'], '')
        # 任何路径都不产生带等级的非故障记录
        self.assertFalse(
            SystemOperationRecord.objects.exclude(fault_level='').exclude(
                record_type='fault').exists())

    def test_retired_level_param_still_rejected(self):
        """旧故障模块退役参数 `level` 出现即拒绝（fault_level 命名避开之）。"""
        res = self.client.post(
            RECORDS_URL + '/', self._payload(level='重大'),
            content_type='application/json')
        self.assertIn('记录已不再支持参数', res.json().get('error', ''))
        self.assertIn('level', res.json().get('error', ''))
        self.assertFalse(SystemOperationRecord.objects.exists())


class FaultLevelEditTest(TestCase):
    """故障等级编辑：修改/清空/保留原值；历史记录兼容。"""

    def setUp(self):
        self.user, self.client, self.system = _make_fault_user()

    def _put(self, record, **overrides):
        payload = {
            'system': self.system.id,
            'occurred_at': '2026-10-23 08:30',
            'title': '发射机故障',
            'description': '面板告警灯亮，输出功率下降。',
        }
        payload.update(overrides)
        return self.client.put(
            RECORDS_URL + '/%d/' % record.id, payload,
            content_type='application/json')

    def test_edit_fault_level_updates(self):
        record = make_record(self.user, self.system, record_type='fault',
                             fault_level='level2')
        res = self._put(record, fault_level='level4')
        self.assertFalse(res.json().get('error'), res.json())
        record.refresh_from_db()
        self.assertEqual(record.fault_level, 'level4')
        self.assertEqual(res.json()['data']['fault_level_text'], '四级')

    def test_edit_without_level_keeps_original(self):
        record = make_record(self.user, self.system, record_type='fault',
                             fault_level='level3')
        res = self._put(record)
        self.assertFalse(res.json().get('error'), res.json())
        record.refresh_from_db()
        self.assertEqual(record.fault_level, 'level3')

    def test_edit_clear_level_with_explicit_empty_string(self):
        record = make_record(self.user, self.system, record_type='fault',
                             fault_level='level1')
        res = self._put(record, fault_level='')
        self.assertFalse(res.json().get('error'), res.json())
        record.refresh_from_db()
        self.assertEqual(record.fault_level, '')

    def test_edit_invalid_level_rejected(self):
        record = make_record(self.user, self.system, record_type='fault',
                             fault_level='level1')
        res = self._put(record, fault_level='urgent')
        self.assertIn('故障等级取值不合法', res.json().get('error', ''))
        record.refresh_from_db()
        self.assertEqual(record.fault_level, 'level1')

    def test_edit_routine_ignores_fault_level(self):
        """一般记录编辑不处理等级字段：历史行为不变，字段保持空串。"""
        record = make_record(self.user, self.system, record_type='routine')
        res = self._put(record, fault_level='level1')
        self.assertFalse(res.json().get('error'), res.json())
        record.refresh_from_db()
        self.assertEqual(record.fault_level, '')

    def test_history_record_without_level_read_and_edit(self):
        """历史记录（无等级）：详情正常返回空等级文本，编辑可补填。"""
        record = make_record(self.user, self.system, record_type='fault')
        self.assertEqual(record.fault_level, '')
        detail = self.client.get(RECORDS_URL + '/%d/' % record.id)
        self.assertFalse(detail.json().get('error'), detail.json())
        self.assertEqual(detail.json()['data']['fault_level'], '')
        self.assertEqual(detail.json()['data']['fault_level_text'], '')
        res = self._put(record, fault_level='level2')
        self.assertFalse(res.json().get('error'), res.json())
        record.refresh_from_db()
        self.assertEqual(record.fault_level, 'level2')


class FaultLevelIdentifyTest(TestCase):
    """异常事件认定设备故障：可选填等级，非法值拒绝。"""

    def setUp(self):
        self.user, self.client, self.system = _make_fault_user()
        self.record = make_record(self.user, self.system, record_type='abnormal')

    def test_identify_with_level(self):
        res = self.client.post(
            RECORDS_URL + '/%d/identify-fault/' % self.record.id,
            {'fault_level': 'level2', 'affects_system_operation': False},
            content_type='application/json')
        self.assertFalse(res.json().get('error'), res.json())
        self.record.refresh_from_db()
        self.assertEqual(self.record.record_type, 'fault')
        self.assertEqual(self.record.status, 'handling')
        self.assertEqual(self.record.fault_level, 'level2')

    def test_identify_without_level_defaults_empty(self):
        res = self.client.post(
            RECORDS_URL + '/%d/identify-fault/' % self.record.id, {},
            content_type='application/json')
        self.assertFalse(res.json().get('error'), res.json())
        self.record.refresh_from_db()
        self.assertEqual(self.record.fault_level, '')

    def test_identify_invalid_level_rejected(self):
        res = self.client.post(
            RECORDS_URL + '/%d/identify-fault/' % self.record.id,
            {'fault_level': 'level9'},
            content_type='application/json')
        self.assertIn('故障等级取值不合法', res.json().get('error', ''))
        self.record.refresh_from_db()
        self.assertEqual(self.record.record_type, 'abnormal')


class FaultLevelExportTest(TestCase):
    """导出 xlsx：故障等级列在位，故障行展示等级文本、非故障行为空。"""

    def setUp(self):
        self.user, self.client, self.system = _make_fault_user()
        self.fault = make_record(self.user, self.system, record_type='fault',
                                 title='主控板异常', fault_level='level1')
        self.routine = make_record(self.user, self.system, record_type='routine',
                                   title='例行检查')

    def _export_sheet(self):
        res = self.client.get(RECORDS_URL + '/export/')
        self.assertEqual(res.status_code, 200)
        return load_workbook(io.BytesIO(res.content)).active

    @staticmethod
    def _find_row(ws, record_no):
        rows = [row for row in range(2, ws.max_row + 1)
                if ws.cell(row=row, column=1).value == record_no]
        assert rows, '导出中未找到记录 %r' % record_no
        return rows[0]

    def test_export_contains_fault_level_column(self):
        ws = self._export_sheet()
        self.assertEqual(ws.cell(row=1, column=8).value, '故障等级')
        self.assertEqual(ws.cell(row=1, column=9).value, '发生时间')
        self.assertEqual(ws.cell(row=1, column=17).value, '处理时间')
        # 基础列扩到 16 列，进展列顺延为 17-19
        self.assertEqual(ws.cell(row=1, column=19).value, '处理内容')

    def test_export_fault_row_shows_level_text(self):
        ws = self._export_sheet()
        row = self._find_row(ws, self.fault.record_no)
        self.assertEqual(ws.cell(row=row, column=8).value, '一级')

    def test_export_non_fault_row_level_empty(self):
        ws = self._export_sheet()
        row = self._find_row(ws, self.routine.record_no)
        self.assertIn(ws.cell(row=row, column=8).value, ('', None))


class FaultLevelListFilterTest(TestCase):
    """列表故障等级筛选：后端查询实现（非前端过滤），只作用于设备故障记录。

    覆盖：四等级逐一筛选、未选等级的历史记录不入任何等级结果、等级筛选
    永不返回非故障记录（含矛盾组合）、与其他筛选正交组合、分页完整性、
    非法值拒绝、租户隔离、查看权限、导出同一筛选口径。
    """

    def setUp(self):
        self.user, self.client, self.system = _make_fault_user()
        self.fault_l1 = make_record(self.user, self.system, record_type='fault',
                                    title='一级故障', fault_level='level1',
                                    occurred_at='2026-10-20 08:00:00')
        self.fault_l2 = make_record(self.user, self.system, record_type='fault',
                                    title='二级故障', fault_level='level2',
                                    occurred_at='2026-10-21 08:00:00')
        # 历史记录语义：字段上线前的故障记录等级为空串
        self.fault_no_level = make_record(
            self.user, self.system, record_type='fault', title='历史故障',
            occurred_at='2026-10-22 08:00:00')
        self.routine = make_record(self.user, self.system, record_type='routine',
                                   title='例行记录', occurred_at='2026-10-23 08:00:00')
        self.abnormal = make_record(self.user, self.system, record_type='abnormal',
                                    title='异常事件', occurred_at='2026-10-23 09:00:00')

    def _list(self, **params):
        res = self.client.get(RECORDS_URL + '/', params)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertFalse(data.get('error'), data)
        return data['data']

    def test_filter_by_each_of_four_levels(self):
        fault_l3 = make_record(self.user, self.system, record_type='fault',
                               title='三级故障', fault_level='level3',
                               occurred_at='2026-10-19 08:00:00')
        fault_l4 = make_record(self.user, self.system, record_type='fault',
                               title='四级故障', fault_level='level4',
                               occurred_at='2026-10-18 08:00:00')
        for level, record in (('level1', self.fault_l1), ('level2', self.fault_l2),
                              ('level3', fault_l3), ('level4', fault_l4)):
            data = self._list(fault_level=level)
            self.assertEqual(data['total'], 1, level)
            self.assertEqual(data['data'][0]['id'], record.id)
        # 无等级的历史故障记录不出现在任何等级筛选结果中
        for level in ('level1', 'level2', 'level3', 'level4'):
            ids = [row['id'] for row in self._list(fault_level=level)['data']]
            self.assertNotIn(self.fault_no_level.id, ids)

    def test_filter_level_never_returns_non_fault_records(self):
        """等级只存在于设备故障：非故障记录与矛盾组合都不出现在结果中。"""
        data = self._list(fault_level='level1')
        self.assertEqual(
            {row['id'] for row in data['data']}, {self.fault_l1.id})
        # 显式矛盾组合（类型=一般记录 + 等级）返回空集，不报错不泄露
        data = self._list(record_type='routine', fault_level='level1')
        self.assertEqual(data['total'], 0)

    def test_filter_level_combined_with_record_type(self):
        data = self._list(record_type='fault', fault_level='level2')
        self.assertEqual(data['total'], 1)
        self.assertEqual(data['data'][0]['id'], self.fault_l2.id)

    def test_filter_level_combined_with_system(self):
        other_system = make_system(self.user, system_code='SYS-B', name='系统B')
        make_record(self.user, other_system, record_type='fault',
                    title='他系统一级故障', fault_level='level1',
                    occurred_at='2026-10-19 08:00:00')
        data = self._list(system=self.system.id, fault_level='level1')
        self.assertEqual(data['total'], 1)
        self.assertEqual(data['data'][0]['id'], self.fault_l1.id)

    def test_filter_level_with_pagination(self):
        """分页由后端切片完成：翻页取全等级命中记录且不重复。"""
        extra_a = make_record(self.user, self.system, record_type='fault',
                              title='补充一级A', fault_level='level1',
                              occurred_at='2026-10-19 08:00:00')
        extra_b = make_record(self.user, self.system, record_type='fault',
                              title='补充一级B', fault_level='level1',
                              occurred_at='2026-10-18 08:00:00')
        page1 = self._list(fault_level='level1', page=1, page_size=2)
        page2 = self._list(fault_level='level1', page=2, page_size=2)
        self.assertEqual(page1['total'], 3)
        self.assertEqual(len(page1['data']), 2)
        self.assertEqual(len(page2['data']), 1)
        self.assertEqual(
            {row['id'] for row in page1['data'] + page2['data']},
            {self.fault_l1.id, extra_a.id, extra_b.id})

    def test_filter_invalid_level_rejected(self):
        for bad in ('level5', '一级', 'high'):
            res = self.client.get(RECORDS_URL + '/', {'fault_level': bad})
            self.assertIn('故障等级筛选值不合法', res.json().get('error', ''), bad)

    def test_filter_level_tenant_isolation(self):
        other_user = make_user('fault_level_t2', perms=ALL_OPERATION_PERMS,
                               tenant_id='tenant_2')
        other_system = make_system(other_user, tenant_id='tenant_2')
        make_record(other_user, other_system, record_type='fault',
                    tenant_id='tenant_2',
                    title='他租户一级故障', fault_level='level1')
        data = self._list(fault_level='level1')
        self.assertEqual(data['total'], 1)
        self.assertEqual(data['data'][0]['id'], self.fault_l1.id)

    def test_filter_level_requires_view_permission(self):
        outsider = make_user('fault_level_no_perm', perms=[])
        client = make_client(outsider)
        res = client.get(RECORDS_URL + '/', {'fault_level': 'level1'})
        self.assertIn('权限拒绝', res.json().get('error', ''))

    def test_export_honors_level_filter(self):
        """导出与列表同一 _filter_records 口径：等级命中记录才进工作簿。"""
        res = self.client.get(RECORDS_URL + '/export/', {'fault_level': 'level2'})
        self.assertEqual(res.status_code, 200)
        ws = load_workbook(io.BytesIO(res.content)).active
        record_nos = {ws.cell(row=row, column=1).value
                      for row in range(2, ws.max_row + 1)}
        self.assertEqual(record_nos, {self.fault_l2.record_no})
