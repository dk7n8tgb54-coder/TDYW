# -*- coding: utf-8 -*-
"""模板打印设置「应用当前设置并预览」契约（2026-10-19，真实 PDF 转换）。

背景：打印参数此前只能先保存（写入草稿）再经空白预览验证，调整参数
必须落库；发布前的最终验收标准应是"平台生成的 PDF"。本套件固化：

1. 未保存参数预览：POST preview/pdf 携带 {"pdf_params": {...}} 时，
   参数整体校验后仅作用于本次预览副本——纸张方向按未保存值生效
   （纸张响应头反映真实页面宽高），草稿 pdf_params 不被改写；
2. 临时样例内容预览：{"sample_values": {field_id: 值}} 经字段绑定与
   类型校验（未知 field_id 拒绝、受控类型必须是候选项）后写入预览
   副本，真实 PDF 中可检索样例文本；必填不强制（预览允许留空）；
3. 只读副作用：预览不写草稿布局/绑定/源文件字节、不创建版本、
   不创建台账记录（数据库计数前后一致）；
4. 预览反馈响应头：X-El-Pdf-Pages / X-El-Pdf-Paper 来自真实生成的
   PDF（pypdf 读取），X-El-Pdf-Warnings 为按本次实际参数的版式预检
   提示（percent 编码 JSON），不伪造；
5. 范围列宽回归：A:C 由单条 <col min="1" max="3"> 统一设定 + 合并
   标题 + 高行 + 长文本的工作簿，预览正常出 PDF（列宽解析修复的
   端到端底稿，解析细节见 test_excel_column_range_layout）。

全部走真实 HTTP 路径与真实 LibreOffice 转换（环境缺失显式 NOT_RUN），
不做源码字符串断言。
"""
import io
import json
import os
import shutil
import subprocess
import tempfile
import uuid
from unittest import skipUnless

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from openpyxl import Workbook
from openpyxl.worksheet.dimensions import ColumnDimension

from apps.equipment_lifecycle.excel_models import (
    OperationLedgerExcelRecord,
    OperationLedgerExcelTemplate,
    OperationLedgerExcelTemplateVersion,
)
from apps.equipment_lifecycle.excel_pdf_service import detect_soffice
from apps.equipment_lifecycle.tests.base import URL_PREFIX
from apps.equipment_lifecycle.tests.stable_contract.test_excel_ledger_templates import (
    ExcelLedgerContractBase,
)

TPL_URL = URL_PREFIX + '/excel-templates'
SHEET_NAME = '信道维护记录'
RANGE_WIDTH = 36.21
# 足够长的样例文本（验证换行与真实渲染；pdftotext 检索短子串）
SAMPLE_TEXT = ('设备主电台发射功率异常，已切换备用信道并完成链路测试，'
               '天馈驻波比恢复正常，待观察一昼夜后复测。')
SAMPLE_MARKER = '发射功率异常'


def build_narrow_xlsx():
    """窄内容工作簿（内容宽约 7cm）：方向自动解析为纵向；第 2 行留可绑定格。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws['A1'] = '日检表'
    ws['B1'] = '备注'
    ws['A2'] = '记录'
    ws.column_dimensions['A'] = ColumnDimension(
        ws, index='A', width=8.43, customWidth=True, min=1, max=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_range_tall_xlsx():
    """A:C 范围列宽 + 合并标题 + 高行 + 多行内容（长文本预览回归底稿）。

    约 100 行内容保证 fit_page 单页适配缩放远低于可读阈值（告警确定）。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws['A1'] = '信道维护记录表'
    ws.merge_cells('A1:C1')
    ws['A2'] = '维护记录'
    ws.column_dimensions['A'] = ColumnDimension(
        ws, index='A', width=RANGE_WIDTH, customWidth=True, min=1, max=3)
    ws.row_dimensions[3].height = 60
    for row in range(4, 104):
        ws['A%d' % row] = '第%d行例行数据' % row
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdftotext(pdf_bytes):
    """真实 PDF 文本提取（pdftotext；调用方先以 shutil.which 判定可用）。"""
    with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as f:
        f.write(pdf_bytes)
        pdf_path = f.name
    try:
        result = subprocess.run(['pdftotext', pdf_path, '-'],
                                capture_output=True)
        return result.stdout.decode('utf-8', 'replace')
    finally:
        os.remove(pdf_path)


def _pdftotext_available():
    return bool(detect_soffice() and shutil.which('pdftotext'))


class PreviewFeedbackUnitTests(SimpleTestCase):
    """PDF 反馈信息读取（纯函数，无数据库）：解析失败不伪造结果。"""

    def test_pdf_page_info_returns_none_on_unparsable_content(self):
        from apps.equipment_lifecycle.excel_pdf_service import pdf_page_info
        self.assertIsNone(pdf_page_info(b'broken')['pages'])
        self.assertIsNone(pdf_page_info(b'')['paper'])


class PreviewFeedbackContractTests(ExcelLedgerContractBase):
    """未保存参数 + 临时样例值预览（真实 HTTP + 真实 LibreOffice PDF）。"""

    def setUp(self):
        super().setUp()
        self.template = self._upload(build_narrow_xlsx(), '窄表模板')

    def _upload(self, data, name):
        response = self.owner_client.post(TPL_URL + '/', {
            'file': SimpleUploadedFile('preview-feedback.xlsx', data),
            'name': name,
            'ledger_type': 'week_month',
        })
        payload = response.json()
        self.assertEqual(payload.get('error', ''), '')
        return payload['data']['data']

    def _bind(self, template_id, coordinate, field_name='维护记录',
              input_type='text', required=True):
        binding = {
            'field_id': str(uuid.uuid4()),
            'sheet_name': SHEET_NAME,
            'cell_coordinate': coordinate,
            'field_name': field_name,
            'input_type': input_type,
            'is_required': required,
            'unit': '',
            'options': [],
        }
        response = self.owner_client.put(
            '%s/%s/bindings/' % (TPL_URL, template_id),
            json.dumps({'field_bindings': [binding]}),
            content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        return binding

    def _preview(self, template_id, body=None, client=None):
        client = client or self.owner_client
        kwargs = {}
        if body is not None:
            kwargs['data'] = json.dumps(body)
            kwargs['content_type'] = 'application/json'
        return client.post('%s/%s/preview/pdf/' % (TPL_URL, template_id),
                           **kwargs)

    @staticmethod
    def _draft_pdf_params(template_id):
        part = json.loads(
            OperationLedgerExcelTemplate.objects.get(pk=template_id).draft_layout)
        return part.get('pdf_params') or {}

    def _draft_bytes(self, template_id):
        from apps.equipment_lifecycle import excel_template_storage
        with open(excel_template_storage.draft_path(template_id), 'rb') as f:
            return f.read()

    @staticmethod
    def _decode_warnings(response):
        import urllib.parse
        return json.loads(urllib.parse.unquote(response['X-El-Pdf-Warnings']))

    @skipUnless(detect_soffice(),
                'LibreOffice 不可用，预览反馈用例 NOT_RUN（环境限制）')
    def test_unsaved_params_apply_without_saving(self):
        """未保存横向参数即时生效（纸张头 297x210），草稿仍为播种参数。"""
        before = self._draft_pdf_params(self.template['id'])
        response = self._preview(self.template['id'], body={
            'pdf_params': {'preset': 'custom', 'paper_size': 'A4',
                           'orientation': 'landscape'},
        })
        self.assertIn('application/pdf', response['Content-Type'])
        self.assertTrue(response.content.startswith(b'%PDF-'))
        self.assertIn('297x210', response['X-El-Pdf-Paper'])
        self.assertEqual(int(response['X-El-Pdf-Pages']), 1)
        # 草稿参数未被改写（仍为上传播种的 a4_auto/自动方向）
        self.assertEqual(self._draft_pdf_params(self.template['id']), before)

    @skipUnless(detect_soffice(),
                'LibreOffice 不可用，预览反馈用例 NOT_RUN（环境限制）')
    def test_blank_preview_uses_saved_params_portrait(self):
        """无 body 的空白预览沿用已保存参数：窄表自动方向解析为纵向。"""
        response = self._preview(self.template['id'])
        self.assertIn('application/pdf', response['Content-Type'])
        self.assertIn('210x297', response['X-El-Pdf-Paper'])
        self.assertGreaterEqual(int(response['X-El-Pdf-Pages']), 1)

    @skipUnless(_pdftotext_available(),
                'LibreOffice 或 pdftotext 不可用，样例文本渲染用例 NOT_RUN')
    def test_sample_values_render_into_pdf_without_side_effects(self):
        """样例长文本写入预览副本并真实渲染；数据库零写入副作用。"""
        template = self._upload(build_range_tall_xlsx(), '范围列宽高行模板')
        binding = self._bind(template['id'], 'A3')
        draft_before = self._draft_bytes(template['id'])
        versions_before = OperationLedgerExcelTemplateVersion.objects.count()
        records_before = OperationLedgerExcelRecord.objects.count()
        response = self._preview(template['id'], body={
            'sample_values': {binding['field_id']: SAMPLE_TEXT},
        })
        self.assertIn('application/pdf', response['Content-Type'])
        self.assertGreaterEqual(int(response['X-El-Pdf-Pages']), 1)
        text = _pdftotext(response.content)
        self.assertIn(SAMPLE_MARKER, text.replace('\n', '').replace(' ', ''))
        # 必填字段留空的空白预览同样可生成（预览不强制必填）
        blank = self._preview(template['id'])
        self.assertIn('application/pdf', blank['Content-Type'])
        # 只读：草稿字节、版本数、台账记录数全部不变
        self.assertEqual(self._draft_bytes(template['id']), draft_before)
        self.assertEqual(OperationLedgerExcelTemplateVersion.objects.count(),
                         versions_before)
        self.assertEqual(OperationLedgerExcelRecord.objects.count(),
                         records_before)

    def test_sample_values_reject_unknown_field_and_bad_type(self):
        """未知 field_id 与非法受控值拒绝（复用绑定与类型校验，无需 LO）。"""
        response = self._preview(self.template['id'], body={
            'sample_values': {str(uuid.uuid4()): '任意'},
        })
        self.assertIn('未绑定', response.json()['error'])
        binding = self._bind(self.template['id'], 'B2', '结论',
                             input_type='status', required=False)
        response = self._preview(self.template['id'], body={
            'sample_values': {binding['field_id']: '良好'},
        })
        self.assertIn('候选项', response.json()['error'])
        response = self._preview(self.template['id'], body={
            'sample_values': ['不是对象'],
        })
        self.assertIn('预览数据必须是对象', response.json()['error'])

    def test_unsaved_params_validation_errors(self):
        """参数整体校验：非法纸张与打印区域越界逐项拒绝（无需 LO）。"""
        response = self._preview(self.template['id'], body={
            'pdf_params': {'paper_size': 'X9'},
        })
        self.assertIn('纸张大小', response.json()['error'])
        response = self._preview(self.template['id'], body={
            'pdf_params': {'print_area': 'A1:Z99'},
        })
        self.assertIn('超出工作表', response.json()['error'])

    @skipUnless(detect_soffice(),
                'LibreOffice 不可用，版式预检反馈用例 NOT_RUN（环境限制）')
    def test_warnings_header_reports_scale_hint_by_actual_params(self):
        """单页适配 + A5：按本次实际参数给出缩放过小提示（percent 编码）。"""
        template = self._upload(build_range_tall_xlsx(), '高行多行模板')
        response = self._preview(template['id'], body={
            'pdf_params': {'preset': 'single_page', 'paper_size': 'A5',
                           'scale_mode': 'fit_page'},
        })
        self.assertIn('application/pdf', response['Content-Type'])
        self.assertTrue(response['X-El-Pdf-Paper'].startswith('A5'))
        warnings = self._decode_warnings(response)
        self.assertIsInstance(warnings, list)
        self.assertTrue(any('缩放' in w or '字号' in w for w in warnings),
                        warnings)
        # 常规参数时提示头始终可解析（空列表也为合法 JSON）
        blank = self._preview(template['id'])
        self.assertEqual(self._decode_warnings(blank), [])
