# -*- coding: utf-8 -*-
"""运行处置记录——2026-10-14 体验二轮行为契约（真实 HTTP + 隔离测试数据库）。

覆盖：
1. 发生时间范围边界：结束时间按用户所选整分钟包含该分钟记录（列表口径，
   导出/统计/系统页签共用同一 _apply_occurred_range）；显式携带秒的取值
   保持排他边界；仅日期按当天结束处理。
2. 待办：跟进中异常事件 + 处置中设备故障（真正需要持续处理的两类记录），
   分组计数 counts 不随 record_type 筛选变化；分页全量可达，无静默截断。
3. 导出筛选：导出范围与列表筛选完全一致（关键词/系统/时间范围生效，
   默认仍排除已撤销，非法筛选值拒绝）。
4. 统计时间边界与分组计数口径。
"""
import io

from django.test import TestCase

from openpyxl import load_workbook

from apps.equipment_lifecycle.tests.base import (
    ALL_OPERATION_PERMS, RECORDS_URL, make_client, make_record, make_system,
    make_user,
)


class OperationRecordTimeBoundaryTest(TestCase):
    """发生时间范围边界：结束时间按用户所选整分钟包含该分钟记录。

    用户按分钟粒度选择截止时间（前端 RangePicker 固定 HH:mm）时，
    该分钟内发生的记录必须全部命中；显式携带秒的取值保持排他边界。
    """

    def setUp(self):
        self.user = make_user('time_boundary_user', perms=ALL_OPERATION_PERMS)
        self.client = make_client(self.user)
        self.system = make_system(self.user)
        make_record(self.user, self.system, occurred_at='2026-10-11 10:30:30')
        make_record(self.user, self.system, occurred_at='2026-10-11 10:31:00')

    def _list(self, **params):
        res = self.client.get(RECORDS_URL + '/', params)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertFalse(data.get('error'), data)
        return data['data']

    def test_end_time_includes_selected_minute(self):
        data = self._list(occurred_from='2026-10-11 00:00', occurred_to='2026-10-11 10:30')
        self.assertEqual(data['total'], 1)

    def test_next_minute_boundary(self):
        # 截止 10:31：10:31:00 属于所选分钟内（命中），10:32 起排除
        data = self._list(occurred_from='2026-10-11 00:00', occurred_to='2026-10-11 10:31')
        self.assertEqual(data['total'], 2)

    def test_explicit_seconds_boundary_stays_exclusive(self):
        data = self._list(occurred_from='2026-10-11 00:00:00',
                          occurred_to='2026-10-11 10:30:30')
        self.assertEqual(data['total'], 0)

    def test_end_time_date_only_includes_whole_day(self):
        make_record(self.user, self.system, occurred_at='2026-10-11 23:59:59')
        data = self._list(occurred_from='2026-10-11 00:00', occurred_to='2026-10-11')
        self.assertEqual(data['total'], 3)


class OperationTodoTest(TestCase):
    """待办契约：跟进中异常事件 + 处置中设备故障，分组计数/类型筛选/分页。

    待办口径 = 真正需要持续处理的两类记录；一般记录（无跟进语义）、
    已完成/已关闭/已撤销不进入待办。counts 为全量分组计数，不随
    record_type 筛选变化；分页全量可达，无静默截断。
    """

    def setUp(self):
        self.user = make_user('todo_user', perms=ALL_OPERATION_PERMS)
        self.client = make_client(self.user)
        self.system = make_system(self.user)
        self.following = make_record(self.user, self.system, record_type='abnormal',
                                     title='待跟进异常', status='following_up')
        self.handling_fault = make_record(self.user, self.system, record_type='fault',
                                          title='处置中故障')
        make_record(self.user, self.system, record_type='routine', title='一般记录')
        make_record(self.user, self.system, record_type='abnormal', title='已完成',
                    status='completed')
        make_record(self.user, self.system, record_type='fault', title='已关闭',
                    status='closed')
        make_record(self.user, self.system, record_type='fault', title='已撤销',
                    status='cancelled')

    def _todo(self, **params):
        res = self.client.get(RECORDS_URL + '/todo/', params)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertFalse(data.get('error'), data)
        return data['data']

    def test_todo_contains_following_and_handling_only(self):
        payload = self._todo()
        nos = [row['record_no'] for row in payload['data']]
        self.assertIn(self.following.record_no, nos)
        self.assertIn(self.handling_fault.record_no, nos)
        self.assertEqual(payload['total'], 2)
        self.assertEqual(payload['counts']['following_up'], 1)
        self.assertEqual(payload['counts']['handling'], 1)
        self.assertEqual(payload['counts']['total'], 2)

    def test_todo_record_type_filter_keeps_counts(self):
        payload = self._todo(record_type='fault')
        self.assertEqual(payload['total'], 1)
        # counts 不随 record_type 筛选变化（前端始终展示全量分组数量）
        self.assertEqual(payload['counts']['total'], 2)
        payload = self._todo(record_type='routine')
        self.assertEqual(payload['total'], 0)

    def test_todo_invalid_record_type_rejected(self):
        res = self.client.get(RECORDS_URL + '/todo/', {'record_type': 'xxx'})
        self.assertTrue(res.json().get('error'))

    def test_todo_pagination_full_reachable(self):
        for i in range(25):
            make_record(self.user, self.system, record_type='abnormal',
                        status='following_up', title='积压事件%d' % i)
        page1 = self._todo(page=1, page_size=10)
        page2 = self._todo(page=2, page_size=10)
        self.assertEqual(page1['total'], 27)
        self.assertEqual(len(page1['data']), 10)
        self.assertEqual(len(page2['data']), 10)
        self.assertEqual(page2['page'], 2)
        # 翻页可达全部记录（无静默截断）：两页编号不重叠
        nos1 = {row['record_no'] for row in page1['data']}
        nos2 = {row['record_no'] for row in page2['data']}
        self.assertFalse(nos1 & nos2)


class OperationExportFilterTest(TestCase):
    """导出筛选契约：导出范围与列表筛选完全一致。

    关键词/系统/时间范围（含所选整分钟）与列表同一 _filter_records 口径；
    非法筛选值拒绝。默认仍排除已撤销。
    """

    def setUp(self):
        self.user = make_user('export_filter_user', perms=ALL_OPERATION_PERMS)
        self.client = make_client(self.user)
        self.system_a = make_system(self.user, system_code='SYS-A', name='系统A')
        self.system_b = make_system(self.user, system_code='SYS-B', name='系统B')
        self.matched = make_record(self.user, self.system_a,
                                   occurred_at='2026-10-11 10:30:30',
                                   title='主机异常告警')
        make_record(self.user, self.system_b, occurred_at='2026-10-12 09:00:00',
                    title='例行检查')
        make_record(self.user, self.system_a, status='cancelled',
                    record_type='fault', title='已撤销重复登记')

    def _nos(self, **params):
        res = self.client.get(RECORDS_URL + '/export/', params)
        self.assertEqual(res.status_code, 200)
        self.assertIn('spreadsheetml.sheet', res['Content-Type'])
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb.active
        return [ws.cell(row=row, column=1).value
                for row in range(2, ws.max_row + 1)]

    def test_export_keyword_filter(self):
        nos = self._nos(keyword='主机异常')
        self.assertEqual(nos, [self.matched.record_no])

    def test_export_system_filter(self):
        nos = self._nos(system=self.system_b.id)
        self.assertEqual(len(nos), 1)
        self.assertNotIn(self.matched.record_no, nos)

    def test_export_occurred_range_includes_selected_minute(self):
        nos = self._nos(occurred_from='2026-10-11 00:00',
                        occurred_to='2026-10-11 10:30')
        self.assertEqual(nos, [self.matched.record_no])

    def test_export_default_still_excludes_cancelled(self):
        nos = self._nos()
        self.assertIn(self.matched.record_no, nos)
        # 系统B的例行检查 + 主机异常告警，已撤销不导出
        self.assertEqual(len(nos), 2)

    def test_export_invalid_filter_rejected(self):
        res = self.client.get(RECORDS_URL + '/export/', {'status': 'xxx'})
        self.assertTrue(res.json().get('error'))
        res = self.client.get(RECORDS_URL + '/export/', {'record_type': 'xxx'})
        self.assertTrue(res.json().get('error'))


class OperationStatisticsBoundaryTest(TestCase):
    """统计时间边界与口径：截止时间按所选整分钟包含该分钟记录。"""

    def setUp(self):
        self.user = make_user('stats_user', perms=ALL_OPERATION_PERMS)
        self.client = make_client(self.user)
        self.system = make_system(self.user)
        make_record(self.user, self.system, occurred_at='2026-10-11 10:30:30')
        make_record(self.user, self.system, occurred_at='2026-10-11 11:00:00')
        make_record(self.user, self.system, record_type='abnormal',
                    status='following_up', occurred_at='2026-10-11 12:00:00')

    def _stats(self, **params):
        res = self.client.get(RECORDS_URL + '/statistics/', params)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertFalse(data.get('error'), data)
        return data['data']

    def test_statistics_occurred_to_includes_minute(self):
        stats = self._stats(occurred_from='2026-10-11 00:00',
                            occurred_to='2026-10-11 10:30')
        self.assertEqual(stats['totals']['total'], 1)

    def test_statistics_counts_groups(self):
        stats = self._stats(occurred_from='2026-10-11 00:00',
                            occurred_to='2026-10-12 00:00')
        self.assertEqual(stats['totals']['total'], 3)
        self.assertEqual(stats['totals']['following_up'], 1)
        self.assertEqual(stats['totals']['handling_fault'], 0)
