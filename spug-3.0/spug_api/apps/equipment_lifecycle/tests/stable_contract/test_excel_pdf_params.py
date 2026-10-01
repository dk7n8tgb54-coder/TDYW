# -*- coding: utf-8 -*-
"""Excel 路线模板级 PDF 打印参数稳定契约（2026-09-30）。

覆盖任务书「PDF 参数保存和默认使用」验收项（真实 HTTP 路径
URL→View→Service→Model，转换边界打桩的用例不依赖 LibreOffice）：

1. 上传解析播种：pdf_params 由文件打印设置自动生成（纸张/方向/缩放/
   页边距），保存在草稿布局 JSON 内（零迁移）；
2. 保存与回显：PUT pdf-params 全量校验后整体替换，重新打开详情正确
   回显，未传字段回落缺省；
3. 非法拒绝：方向/纸张/百分比/重复表头行/手动分页行/页边距越界逐一
   拒绝且草稿不变；重复表头行与分页行超出主工作表行数拒绝；
4. 发布快照：发布后版本详情携带 pdf_params；修改草稿参数再发布 v2，
   v1 快照不变（历史台账不受模板后续修改影响）；
5. 删除拦截与跨租户：停用模板拒绝保存，跨租户统一按不存在拒绝；
6. 模板四态：draft/published/used/deleted 与 ledgers_count 随列表和
   详情返回；
7. 参数应用：apply_pdf_params 写入工作簿打印设置（纸张编号/方向/页序/
   缩放与 fit 互斥/页边距英寸换算/打印标题行/手动分页行去重升序），
   分页细节只写主工作表；
8. 导出默认使用版本快照参数（转换边界捕获 xlsx 验证），可显式覆盖
   （仅当次导出生效）；历史台账导出继续使用创建时版本参数；
9. 导出预检：必填未填/横向溢出提示，填写齐全且宽度适配时无提示；
   LibreOffice 可用时真实转换成功；预检与导出权限独立校验；
10. 打印模式扩展（2026-10-01）：original/a4_auto/single_page/custom
    四档 preset 播种与快照（样板纸张方向不触发 original）、方向自动
    按内容宽度逐表解析（宽表横向/窄表纵向）、original 整体跳过参数
    应用、打印区域与水平垂直居中应用及作用域、单页缩放过小预检告警、
    历史台账固定创建时版本的模式快照。
"""
import io
import json
import uuid
from unittest import mock, skipUnless

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.equipment_lifecycle.excel_pdf_service import detect_soffice
from apps.equipment_lifecycle.tests.base import make_client, make_user
from apps.equipment_lifecycle.tests.stable_contract.test_excel_ledger_templates import (
    LEDGER_URL,
    P_LEDGER,
    P_TPL,
    SHEET_MAIN,
    TPL_URL,
    ExcelLedgerContractBase,
)

# A3 横向 + 按百分比缩放 + 重复表头 + 手动分页（样例主表行数内合法）
LANDSCAPE_PARAMS = {
    'paper_size': 'A3',
    'orientation': 'landscape',
    'scale_mode': 'percent',
    'scale_percent': 80,
    'repeat_header_rows': '1:2',
    'row_breaks': [4],
    'page_order': 'over_then_down',
}

# 样例内容宽度约 20.7cm，A4 纵向 + 原尺寸必然横向分页（预检用例基准）
OVERFLOW_PARAMS = {'paper_size': 'A4', 'orientation': 'portrait',
                   'scale_mode': 'actual'}


class PdfParamsContractTests(ExcelLedgerContractBase):
    """模板级 PDF 参数：播种/保存回显/校验/发布快照/删除与租户边界/四态。"""

    def _save(self, template_id, params, client=None):
        client = client or self.owner_client
        return client.put('%s/%s/pdf-params/' % (TPL_URL, template_id),
                          json.dumps({'pdf_params': params}),
                          content_type='application/json')

    def _list_rows_by_id(self, client=None):
        client = client or self.owner_client
        response = client.get(TPL_URL + '/')
        self.assertEqual(response.json().get('error', ''), '')
        return {row['id']: row for row in response.json()['data']['data']}

    def test_upload_seeds_pdf_params_from_file(self):
        """样例打印设置（A4 纵向 fitToWidth）播种为模板级参数缺省值。"""
        template = self.upload_sample()
        params = template['draft']['pdf_params']
        self.assertEqual(params['paper_size'], 'A4')
        self.assertEqual(params['orientation'], 'portrait')
        self.assertEqual(params['scale_mode'], 'fit_width')
        self.assertEqual(params['margins']['left'], 1.52)  # 0.6in → 1.52cm
        self.assertEqual(params['row_breaks'], [])

    def test_save_and_restore_pdf_params(self):
        template = self.upload_sample()
        response = self._save(template['id'], LANDSCAPE_PARAMS)
        self.assertEqual(response.json().get('error', ''), '')
        saved = response.json()['data']['data']['draft']['pdf_params']
        self.assertEqual(saved['paper_size'], 'A3')
        self.assertEqual(saved['orientation'], 'landscape')
        self.assertEqual(saved['scale_mode'], 'percent')
        self.assertEqual(saved['scale_percent'], 80)
        self.assertEqual(saved['repeat_header_rows'], '1:2')
        self.assertEqual(saved['row_breaks'], [4])
        self.assertEqual(saved['page_order'], 'over_then_down')
        # 重新打开（GET 详情）回显；未传字段回落缺省
        response = self.owner_client.get('%s/%s/' % (TPL_URL, template['id']))
        again = response.json()['data']['data']['draft']['pdf_params']
        self.assertEqual(again['orientation'], 'landscape')
        self.assertEqual(again['margins']['left'], 1.78)

    def test_invalid_pdf_params_rejected_and_draft_unchanged(self):
        template = self.upload_sample()
        bad_inputs = (
            {'orientation': 'diagonal'}, {'paper_size': 'A8'},
            {'scale_mode': 'shrink'}, {'scale_percent': 5},
            {'repeat_header_rows': '2:1'}, {'row_breaks': [0]},
            {'row_breaks': '4'}, {'margins': {'top': 20}},
        )
        for bad in bad_inputs:
            response = self._save(template['id'], bad)
            self.assertTrue(response.json().get('error'), '应拒绝：%s' % bad)
        response = self.owner_client.get('%s/%s/' % (TPL_URL, template['id']))
        params = response.json()['data']['data']['draft']['pdf_params']
        self.assertEqual(params['paper_size'], 'A4')
        self.assertEqual(params['orientation'], 'portrait')

    def test_pagination_rows_beyond_main_sheet_rejected(self):
        template = self.upload_sample()
        response = self._save(template['id'], {'repeat_header_rows': '1:99'})
        self.assertIn('超出', response.json()['error'])
        response = self._save(template['id'], {'row_breaks': [99]})
        self.assertIn('超出', response.json()['error'])

    def test_publish_snapshots_pdf_params_and_keeps_old_version(self):
        template = self.upload_sample()
        bindings = [self.make_binding('C4', '第1信道-主电台-PS模块')]
        response = self.save_bindings(template['id'], bindings)
        self.assertEqual(response.json().get('error', ''), '')
        response = self._save(template['id'], LANDSCAPE_PARAMS)
        self.assertEqual(response.json().get('error', ''), '')
        version = self.publish(template['id'])
        self.assertEqual(version['version_no'], 1)
        response = self.owner_client.get(
            '%s/%s/versions/1/' % (TPL_URL, template['id']))
        v1 = response.json()['data']['data']
        self.assertEqual(v1['pdf_params']['orientation'], 'landscape')
        self.assertEqual(v1['pdf_params']['paper_size'], 'A3')
        # 修改草稿参数并发布 v2：v1 快照不变
        response = self._save(template['id'], {'orientation': 'portrait'})
        self.assertEqual(response.json().get('error', ''), '')
        self.publish(template['id'])
        response = self.owner_client.get(
            '%s/%s/versions/1/' % (TPL_URL, template['id']))
        self.assertEqual(
            response.json()['data']['data']['pdf_params']['orientation'],
            'landscape')
        response = self.owner_client.get(
            '%s/%s/versions/2/' % (TPL_URL, template['id']))
        self.assertEqual(
            response.json()['data']['data']['pdf_params']['orientation'],
            'portrait')

    def _delete_template(self, template_id):
        """del 权限独立校验：owner 无 del 权限，用专门的删除用户执行。"""
        deleter = make_user('excel_del_%s' % uuid.uuid4().hex[:8],
                            perms=P_TPL + ['equipment_lifecycle.ledger_template.del'])
        response = make_client(deleter).delete('%s/%s/' % (TPL_URL, template_id))
        self.assertEqual(response.json().get('error', ''), '')
        return response

    def test_save_pdf_params_blocked_after_delete(self):
        template = self.upload_sample()
        self._delete_template(template['id'])
        response = self._save(template['id'], LANDSCAPE_PARAMS)
        self.assertIn('已删除', response.json()['error'])

    def test_pdf_params_cross_tenant_rejected(self):
        template = self.upload_sample()
        response = self.other_client.put(
            '%s/%s/pdf-params/' % (TPL_URL, template['id']),
            json.dumps({'pdf_params': LANDSCAPE_PARAMS}),
            content_type='application/json')
        self.assertEqual(response.json()['error'], 'Excel台账模板不存在')

    def test_template_status_and_ledgers_count_lifecycle(self):
        draft = self.upload_sample()
        template, version, _ = self.create_published_version()
        rows = self._list_rows_by_id()
        self.assertEqual(rows[template['id']]['status'], 'published')
        self.assertEqual(rows[template['id']]['ledgers_count'], 0)
        self.assertEqual(rows[draft['id']]['status'], 'draft')
        response = self.owner_client.post(
            LEDGER_URL + '/',
            json.dumps({'version_id': version['id'],
                        'executed_date': '2026-09-30'}),
            content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        rows = self._list_rows_by_id()
        self.assertEqual(rows[template['id']]['status'], 'used')
        self.assertEqual(rows[template['id']]['ledgers_count'], 1)
        response = self.owner_client.get('%s/%s/' % (TPL_URL, template['id']))
        detail = response.json()['data']['data']
        self.assertEqual(detail['status'], 'used')
        self.assertEqual(detail['ledgers_count'], 1)
        self._delete_template(template['id'])
        rows = self._list_rows_by_id()
        self.assertEqual(rows[template['id']]['status'], 'deleted')


class PdfApplyAndCheckTests(ExcelLedgerContractBase):
    """参数应用与导出预检（默认走版本快照参数，可覆盖，仅当次生效）。"""

    def _publish_with_params(self, params=None):
        template = self.upload_sample()
        bindings = [self.make_binding('C4', '第1信道-主电台-PS模块')]
        response = self.save_bindings(template['id'], bindings)
        self.assertEqual(response.json().get('error', ''), '')
        if params is not None:
            response = self._save_params(template['id'], params)
            self.assertEqual(response.json().get('error', ''), '')
        version = self.publish(template['id'])
        return template, version, bindings

    def _save_params(self, template_id, params):
        return self.owner_client.put(
            '%s/%s/pdf-params/' % (TPL_URL, template_id),
            json.dumps({'pdf_params': params}),
            content_type='application/json')

    def _create_filled_ledger(self, version_id, bindings):
        response = self.owner_client.post(
            LEDGER_URL + '/',
            json.dumps({'version_id': version_id,
                        'executed_date': '2026-09-30'}),
            content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        ledger_id = response.json()['data']['data']['id']
        values = {binding['field_id']: '正常' for binding in bindings}
        response = self.owner_client.put(
            '%s/%s/' % (LEDGER_URL, ledger_id),
            json.dumps({'values': values}), content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        return ledger_id

    def _capture_pdf_workbook(self, ledger_id, body=None):
        from apps.equipment_lifecycle import excel_template_services
        with mock.patch.object(excel_template_services, '_xlsx_bytes_to_pdf',
                               return_value=b'%PDF-captured') as fake:
            if body is None:
                response = self.owner_client.post(
                    '%s/%s/export/pdf/' % (LEDGER_URL, ledger_id))
            else:
                response = self.owner_client.post(
                    '%s/%s/export/pdf/' % (LEDGER_URL, ledger_id),
                    json.dumps(body), content_type='application/json')
            self.assertIn('application/pdf', response['Content-Type'])
            return load_workbook(io.BytesIO(fake.call_args[0][0]))

    def _check_warnings(self, ledger_id, body=None):
        if body is None:
            response = self.owner_client.post(
                '%s/%s/export/check/' % (LEDGER_URL, ledger_id))
        else:
            response = self.owner_client.post(
                '%s/%s/export/check/' % (LEDGER_URL, ledger_id),
                json.dumps(body), content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        return response.json()['data']['data']['warnings']

    def test_apply_pdf_params_sets_print_setup(self):
        from apps.equipment_lifecycle.excel_pdf_params import (
            apply_pdf_params,
            validate_pdf_params,
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_MAIN
        sheet['A1'] = '标题'
        sheet.page_setup.scale = 100
        params = validate_pdf_params({
            'paper_size': 'A3', 'orientation': 'landscape',
            'margins': {'top': 2.5, 'left': 1.0},
            'scale_mode': 'percent', 'scale_percent': 80,
            'repeat_header_rows': '1:2', 'row_breaks': [5, 3, 5],
            'page_order': 'over_then_down',
        }, max_row=6)
        self.assertEqual(params['row_breaks'], [3, 5])
        apply_pdf_params(workbook, params, sheet_name=SHEET_MAIN)
        self.assertEqual(int(sheet.page_setup.paperSize), 8)
        self.assertEqual(sheet.page_setup.orientation, 'landscape')
        self.assertEqual(sheet.page_setup.pageOrder, 'overThenDown')
        self.assertEqual(sheet.page_setup.scale, 80)
        self.assertIsNone(sheet.page_setup.fitToWidth)
        self.assertEqual(sheet.page_margins.top, round(2.5 / 2.54, 4))
        self.assertEqual(sheet.page_margins.left, round(1.0 / 2.54, 4))
        self.assertEqual(sheet.print_title_rows, '1:2')
        self.assertEqual([brk.id for brk in sheet.row_breaks.brk], [3, 5])

    def test_apply_fit_width_sets_fit_and_pagination_scoped_to_main_sheet(self):
        from apps.equipment_lifecycle.excel_pdf_params import (
            apply_pdf_params,
            validate_pdf_params,
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_MAIN
        sheet['A1'] = '标题'
        sheet.page_setup.scale = 50
        params = validate_pdf_params({'scale_mode': 'fit_width',
                                      'repeat_header_rows': '1:1'})
        apply_pdf_params(workbook, params, sheet_name=SHEET_MAIN)
        self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(sheet.page_setup.fitToHeight, 0)
        self.assertIsNone(sheet.page_setup.scale)
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(sheet.print_title_rows, '1:1')
        other = workbook.create_sheet('其他')
        other['A1'] = '其他表'
        apply_pdf_params(workbook, params, sheet_name=SHEET_MAIN)
        self.assertIsNone(other.print_title_rows)

    def test_export_pdf_defaults_to_version_snapshot_params(self):
        template, version, bindings = self._publish_with_params(LANDSCAPE_PARAMS)
        ledger_id = self._create_filled_ledger(version['id'], bindings)
        workbook = self._capture_pdf_workbook(ledger_id)
        self.assertEqual(workbook[SHEET_MAIN].page_setup.orientation, 'landscape')
        self.assertEqual(int(workbook[SHEET_MAIN].page_setup.paperSize), 8)
        self.assertEqual(workbook[SHEET_MAIN].page_setup.scale, 80)
        self.assertEqual(workbook[SHEET_MAIN].print_title_rows, '1:2')
        self.assertEqual([brk.id for brk in workbook[SHEET_MAIN].row_breaks.brk],
                         [4])

    def test_export_pdf_override_params_only_for_this_call(self):
        template, version, bindings = self._publish_with_params(LANDSCAPE_PARAMS)
        ledger_id = self._create_filled_ledger(version['id'], bindings)
        workbook = self._capture_pdf_workbook(
            ledger_id, {'pdf_params': {'orientation': 'portrait',
                                       'scale_mode': 'fit_width'}})
        self.assertEqual(workbook[SHEET_MAIN].page_setup.orientation, 'portrait')
        self.assertEqual(workbook[SHEET_MAIN].page_setup.fitToWidth, 1)
        # 覆盖仅当次生效：默认导出仍用版本快照参数
        workbook = self._capture_pdf_workbook(ledger_id)
        self.assertEqual(workbook[SHEET_MAIN].page_setup.orientation, 'landscape')

    def test_history_ledger_keeps_own_version_params(self):
        template, version, bindings = self._publish_with_params(LANDSCAPE_PARAMS)
        ledger_id = self._create_filled_ledger(version['id'], bindings)
        # 模板后续修改参数并发布 v2（纵向），v1 台账导出仍用 v1 快照
        response = self._save_params(template['id'], {'orientation': 'portrait'})
        self.assertEqual(response.json().get('error', ''), '')
        self.publish(template['id'])
        workbook = self._capture_pdf_workbook(ledger_id)
        self.assertEqual(workbook[SHEET_MAIN].page_setup.orientation, 'landscape')

    def test_export_check_clean_when_fit_width_and_filled(self):
        template, version, bindings = self._publish_with_params(None)
        ledger_id = self._create_filled_ledger(version['id'], bindings)
        self.assertEqual(self._check_warnings(ledger_id), [])

    def test_export_check_reports_missing_required_and_overflow(self):
        template, version, bindings = self._publish_with_params(None)
        ledger_id = self._create_filled_ledger(version['id'], [])
        warnings = self._check_warnings(ledger_id)
        self.assertTrue(any('必填字段未填写' in item for item in warnings))
        # 样例内容宽于 A4 一页：覆盖为原尺寸 + A4 纵向应给出横向分页提示
        warnings = self._check_warnings(ledger_id, {'pdf_params': OVERFLOW_PARAMS})
        self.assertTrue(any('横向分页' in item for item in warnings))

    @skipUnless(detect_soffice(), 'LibreOffice 不可用，真实转换用例 NOT_RUN')
    def test_export_pdf_real_conversion_with_params(self):
        template, version, bindings = self._publish_with_params(LANDSCAPE_PARAMS)
        ledger_id = self._create_filled_ledger(version['id'], bindings)
        response = self.owner_client.post(
            '%s/%s/export/pdf/' % (LEDGER_URL, ledger_id))
        self.assertIn('application/pdf', response['Content-Type'])
        self.assertTrue(response.content.startswith(b'%PDF-'))
        self.assertGreater(len(response.content), 1000)

    def test_export_check_requires_export_permission(self):
        template, version, bindings = self._publish_with_params(None)
        ledger_id = self._create_filled_ledger(version['id'], bindings)
        limited = make_user('excel_noexp_%s' % uuid.uuid4().hex[:8],
                            perms=P_TPL + [P_LEDGER[0], P_LEDGER[1]])
        response = make_client(limited).post(
            '%s/%s/export/check/' % (LEDGER_URL, ledger_id))
        self.assertEqual(response.json()['error'], '权限拒绝')

    def test_export_check_cross_tenant_rejected(self):
        template, version, bindings = self._publish_with_params(None)
        ledger_id = self._create_filled_ledger(version['id'], bindings)
        response = self.other_client.post(
            '%s/%s/export/check/' % (LEDGER_URL, ledger_id))
        self.assertEqual(response.json()['error'], '台账不存在')


def _narrow_workbook():
    """窄表（内容约 3cm，纵向一页放得下）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_MAIN
    sheet['A1'] = '窄表'
    sheet.column_dimensions['A'].width = 10
    return workbook


def _wide_workbook():
    """宽表（两列各 80 字符，内容约 30cm，纵向放不下）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_MAIN
    sheet['A1'] = '宽表甲'
    sheet['B1'] = '宽表乙'
    sheet.column_dimensions['A'].width = 80
    sheet.column_dimensions['B'].width = 80
    return workbook


class PdfPresetModeTests(ExcelLedgerContractBase):
    """打印模式扩展（2026-10-01）：三档预设/自动方向/居中/打印区域。"""

    def _save(self, template_id, params):
        return self.owner_client.put(
            '%s/%s/pdf-params/' % (TPL_URL, template_id),
            json.dumps({'pdf_params': params}),
            content_type='application/json')

    def _capture_pdf_workbook(self, ledger_id):
        """导出 PDF 打桩：捕获 LibreOffice 转换前的 xlsx 字节。"""
        from apps.equipment_lifecycle import excel_template_services
        with mock.patch.object(excel_template_services, '_xlsx_bytes_to_pdf',
                               return_value=b'%PDF-captured') as fake:
            response = self.owner_client.post(
                '%s/%s/export/pdf/' % (LEDGER_URL, ledger_id))
            self.assertIn('application/pdf', response['Content-Type'])
            return load_workbook(io.BytesIO(fake.call_args[0][0]))

    def _upload_bytes(self, workbook, name='preset.xlsx'):
        buffer = io.BytesIO()
        workbook.save(buffer)
        response = self.owner_client.post(TPL_URL + '/', {
            'file': SimpleUploadedFile(name, buffer.getvalue()),
            'name': '打印模式模板', 'ledger_type': 'week_month',
        })
        payload = response.json()
        self.assertEqual(payload.get('error', ''), '')
        return payload['data']['data']

    def test_seed_preset_original_for_file_with_print_config(self):
        """样例带 fitToWidth（用户配置过打印）→ original 模式。"""
        template = self.upload_sample()
        params = template['draft']['pdf_params']
        self.assertEqual(params['preset'], 'original')
        self.assertEqual(params['paper_size'], 'A4')
        self.assertEqual(params['orientation'], 'portrait')

    def test_seed_preset_a4_auto_for_boilerplate_only_file(self):
        """仅纸张/方向样板值（WPS 总是序列化）→ a4_auto 模式。"""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_MAIN
        sheet['A1'] = '标题'
        # WPS 样板值：A4 纵向 + 默认页边距，无 fit/scale/打印区域
        sheet.page_setup.orientation = 'portrait'
        sheet.page_setup.paperSize = 9
        template = self._upload_bytes(workbook)
        params = template['draft']['pdf_params']
        self.assertEqual(params['preset'], 'a4_auto')
        self.assertEqual(params['paper_size'], 'A4')
        self.assertEqual(params['orientation'], 'auto')
        self.assertEqual(params['scale_mode'], 'fit_width')

    def test_auto_orientation_resolves_by_content_width(self):
        """方向自动：纵向放不下的宽表转横向，窄表保持纵向。"""
        from apps.equipment_lifecycle.excel_pdf_params import (
            apply_pdf_params,
            validate_pdf_params,
        )
        params = validate_pdf_params({'preset': 'a4_auto',
                                      'orientation': 'auto',
                                      'scale_mode': 'fit_width'})
        wide = _wide_workbook()
        apply_pdf_params(wide, params, sheet_name=SHEET_MAIN)
        self.assertEqual(wide[SHEET_MAIN].page_setup.orientation, 'landscape')
        self.assertEqual(wide[SHEET_MAIN].page_setup.fitToWidth, 1)
        self.assertEqual(wide[SHEET_MAIN].page_setup.fitToHeight, 0)
        narrow = _narrow_workbook()
        apply_pdf_params(narrow, params, sheet_name=SHEET_MAIN)
        self.assertEqual(narrow[SHEET_MAIN].page_setup.orientation, 'portrait')

    def test_original_preset_skips_page_setup_apply(self):
        """original 模式：源文件页面设置原样保留（完全尊重上传文件）。"""
        from apps.equipment_lifecycle.excel_pdf_params import apply_pdf_params
        workbook = _narrow_workbook()
        sheet = workbook[SHEET_MAIN]
        sheet.page_setup.scale = 55
        sheet.page_setup.paperSize = 8
        sheet.page_setup.orientation = 'landscape'
        apply_pdf_params(workbook, {'preset': 'original'}, sheet_name=SHEET_MAIN)
        self.assertEqual(sheet.page_setup.scale, 55)
        self.assertEqual(int(sheet.page_setup.paperSize), 8)
        self.assertEqual(sheet.page_setup.orientation, 'landscape')

    def test_apply_print_area_and_centering(self):
        """打印区域与水平/垂直居中：主工作表生效，分页细节不外溢。"""
        from apps.equipment_lifecycle.excel_pdf_params import (
            apply_pdf_params,
            validate_pdf_params,
        )
        workbook = _narrow_workbook()
        other = workbook.create_sheet('其他')
        other['A1'] = '其他表'
        params = validate_pdf_params({'print_area': 'a1:b2',
                                      'horizontal_centered': True,
                                      'vertical_centered': True})
        self.assertEqual(params['print_area'], 'A1:B2')
        apply_pdf_params(workbook, params, sheet_name=SHEET_MAIN)
        sheet = workbook[SHEET_MAIN]
        self.assertEqual(list(sheet.print_area), ['$A$1:$B$2'])
        self.assertTrue(sheet.print_options.horizontalCentered)
        self.assertTrue(sheet.print_options.verticalCentered)
        # 打印区域只作用于主工作表；居中作用于全部工作表
        self.assertFalse(other.print_area)
        self.assertTrue(other.print_options.horizontalCentered)

    def test_print_area_validation_rejects_out_of_bounds(self):
        template = self.upload_sample()
        for bad in ('A1:ZZ99', 'A99:B200', 'C1:A3', '不是区域'):
            response = self._save(template['id'], {'print_area': bad})
            self.assertTrue(response.json().get('error'), '应拒绝：%s' % bad)
        response = self._save(template['id'], {'print_area': 'A1:H6'})
        self.assertEqual(response.json().get('error', ''), '')
        saved = response.json()['data']['data']['draft']['pdf_params']
        self.assertEqual(saved['print_area'], 'A1:H6')

    def test_fit_page_small_scale_warning(self):
        """单页表单：内容高度超页导致缩放过小时预检告警，可读时不告警。"""
        from apps.equipment_lifecycle.excel_pdf_params import estimate_pdf_warnings
        meta = {
            'sheet_name': 'Sheet1',
            'field_bindings': [],
            'layout': {'sheets': [{
                'name': 'Sheet1', 'max_row': 200, 'max_column': 2,
                'cells': {'A1': {'v': '标题'}},
                'column_widths': {'A': 12, 'B': 12},
                'row_heights': {}, 'hidden_rows': [], 'hidden_columns': [],
            }]},
        }
        params = {'preset': 'single_page', 'scale_mode': 'fit_page'}
        warnings = estimate_pdf_warnings(meta, {}, params)
        self.assertTrue(any('缩放' in item and '过小' in item for item in warnings))
        short_meta = json.loads(json.dumps(meta))
        short_meta['layout']['sheets'][0]['max_row'] = 40
        self.assertEqual(
            [w for w in estimate_pdf_warnings(short_meta, {}, params)
             if '过小' in w], [])

    def test_export_uses_preset_and_auto_orientation_from_snapshot(self):
        """发布快照携带 preset/auto：导出按快照解析实际方向（宽表横向）。"""
        template = self.upload_sample()
        bindings = [self.make_binding('C4', '第1信道-主电台-PS模块')]
        self.assertEqual(self.save_bindings(template['id'], bindings)
                         .json().get('error', ''), '')
        response = self._save(template['id'], {
            'preset': 'a4_auto', 'paper_size': 'A4', 'orientation': 'auto',
            'scale_mode': 'fit_width', 'horizontal_centered': True})
        self.assertEqual(response.json().get('error', ''), '')
        version = self.publish(template['id'])
        response = self.owner_client.get(
            '%s/%s/versions/1/' % (TPL_URL, template['id']))
        snapshot = response.json()['data']['data']
        self.assertEqual(snapshot['pdf_params']['preset'], 'a4_auto')
        self.assertEqual(snapshot['pdf_params']['orientation'], 'auto')
        ledger_id = self._create_filled_ledger_by_id(version['id'], bindings)
        workbook = self._capture_pdf_workbook(ledger_id)
        # 样例主表内容约 20.7cm > A4 纵向可打印宽度 → 自动解析为横向
        self.assertEqual(workbook[SHEET_MAIN].page_setup.orientation, 'landscape')
        self.assertEqual(workbook[SHEET_MAIN].page_setup.fitToWidth, 1)
        self.assertEqual(workbook[SHEET_MAIN].page_setup.fitToHeight, 0)
        self.assertTrue(workbook[SHEET_MAIN].print_options.horizontalCentered)

    def test_history_ledger_keeps_own_preset_snapshot(self):
        """历史台账不受新版本影响：v1 单页表单，v2 改 a4_auto 后 v1 台账
        导出仍为 1 页 × 1 页（fitToHeight=1）。"""
        template = self.upload_sample()
        bindings = [self.make_binding('C4', '第1信道-主电台-PS模块')]
        self.assertEqual(self.save_bindings(template['id'], bindings)
                         .json().get('error', ''), '')
        response = self._save(template['id'], {
            'preset': 'single_page', 'paper_size': 'A4',
            'orientation': 'portrait', 'scale_mode': 'fit_page'})
        self.assertEqual(response.json().get('error', ''), '')
        version_v1 = self.publish(template['id'])
        ledger_id = self._create_filled_ledger_by_id(version_v1['id'], bindings)
        response = self._save(template['id'], {
            'preset': 'a4_auto', 'orientation': 'auto',
            'scale_mode': 'fit_width'})
        self.assertEqual(response.json().get('error', ''), '')
        self.publish(template['id'])
        workbook = self._capture_pdf_workbook(ledger_id)
        self.assertEqual(workbook[SHEET_MAIN].page_setup.fitToHeight, 1)
        self.assertEqual(workbook[SHEET_MAIN].page_setup.fitToWidth, 1)

    def _create_filled_ledger_by_id(self, version_id, bindings):
        response = self.owner_client.post(
            LEDGER_URL + '/',
            json.dumps({'version_id': version_id,
                        'executed_date': '2026-10-01'}),
            content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        ledger_id = response.json()['data']['data']['id']
        values = {binding['field_id']: '正常' for binding in bindings}
        response = self.owner_client.put(
            '%s/%s/' % (LEDGER_URL, ledger_id),
            json.dumps({'values': values}), content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        return ledger_id
