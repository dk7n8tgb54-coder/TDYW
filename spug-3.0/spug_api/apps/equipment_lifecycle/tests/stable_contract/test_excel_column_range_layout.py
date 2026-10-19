# -*- coding: utf-8 -*-
"""Excel 模板列宽范围（<col min max>）解析与旧快照读取端修正契约（2026-10-19）。

缺陷背景（真实模板 21 v3 复现）：源文件 sheet1.xml 用单条
<col min="1" max="3" width="36.21"> 统一设定 A:C 三列宽。旧解析按字母
逐列访问 ws.column_dimensions[letter]：openpyxl 只把范围记在起始列 'A'
的 min/max 属性上，B/C 缺键、按下标访问被 openpyxl 用内置缺省宽度
凭空造出 width=13 的假维度——落库快照 A=36.21、B=13、C=13，网页把
B/C 画窄，表格从左页边距开始显示时右侧留白（Excel 打印预览三列同宽，
导出 PDF 以源文件为事实不受影响）。

契约：
1. 解析展开 min/max 范围：A:C 同宽（单条 col 覆盖多列）、单列设置
   （min=max）、隐藏列范围（hidden 覆盖多列）全部逐列生效；
2. 默认列宽：sheetFormatPr.defaultColWidth 显式声明时作为未设列的
   缺省宽度；未声明时不伪造（前端按 Excel 内置缺省 8.43 兜底）；
3. 范围覆盖顺序：重叠 <col> 按 XML 文档顺序后者覆盖前者；
4. 旧快照读取端修正：发布版本不可变，历史快照的坏列宽在读取时以
   版本源文件为准重新推导（版本详情/台账详情/草稿工作区同一出口），
   快照本身与传入 dict 不被改写；源文件缺失原样返回；新快照修正幂等；
5. 同类修正：PDF 参数方向自动（excel_pdf_params）同样按展开后列宽
   估算内容宽度，A:C 同宽范围纵向放不下时正确转横向。

测试工作簿以 openpyxl 写出与真实文件同形的 <col min max> 字节，
全部走真实解析与真实 HTTP 路径，不做源码字符串断言。
"""
import io
import json
import os
import shutil
import tempfile
import uuid

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.dimensions import ColumnDimension

from apps.equipment_lifecycle.excel_models import (
    OperationLedgerExcelTemplate,
    OperationLedgerExcelTemplateVersion,
)
from apps.equipment_lifecycle.excel_pdf_params import apply_pdf_params
from apps.equipment_lifecycle.excel_template_services import (
    parse_workbook,
    refresh_column_layout,
)
from apps.equipment_lifecycle.tests.base import URL_PREFIX
from apps.equipment_lifecycle.tests.stable_contract.test_excel_ledger_templates import (
    LEDGER_URL,
    TPL_URL,
    ExcelLedgerContractBase,
)

SHEET_NAME = '信道维护记录'
RANGE_WIDTH = 36.21
LEGACY_BROKEN_WIDTHS = {'A': 36.21, 'B': 13.0, 'C': 13.0}


def build_range_width_xlsx():
    """A:C 三列由单条 <col min="1" max="3" width="36.21"> 统一设定。

    与模板 21 v3 真实文件同形；内容覆盖到 C 列保证 max_column=3。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws['A1'] = '维护台账'
    ws['C1'] = '备注'
    ws.column_dimensions['A'] = ColumnDimension(
        ws, index='A', width=RANGE_WIDTH, customWidth=True, min=1, max=3)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_hidden_range_xlsx():
    """B:D 由单条 <col min="2" max="4" hidden="1"> 隐藏（内容覆盖到 D 列）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for letter in ('A', 'B', 'C', 'D'):
        ws['%s1' % letter] = '%s列' % letter
    ws.column_dimensions['B'] = ColumnDimension(
        ws, index='B', hidden=True, min=2, max=4)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_overlapping_range_xlsx():
    """两条重叠范围：<col min="1" max="3" width="36.21"> + <col min="3" max="5" width="10">。

    文档顺序后者覆盖前者：C 列归后者（10），A/B 归前者（36.21）。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for letter in ('A', 'B', 'C', 'D', 'E'):
        ws['%s1' % letter] = '%s列' % letter
    ws.column_dimensions['A'] = ColumnDimension(
        ws, index='A', width=RANGE_WIDTH, customWidth=True, min=1, max=3)
    ws.column_dimensions['C'] = ColumnDimension(
        ws, index='C', width=10.0, customWidth=True, min=3, max=5)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_default_width_xlsx(with_default):
    """单列设置 + 可选 sheetFormatPr.defaultColWidth（内容覆盖到 C 列）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for letter in ('A', 'B', 'C'):
        ws['%s1' % letter] = '%s列' % letter
    ws.column_dimensions['B'] = ColumnDimension(
        ws, index='B', width=10.0, customWidth=True, min=2, max=2)
    if with_default:
        ws.sheet_format.defaultColWidth = 9
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ColumnLayoutParseTests(SimpleTestCase):
    """真实字节解析：<col min max> 范围展开、默认列宽与覆盖顺序。"""

    def _parse(self, data):
        return parse_workbook('column-range.xlsx', data)

    def _source_file(self, data):
        directory = tempfile.mkdtemp(prefix='el_column_range_test_')
        self.addCleanup(shutil.rmtree, directory, ignore_errors=True)
        path = '%s/source.xlsx' % directory
        with open(path, 'wb') as f:
            f.write(data)
        return path

    def test_range_width_expanded_to_all_covered_columns(self):
        """单条 <col min="1" max="3" width="36.21">：A:C 三列全部 36.21。"""
        sheet = self._parse(build_range_width_xlsx())['layout']['sheets'][0]
        self.assertEqual(sheet['max_column'], 3)
        self.assertEqual(
            sheet['column_widths'],
            {'A': RANGE_WIDTH, 'B': RANGE_WIDTH, 'C': RANGE_WIDTH})

    def test_hidden_range_expanded_to_all_covered_columns(self):
        """单条 <col min="2" max="4" hidden="1">：B/C/D 全部隐藏。"""
        sheet = self._parse(build_hidden_range_xlsx())['layout']['sheets'][0]
        self.assertEqual(sheet['hidden_columns'], ['B', 'C', 'D'])

    def test_single_column_setting_stays_single(self):
        """单列设置（min=max=2）：仅 B 列 10，其余列不伪造宽度。"""
        sheet = self._parse(build_default_width_xlsx(with_default=False))['layout']['sheets'][0]
        self.assertEqual(sheet['column_widths'], {'B': 10.0})

    def test_sheet_default_col_width_fills_unspecified_columns(self):
        """defaultColWidth=9 显式声明：未设宽度的 A/C 列按 9 落库。"""
        sheet = self._parse(build_default_width_xlsx(with_default=True))['layout']['sheets'][0]
        self.assertEqual(sheet['column_widths'], {'A': 9.0, 'B': 10.0, 'C': 9.0})

    def test_overlapping_ranges_last_definition_wins(self):
        """重叠范围按 XML 文档顺序后者覆盖前者：C 列归后者（10）。"""
        sheet = self._parse(build_overlapping_range_xlsx())['layout']['sheets'][0]
        self.assertEqual(
            sheet['column_widths'],
            {'A': RANGE_WIDTH, 'B': RANGE_WIDTH, 'C': 10.0, 'D': 10.0, 'E': 10.0})

    def test_auto_orientation_uses_expanded_range_widths(self):
        """方向自动：A:C 同宽范围按展开后列宽估算，纵向放不下转横向。

        旧估算（excel_pdf_params 同类缺陷）把 B/C 按缺省宽度低估，
        内容宽度约 10cm 判纵向；展开后约 20.5cm 超过 A4 纵向可打印
        宽度，正确转横向。
        """
        meta = self._parse(build_range_width_xlsx())
        workbook = load_workbook(io.BytesIO(build_range_width_xlsx()))
        apply_pdf_params(workbook, meta['pdf_params'], meta['sheet_name'])
        self.assertEqual(workbook.active.page_setup.orientation, 'landscape')

    def test_refresh_fixes_legacy_snapshot_without_mutation(self):
        """旧快照（A=36.21/B=13/C=13）读取修正；传入 dict 原样不被改写。"""
        source = self._source_file(build_range_width_xlsx())
        layout = {'active_sheet': SHEET_NAME, 'sheets': [
            {'name': SHEET_NAME, 'column_widths': dict(LEGACY_BROKEN_WIDTHS),
             'hidden_columns': []}]}
        corrected = refresh_column_layout(layout, source)
        self.assertEqual(
            corrected['sheets'][0]['column_widths'],
            {'A': RANGE_WIDTH, 'B': RANGE_WIDTH, 'C': RANGE_WIDTH})
        self.assertEqual(layout['sheets'][0]['column_widths'],
                         LEGACY_BROKEN_WIDTHS)

    def test_refresh_idempotent_for_new_snapshot(self):
        """新解析快照（已展开）再修正结果不变。"""
        source = self._source_file(build_range_width_xlsx())
        meta = parse_workbook('column-range.xlsx', build_range_width_xlsx())
        layout = meta['layout']
        self.assertIs(refresh_column_layout(layout, source)['sheets'][0],
                      layout['sheets'][0])

    def test_refresh_returns_snapshot_unchanged_when_source_missing(self):
        """源文件缺失/损坏：原样返回快照，不阻断详情。"""
        layout = {'active_sheet': SHEET_NAME, 'sheets': [
            {'name': SHEET_NAME, 'column_widths': dict(LEGACY_BROKEN_WIDTHS),
             'hidden_columns': []}]}
        self.assertIs(refresh_column_layout(layout, '/tmp/not-exist.xlsx'),
                      layout)
        directory = tempfile.mkdtemp(prefix='el_column_range_test_')
        self.addCleanup(shutil.rmtree, directory, ignore_errors=True)
        broken_path = os.path.join(directory, 'broken.xlsx')
        with open(broken_path, 'wb') as f:
            f.write(b'broken')
        self.assertIs(refresh_column_layout(layout, broken_path), layout)


class ColumnRangeContractTests(ExcelLedgerContractBase):
    """范围列宽全链路（真实 HTTP：上传解析→发布快照→版本/台账/草稿读取）。"""

    def _upload_range_template(self):
        """上传 A:C 同宽范围工作簿并绑定 B1 → 发布，返回 (template, version)。"""
        response = self.owner_client.post(TPL_URL + '/', {
            'file': SimpleUploadedFile('column-range.xlsx',
                                       build_range_width_xlsx()),
            'name': '同宽范围列宽模板',
            'ledger_type': 'week_month',
        })
        payload = response.json()
        self.assertEqual(payload.get('error', ''), '')
        template = payload['data']['data']
        binding = {
            'field_id': str(uuid.uuid4()),
            'sheet_name': SHEET_NAME,
            'cell_coordinate': 'B1',
            'field_name': '备注',
            'input_type': 'status',
            'is_required': True,
            'unit': '',
            'options': [],
        }
        response = self.owner_client.put(
            '%s/%s/bindings/' % (TPL_URL, template['id']),
            json.dumps({'field_bindings': [binding]}),
            content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        response = self.owner_client.post(
            '%s/%s/publish/' % (TPL_URL, template['id']))
        payload = response.json()
        self.assertEqual(payload.get('error', ''), '')
        return template, payload['data']['data']

    def _create_ledger(self, version_id):
        response = self.owner_client.post(
            LEDGER_URL + '/',
            json.dumps({'version_id': version_id, 'ledger_type': 'week_month'}),
            content_type='application/json')
        payload = response.json()
        self.assertEqual(payload.get('error', ''), '')
        return payload['data']['data']

    def _corrupt_version_layout(self, version_id):
        """把版本快照改写为旧解析产物的坏列宽（模拟修复前发布的数据）。"""
        version = OperationLedgerExcelTemplateVersion.objects.get(pk=version_id)
        part = version.get_layout()
        part['layout']['sheets'][0]['column_widths'] = dict(LEGACY_BROKEN_WIDTHS)
        OperationLedgerExcelTemplateVersion.objects.filter(pk=version_id).update(
            layout=json.dumps(part, ensure_ascii=False, separators=(',', ':')))

    def test_version_and_ledger_detail_carry_expanded_widths(self):
        """发布快照即正确：版本详情与台账详情的 A:C 三列同宽。"""
        template, version = self._upload_range_template()
        detail = self.owner_client.get(
            '%s/%s/versions/%s/' % (TPL_URL, template['id'], version['version_no'])
        ).json()['data']['data']
        sheet = detail['layout']['sheets'][0]
        self.assertEqual(
            sheet['column_widths'],
            {'A': RANGE_WIDTH, 'B': RANGE_WIDTH, 'C': RANGE_WIDTH})
        ledger = self._create_ledger(version['id'])
        excel = self.owner_client.get(
            '%s/%s/' % (LEDGER_URL, ledger['id'])).json()['data']['data']['excel']
        self.assertEqual(
            excel['version']['layout']['sheets'][0]['column_widths'],
            {'A': RANGE_WIDTH, 'B': RANGE_WIDTH, 'C': RANGE_WIDTH})

    def test_legacy_broken_snapshot_corrected_on_read_only(self):
        """历史坏快照：台账/版本详情读取时修正，快照本身不被改写。"""
        template, version = self._upload_range_template()
        ledger = self._create_ledger(version['id'])
        self._corrupt_version_layout(version['id'])
        excel = self.owner_client.get(
            '%s/%s/' % (LEDGER_URL, ledger['id'])).json()['data']['data']['excel']
        self.assertEqual(
            excel['version']['layout']['sheets'][0]['column_widths'],
            {'A': RANGE_WIDTH, 'B': RANGE_WIDTH, 'C': RANGE_WIDTH})
        detail = self.owner_client.get(
            '%s/%s/versions/%s/' % (TPL_URL, template['id'], version['version_no'])
        ).json()['data']['data']
        self.assertEqual(
            detail['layout']['sheets'][0]['column_widths'],
            {'A': RANGE_WIDTH, 'B': RANGE_WIDTH, 'C': RANGE_WIDTH})
        stored = OperationLedgerExcelTemplateVersion.objects.get(
            pk=version['id']).get_layout()
        self.assertEqual(stored['layout']['sheets'][0]['column_widths'],
                         LEGACY_BROKEN_WIDTHS)

    def test_legacy_broken_draft_corrected_on_template_detail(self):
        """历史坏草稿：模板详情（绑定工作区数据源）读取时修正，草稿不被改写。"""
        template, _version = self._upload_range_template()
        tpl = OperationLedgerExcelTemplate.objects.get(pk=template['id'])
        part = json.loads(tpl.draft_layout)
        part['layout']['sheets'][0]['column_widths'] = dict(LEGACY_BROKEN_WIDTHS)
        OperationLedgerExcelTemplate.objects.filter(pk=template['id']).update(
            draft_layout=json.dumps(part, ensure_ascii=False, separators=(',', ':')))
        detail = self.owner_client.get(
            '%s/%s/' % (TPL_URL, template['id'])).json()['data']['data']
        self.assertEqual(
            detail['draft']['layout']['sheets'][0]['column_widths'],
            {'A': RANGE_WIDTH, 'B': RANGE_WIDTH, 'C': RANGE_WIDTH})
        stored = json.loads(
            OperationLedgerExcelTemplate.objects.get(pk=template['id']).draft_layout)
        self.assertEqual(stored['layout']['sheets'][0]['column_widths'],
                         LEGACY_BROKEN_WIDTHS)
