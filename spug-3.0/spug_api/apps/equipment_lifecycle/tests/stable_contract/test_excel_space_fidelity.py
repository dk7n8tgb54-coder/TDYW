# -*- coding: utf-8 -*-
"""Excel 路线空格保真稳定契约（2026-10-01）。

覆盖任务书「空格保真」验收项（真实代码路径：上传解析 → 布局 JSON →
填写值保存 → 导出 Excel → 导出 PDF，转换边界不打桩）：

1. 上传解析：模板单元格中的中间空格、连续空格、首部/尾随空格、换行符
   原样进入布局 JSON（逐字符断言，不 normalize）；
2. 填写值：文本类型的填写值（含首尾空格/连续空格）保存后原样返回，
   服务端不 strip 文本填写值（受控类型单选仍精确匹配候选原文）；
3. 导出 Excel：模板文本与填写值写回后与原始字符串逐字符一致；
4. 导出 PDF：LibreOffice 转换后按 PDF 字形逐字符断言空格存在（需要
   pdfminer.six 与 LibreOffice，缺失时显式 NOT_RUN）。

样本单元与「ICS 200/60 内话系统周月维护记录表」同型：标题中间空格
（ICS 200/60）、「温 度」、「年 月 日」前后空格、维护记录留空线式
连续空格。
"""
import io
import json
import uuid
from unittest import skipUnless

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from apps.equipment_lifecycle.excel_pdf_service import detect_soffice
from apps.equipment_lifecycle.tests.stable_contract.test_excel_ledger_templates import (
    LEDGER_URL,
    TPL_URL,
    ExcelLedgerContractBase,
)

try:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LTChar, LTTextContainer, LTTextLine
    HAS_PDFMINER = True
except ImportError:
    HAS_PDFMINER = False

# 与 ICS 200/60 内话系统周月维护记录表同型的空白场景
TITLE_WITH_SPACE = 'ICS 200/60内话系统周月维护记录表'
TEMP_WITH_SPACE = '温 度'
RUN_OF_30_SPACES = 'A' + ' ' * 30 + 'B'
DATE_WITH_LEADING = '     年   月   日'
TRAILING_SPACE = '尾随空格 '
MULTI_LINE = 'line1\nline2'

# 填写值（文本字段）：首尾空格 + 连续空格全部保留
FILLED_VALUE = '  温 度 25.5  '

TEMPLATE_CELLS = {
    'A1': TITLE_WITH_SPACE,
    'B1': TEMP_WITH_SPACE,
    'C1': RUN_OF_30_SPACES,
    'D1': DATE_WITH_LEADING,
    'E1': TRAILING_SPACE,
    'F1': MULTI_LINE,
}

SHEET_NAME = '空白保真'


def build_space_sample_xlsx():
    """构造含各类空白场景的最小模板（A1-F1 模板文本 + A3/C3 可绑定空格）。

    F1 开启 wrap_text（Excel 中换行内容本就需要换行显示，LibreOffice
    才会把两行都画出来）；列宽保证文本不溢出裁剪。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    for coordinate, value in TEMPLATE_CELLS.items():
        sheet[coordinate] = value
    sheet['F1'].alignment = Alignment(wrap_text=True)
    sheet['A5'] = '填写区'  # 扩展行数，A3/C3 保持空单元格可绑定
    sheet.column_dimensions['A'].width = 24
    sheet.column_dimensions['C'].width = 30
    sheet.column_dimensions['D'].width = 18
    sheet.column_dimensions['E'].width = 16
    sheet.column_dimensions['F'].width = 14
    return workbook


def _upload_space_template(client):
    buffer = io.BytesIO()
    build_space_sample_xlsx().save(buffer)
    data = buffer.getvalue()
    response = client.post(TPL_URL + '/', {
        'file': SimpleUploadedFile('空白保真样例.xlsx', data),
        'name': '空白保真模板',
        'ledger_type': 'week_month',
    })
    payload = response.json()
    return payload, data


def _text_binding(coordinate='A3', field_name='温度记录'):
    return {
        'field_id': str(uuid.uuid4()),
        'sheet_name': SHEET_NAME,
        'cell_coordinate': coordinate,
        'field_name': field_name,
        'input_type': 'text',
        'is_required': False,
        'unit': '',
        'options': [],
    }


class SpaceFidelityContractTests(ExcelLedgerContractBase):
    """空格保真：解析/填写值保存/导出 Excel 全链路逐字符一致。"""

    def _sheet_meta(self, template):
        for sheet in template['draft']['layout']['sheets']:
            if sheet['name'] == SHEET_NAME:
                return sheet
        self.fail('缺少工作表')

    def _create_ledger_with_value(self, template, binding, value):
        """发布 → 建台账 → 保存填写值，返回台账 id。"""
        response = self.save_bindings(template['id'], [binding])
        self.assertEqual(response.json().get('error', ''), '')
        version = self.publish(template['id'])
        response = self.owner_client.post(
            LEDGER_URL + '/',
            json.dumps({'version_id': version['id'],
                        }),
            content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        ledger_id = response.json()['data']['data']['id']
        response = self.owner_client.put(
            '%s/%s/' % (LEDGER_URL, ledger_id),
            json.dumps({'values': {binding['field_id']: value}}),
            content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        return ledger_id

    def test_upload_parse_preserves_spaces_in_layout(self):
        """后端解析：模板单元格空格逐字符进入布局 JSON（不 normalize）。"""
        payload, _ = _upload_space_template(self.owner_client)
        self.assertEqual(payload.get('error', ''), '')
        template = payload['data']['data']
        cells = self._sheet_meta(template)['cells']
        for coordinate, expected in TEMPLATE_CELLS.items():
            self.assertEqual(
                cells[coordinate]['v'], expected,
                '单元格 %s 空格丢失：%r != %r'
                % (coordinate, cells[coordinate]['v'], expected))

    def test_text_value_spaces_survive_save_and_reload(self):
        """填写值：文本类型不 strip——首尾空格/连续空格保存后原样。"""
        payload, _ = _upload_space_template(self.owner_client)
        self.assertEqual(payload.get('error', ''), '')
        template = payload['data']['data']
        binding = _text_binding()
        ledger_id = self._create_ledger_with_value(template, binding, FILLED_VALUE)
        response = self.owner_client.get('%s/%s/' % (LEDGER_URL, ledger_id))
        self.assertEqual(response.json().get('error', ''), '')
        values = response.json()['data']['data']['excel']['values']
        self.assertEqual(values[binding['field_id']], FILLED_VALUE)

    def test_export_excel_preserves_spaces(self):
        """导出 Excel：模板文本与填写值写回后与原字符串逐字符一致。"""
        payload, _ = _upload_space_template(self.owner_client)
        self.assertEqual(payload.get('error', ''), '')
        template = payload['data']['data']
        binding = _text_binding()
        ledger_id = self._create_ledger_with_value(template, binding, FILLED_VALUE)
        response = self.owner_client.post(
            '%s/%s/export/excel/' % (LEDGER_URL, ledger_id))
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            response['Content-Type'])
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook[SHEET_NAME]
        for coordinate, expected in TEMPLATE_CELLS.items():
            self.assertEqual(
                sheet[coordinate].value, expected,
                '导出 Excel 单元格 %s 空格丢失：%r'
                % (coordinate, sheet[coordinate].value))
        self.assertEqual(sheet['A3'].value, FILLED_VALUE)

    def test_choice_value_still_requires_exact_option(self):
        """受控类型不受影响：单选匹配候选项（匹配副本清理），非法值拒绝。"""
        payload, _ = _upload_space_template(self.owner_client)
        self.assertEqual(payload.get('error', ''), '')
        template = payload['data']['data']
        binding = _text_binding('C3', '维护结论')
        binding['input_type'] = 'choice'
        binding['options'] = ['正常', '异常']
        ledger_id = self._create_ledger_with_value(template, binding, ' 异常 ')
        response = self.owner_client.get('%s/%s/' % (LEDGER_URL, ledger_id))
        values = response.json()['data']['data']['excel']['values']
        # 匹配候选项后落库为候选原文（无首尾空格）
        self.assertEqual(values[binding['field_id']], '异常')
        # 非候选值仍然拒绝
        response = self.owner_client.put(
            '%s/%s/' % (LEDGER_URL, ledger_id),
            json.dumps({'values': {binding['field_id']: '无效值'}}),
            content_type='application/json')
        self.assertIn('候选项之一', response.json().get('error', ''))


@skipUnless(HAS_PDFMINER and detect_soffice(),
            '需要 LibreOffice 与 pdfminer.six，缺失时本用例 NOT_RUN')
class PdfSpaceFidelityTests(ExcelLedgerContractBase):
    """PDF 导出空格保真：按 PDF 字形逐字符断言（真实 LibreOffice 转换）。"""

    def _prepare_filled_ledger(self):
        payload, _ = _upload_space_template(self.owner_client)
        self.assertEqual(payload.get('error', ''), '')
        template = payload['data']['data']
        binding = _text_binding()
        response = self.save_bindings(template['id'], [binding])
        self.assertEqual(response.json().get('error', ''), '')
        version = self.publish(template['id'])
        response = self.owner_client.post(
            LEDGER_URL + '/',
            json.dumps({'version_id': version['id'],
                        }),
            content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        ledger_id = response.json()['data']['data']['id']
        response = self.owner_client.put(
            '%s/%s/' % (LEDGER_URL, ledger_id),
            json.dumps({'values': {binding['field_id']: FILLED_VALUE}}),
            content_type='application/json')
        self.assertEqual(response.json().get('error', ''), '')
        return ledger_id

    def _page_lines(self, pdf_bytes):
        """按行抽取 PDF 字符序列（LTChar 级，保留每个空格字形）。"""
        lines = []
        for page in extract_pages(io.BytesIO(pdf_bytes)):
            for element in page:
                if not isinstance(element, LTTextContainer):
                    continue
                for line in element:
                    if isinstance(line, LTTextLine):
                        lines.append(''.join(
                            char.get_text() for char in line
                            if isinstance(char, LTChar)))
        return lines

    def test_export_pdf_preserves_spaces_glyphwise(self):
        ledger_id = self._prepare_filled_ledger()
        response = self.owner_client.post(
            '%s/%s/export/pdf/' % (LEDGER_URL, ledger_id))
        self.assertIn('application/pdf', response['Content-Type'])
        self.assertTrue(response.content.startswith(b'%PDF-'))
        joined = '\n'.join(self._page_lines(response.content))
        # 中间空格与 30 连续空格（行内逐字形）
        self.assertIn(RUN_OF_30_SPACES, joined)
        self.assertIn(TEMP_WITH_SPACE, joined)
        # 首部 5 空格 + 字间 3 空格
        self.assertIn(DATE_WITH_LEADING, joined)
        # 尾随空格与换行单元格
        self.assertIn(TRAILING_SPACE.strip(), joined)
        self.assertIn('line1', joined)
        self.assertIn('line2', joined)
        # 填写值（首部空格 + 连续空格）逐字形存在
        self.assertIn('  温 度 25.5', joined)
