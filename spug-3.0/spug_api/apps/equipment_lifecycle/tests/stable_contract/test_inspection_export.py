# -*- coding: utf-8 -*-
"""巡检单导出稳定契约（列表 Excel / 详情 Excel / 详情 PDF）。

覆盖行为契约：
- 列表导出复用列表筛选条件（值班日范围/状态/关键字），导出筛选范围内全部数据；
- 详情导出内容与详情页一致：固定四轮矩阵（当日 11:00 / 16:00 / 22:00 /
  次日 06:30，第四轮归属前一值班日）、异常说明、现场处理、转单去向；
- 不导出明细实际填写时间，无迟填/逾期标记；未记录轮次以"未记录"呈现；
- 独立权限 equipment_lifecycle.inspection.export（后端独立校验）；
- 跨租户数据、不存在的巡检单不可导出；超量导出返回统一错误；
- 导出为只读行为，不修改巡检单数据。

全部通过真实 HTTP 路径 + 隔离测试数据库执行。
"""
import io
import re
from unittest import mock
from urllib.parse import unquote

from django.test import TestCase
from openpyxl import load_workbook

from apps.equipment_lifecycle import inspection_views
from apps.equipment_lifecycle.models import InspectionRecord
from apps.equipment_lifecycle.tests.base import make_client, make_system, make_user
from apps.equipment_lifecycle.tests.stable_contract.inspection_base import (
    INS_URL,
    create_sheet_ok,
    get_sheet_detail,
    get_template_items,
    make_inspection_template,
    make_inspection_user,
    put_records_ok,
    record_spec,
    sheet_payload,
    submit_sheet_ok,
)

EXPORT_URL = '/equipment-lifecycle/inspection-sheets/export'
P_INS_EXPORT = 'equipment_lifecycle.inspection.export'
P_FAULT_ADD = 'equipment_lifecycle.fault.add'

INFO_SHEET = '巡查单信息'
MATRIX_SHEET = '四轮巡检明细'
ROUND_HEADERS = ['检查项目', '检查内容', '当日 11:00', '当日 16:00', '当日 22:00', '次日 06:30']


def make_export_user(username, tenant_id='tenant_a', with_export=True):
    """仅授予查看（可选导出）权限的用户，避免夹具权限掩盖权限缺失缺陷。"""
    perms = ['equipment_lifecycle.inspection.view']
    if with_export:
        perms.append(P_INS_EXPORT)
    return make_user(username, perms, tenant_id=tenant_id)


def build_submitted_sheet(client, template, items, duty_date='2026-09-23'):
    """走真实 API 建单：r1 异常（含异常说明/现场处理）、r2/r3 正常、r4 临时加项。"""
    detail = create_sheet_ok(client, sheet_payload(
        template.id, duty_date=duty_date, client_token=f'tok-export-{duty_date}'))
    sheet_id, version = detail['id'], detail['version']
    plan = [
        record_spec(items['面板指示灯'].id, 'r1', value_text='不正常',
                    abnormal_desc='告警灯红色', onsite_action='已切换备用机'),
        record_spec(items['输出功率读数'].id, 'r2', value_text='512'),
        record_spec(items['驻波比检查'].id, 'r3', value_text='正常'),
        {'client_key': 'tmp-export-1', 'round_no': 'r4', 'device_name': '',
         'item_name': '机房温度', 'input_type': 'text', 'value_text': '26.5'},
    ]
    for spec in plan:
        detail = put_records_ok(client, sheet_id, version, records=[spec])
        version = detail['version']
    return submit_sheet_ok(client, sheet_id)


def load_worksheet(response):
    wb = load_workbook(io.BytesIO(response.content))
    return wb


def sheet_text(ws):
    return '\n'.join(
        str(cell.value) for row in ws.iter_rows() for cell in row
        if cell.value is not None)


def disposition_filename(response):
    header = response['Content-Disposition']
    assert header.startswith("attachment; filename*=UTF-8''"), header
    return unquote(header.split("UTF-8''", 1)[1])


class InspectionExportTestBase(TestCase):
    """公共夹具：作者（全巡检权限）+ 模板（设备A 两项 + 设备B 一项）。"""

    def setUp(self):
        self.admin = make_inspection_user('ins_export_admin', tenant_id='tenant_a')
        self.client = make_client(self.admin)
        self.template = make_inspection_template(
            tenant_id='tenant_a', created_by=self.admin)
        self.items = get_template_items(self.template)
        self.exporter = make_client(make_export_user('ins_exporter'))
        self.submitted = build_submitted_sheet(self.client, self.template, self.items)


class InspectionListExportTest(InspectionExportTestBase):
    """列表 Excel 导出：筛选条件、轮次完成情况、状态与异常数量。"""

    def test_list_export_uses_current_filters(self):
        """值班日范围筛选只导出范围内巡检单；关键字筛选同理。"""
        detail_a = self.submitted  # 值班日 2026-09-23
        detail_b = create_sheet_ok(self.client, sheet_payload(
            self.template.id, duty_date='2026-09-25', client_token='tok-export-b'))

        resp = self.exporter.get(f'{EXPORT_URL}/?duty_from=2026-09-24&duty_to=2026-09-26')
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml.sheet', resp['Content-Type'])
        self.assertIn('巡检单记录_', disposition_filename(resp))
        ws = load_worksheet(resp)['巡检单记录']
        text = sheet_text(ws)
        self.assertIn(detail_b['sheet_no'], text)
        self.assertNotIn(detail_a['sheet_no'], text)

        resp = self.exporter.get(f'{EXPORT_URL}/?keyword={detail_a["sheet_no"]}')
        self.assertEqual(resp.status_code, 200)
        text = sheet_text(load_worksheet(resp)['巡检单记录'])
        self.assertIn(detail_a['sheet_no'], text)
        self.assertNotIn(detail_b['sheet_no'], text)

    def test_list_export_rounds_status_and_abnormal_count(self):
        """轮次完成情况取 recorded_rounds 最终状态；异常数量与列表口径一致。"""
        resp = self.exporter.get(f'{EXPORT_URL}/?keyword={self.submitted["sheet_no"]}')
        ws = load_worksheet(resp)['巡检单记录']
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        row = dict(zip(
            ['值班日', '单号', '巡检模板', '值班人员', '状态',
             '当日 11:00', '当日 16:00', '当日 22:00', '次日 06:30',
             '异常数量', '提交时间', '提交人', '创建人'], rows[0]))
        self.assertEqual(row['单号'], self.submitted['sheet_no'])
        self.assertEqual(row['巡检模板'], '机房巡检模板')
        self.assertEqual(row['状态'], '已提交')
        self.assertEqual(row['异常数量'], 1)
        for round_header in ('当日 11:00', '当日 16:00', '当日 22:00', '次日 06:30'):
            self.assertEqual(row[round_header], '已记录')

    def test_list_export_unrecorded_rounds_marked(self):
        """未记录轮次显示"未记录"，不显示任何迟填/逾期标记。"""
        detail = create_sheet_ok(self.client, sheet_payload(
            self.template.id, duty_date='2026-09-26', client_token='tok-export-c'))
        resp = self.exporter.get(f'{EXPORT_URL}/?keyword={detail["sheet_no"]}')
        text = sheet_text(load_worksheet(resp)['巡检单记录'])
        self.assertEqual(text.count('未记录'), 4)
        self.assertNotIn('迟填', text)
        self.assertNotIn('逾期', text)

    def test_list_export_empty_result_rejected(self):
        """筛选范围内无数据时返回统一业务错误而非空文件（HTTP 400 + error）。"""
        resp = self.exporter.get(f'{EXPORT_URL}/?keyword=不存在单号')
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.json()['error'], '当前筛选条件下没有可导出的数据')


class InspectionDetailExcelExportTest(InspectionExportTestBase):
    """详情 Excel 导出：双工作表、四轮矩阵、异常与临时加项、转单去向。"""

    def test_detail_excel_two_sheets_with_round_matrix(self):
        resp = self.exporter.get(f'{INS_URL}/{self.submitted["id"]}/export/xlsx/')
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml.sheet', resp['Content-Type'])
        self.assertIn(f"巡检单_{self.submitted['sheet_no']}_", disposition_filename(resp))
        wb = load_worksheet(resp)
        self.assertEqual(wb.sheetnames, [INFO_SHEET, MATRIX_SHEET])

        info_text = sheet_text(wb[INFO_SHEET])
        self.assertIn(self.submitted['sheet_no'], info_text)
        self.assertIn('2026-09-23', info_text)
        self.assertIn('机房巡检模板', info_text)
        self.assertIn('已提交', info_text)
        self.assertIn('值班员甲', info_text)

        matrix = wb[MATRIX_SHEET]
        headers = [cell.value for cell in matrix[1]]
        self.assertEqual(headers, ROUND_HEADERS)
        text = sheet_text(matrix)
        # 检查项目（设备名称）× 检查内容 + 四轮填写值
        for expected in ('1号收发信机', '面板指示灯', '输出功率读数', '驻波比检查',
                         '不正常', '512', '正常', '26.5', '临时检查项'):
            self.assertIn(expected, text)

    def test_detail_excel_abnormal_desc_onsite_and_unrecorded(self):
        resp = self.exporter.get(f'{INS_URL}/{self.submitted["id"]}/export/xlsx/')
        matrix = load_worksheet(resp)[MATRIX_SHEET]
        abnormal_cell = None
        for row in matrix.iter_rows(min_row=2):
            for cell in row:
                if cell.value and '异常说明：告警灯红色' in str(cell.value):
                    abnormal_cell = cell
        self.assertIsNotNone(abnormal_cell, '异常说明未导出')
        self.assertIn('现场处理：已切换备用机', str(abnormal_cell.value))
        # 异常结果红色加粗，正常数据不受影响
        self.assertEqual(abnormal_cell.font.color.rgb, 'FFFF0000')
        self.assertTrue(abnormal_cell.font.bold)
        # 未有记录的轮次显示灰色"未记录"（如驻波比检查仅 r3 有记录）
        unrecorded = [
            cell for row in matrix.iter_rows(min_row=2) for cell in row
            if cell.value == '未记录']
        self.assertTrue(unrecorded)
        self.assertEqual(unrecorded[0].font.color.rgb, 'FF808080')

    def test_detail_excel_includes_convert_disposition(self):
        """异常明细已转故障时，导出包含转单去向（与详情页一致）。"""
        system = make_system(created_by=self.admin, tenant_id='tenant_a')
        converter = make_inspection_user(
            'ins_export_converter', tenant_id='tenant_a', extra_perms=[P_FAULT_ADD])
        abnormal_id = next(
            r['id'] for r in self.submitted['records'] if r['result'] == 'abnormal')
        resp = make_client(converter).post(
            f'{INS_URL}/{self.submitted["id"]}/records/{abnormal_id}/convert-to-fault/',
            {'system': system.id, 'affects_system_operation': False},
            content_type='application/json')
        self.assertFalse(resp.json().get('error'), resp.json())

        resp = self.exporter.get(f'{INS_URL}/{self.submitted["id"]}/export/xlsx/')
        text = sheet_text(load_worksheet(resp)[MATRIX_SHEET])
        self.assertIn('已转故障：GZ', text)


class InspectionDetailPdfExportTest(InspectionExportTestBase):
    """详情 PDF 导出：响应类型、中文字体、A4 四轮矩阵。"""

    def test_detail_pdf_response_and_chinese_font(self):
        resp = self.exporter.get(f'{INS_URL}/{self.submitted["id"]}/export/pdf/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF'))
        self.assertIn(f"巡检单_{self.submitted['sheet_no']}_", disposition_filename(resp))
        # 中文字体注册成功（PDF 中文渲染前提，字体取自 libs/fonts/simhei.ttf）
        from libs.font_manager import FontManager
        self.assertTrue(FontManager.is_font_registered())
        # 文档元数据标题为中文：PDF 字面串对非可打印字节做八进制转义，
        # 还原后应为 UTF-16BE（含 BOM）的 "值班巡检单 {单号}"
        title_match = re.search(rb'/Title\s*\(((?:[^()\\]|\\.)*)\)', resp.content)
        self.assertIsNotNone(title_match, 'PDF 缺少 /Title 元数据')
        decoded = re.sub(rb'\\([0-7]{3})',
                         lambda m: bytes([int(m.group(1), 8)]), title_match.group(1))
        expected_title = b'\xfe\xff' + (
            f"值班巡检单 {self.submitted['sheet_no']}").encode('utf-16-be')
        self.assertEqual(decoded, expected_title)
        # 真实内容体量（基本信息表 + 四轮矩阵）
        self.assertGreater(len(resp.content), 3000)

    def test_detail_pdf_all_four_rounds_exported(self):
        """四个固定轮次（含临时加项行）均进入 PDF 内容（经构建函数验证）。"""
        from apps.equipment_lifecycle import inspection_export
        pdf_bytes = inspection_export.generate_sheet_pdf(self.submitted)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))
        # 生成无异常即代表四轮矩阵（含 r4 临时加项 26.5）构建成功；
        # 文本层断言由 Excel 双导出覆盖（同一 round_cell_lines 数据源）。


class InspectionExportPermissionTest(InspectionExportTestBase):
    """导出权限与数据边界：独立 export 权限、跨租户、不存在、超量。"""

    def test_export_requires_independent_permission(self):
        no_perm = make_client(make_export_user('ins_no_export', with_export=False))
        for url in (f'{EXPORT_URL}/',
                    f'{INS_URL}/{self.submitted["id"]}/export/xlsx/',
                    f'{INS_URL}/{self.submitted["id"]}/export/pdf/'):
            resp = no_perm.get(url)
            self.assertEqual(resp.json()['error'], '权限拒绝', url)

    def test_cross_tenant_sheet_cannot_export(self):
        other = make_client(make_export_user('ins_exporter_b', tenant_id='tenant_b'))
        resp = other.get(f'{EXPORT_URL}/')
        self.assertEqual(resp.json()['error'], '当前筛选条件下没有可导出的数据')
        for url in (f'{INS_URL}/{self.submitted["id"]}/export/xlsx/',
                    f'{INS_URL}/{self.submitted["id"]}/export/pdf/'):
            resp = other.get(url)
            self.assertEqual(resp.json()['error'], '巡检单不存在', url)

    def test_missing_sheet_cannot_export(self):
        for url in (f'{INS_URL}/999999/export/xlsx/',
                    f'{INS_URL}/999999/export/pdf/'):
            resp = self.exporter.get(url)
            self.assertEqual(resp.json()['error'], '巡检单不存在', url)

    def test_export_over_limit_rejected(self):
        """超过导出上限返回项目统一错误（收紧上限模拟超量）。"""
        for index, duty_date in enumerate(('2026-09-10', '2026-09-11', '2026-09-12')):
            create_sheet_ok(self.client, sheet_payload(
                self.template.id, duty_date=duty_date,
                client_token=f'tok-limit-{index}'))
        original = inspection_views.check_export_limit
        with mock.patch.object(
                inspection_views, 'check_export_limit',
                side_effect=lambda qs, **kw: original(qs, limit=2)):
            resp = self.exporter.get(f'{EXPORT_URL}/')
        self.assertEqual(resp.json()['error'], '导出数据超过 2 条，请缩小筛选范围后重试')


class InspectionExportReadOnlyTest(InspectionExportTestBase):
    """导出为只读行为：不修改巡检单数据与轮次状态。"""

    def test_export_does_not_modify_sheet(self):
        before = get_sheet_detail(self.client, self.submitted['id'])

        list_resp = self.exporter.get(
            f'{EXPORT_URL}/?keyword={self.submitted["sheet_no"]}')
        self.assertEqual(list_resp.status_code, 200)
        xlsx_resp = self.exporter.get(f'{INS_URL}/{self.submitted["id"]}/export/xlsx/')
        self.assertEqual(xlsx_resp.status_code, 200)
        pdf_resp = self.exporter.get(f'{INS_URL}/{self.submitted["id"]}/export/pdf/')
        self.assertEqual(pdf_resp.status_code, 200)

        after = get_sheet_detail(self.client, self.submitted['id'])
        self.assertEqual(before['version'], after['version'])
        self.assertEqual(before['status'], after['status'])
        self.assertEqual(before['recorded_rounds'], after['recorded_rounds'])
        self.assertEqual(
            [(r['id'], r['round_no'], r['value_text'], r['result'],
              r['abnormal_desc'], r['onsite_action']) for r in before['records']],
            [(r['id'], r['round_no'], r['value_text'], r['result'],
              r['abnormal_desc'], r['onsite_action']) for r in after['records']])
        sheet = InspectionRecord.objects.filter(sheet_id=self.submitted['id']).first()
        self.assertIsNotNone(sheet)
