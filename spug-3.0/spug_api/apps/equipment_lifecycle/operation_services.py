# Copyright: (c) OpenSpug Organization. https://github.com/openspug/spug
# Released under the AGPL-3.0 License.
"""运行处置记录统一服务层（2026-10-11 合并改造）。

职责：
- 统一记录编号生成（JL+日期+当日序号，全局锁防撞号）；
- create_operation_record：三种记录类型的唯一创建入口（设备故障可直接登记）；
- identify_record_as_fault：异常事件在原记录上认定为设备故障
  （更新同一条主记录的类型与状态，保留描述/登记人/附件/全部进展，
  不创建第二张单）；
- 系统健康状态统一计算（原 fault_services 口径收窄为统一主表）：
  只有 设备故障 + 处置中（有效未关闭）+ 明确影响系统运行 的记录参与计算；
  一般记录/异常事件/已关闭/已撤销一律不影响健康状态；
- progress_overview_map：列表聚合统一进展概览（单次查询，无 N+1）；
- build_record_export_workbook：统一记录列表导出 xlsx 构建（单工作表，
  记录基础列纵向合并、进展逐行展平）；
- compute_operation_statistics：只读统计。

履历说明：原 SystemEvent 系统履历已整体退役（2026-10-11），本服务不再写任何
履历事件；系统档案直接查询统一主表等业务表。健康状态变化留痕由业务审计
（record_audit_event）承载。
"""
import io
import math
from contextlib import contextmanager
from datetime import timedelta

from django.db import connection
from django.db.models import Count, OuterRef, Q, Subquery
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.evidence.models import EvidenceAttachment
from libs.tenant_utils import apply_tenant_filter
from apps.equipment_lifecycle.models import LifecycleSystem
from apps.equipment_lifecycle.operation_models import (
    FAULT_IMPACT_TEXT_MAP,
    FAULT_IMPACT_VALUES,
    RECORD_NO_PREFIX,
    RECORD_NO_SEQ_LIMIT,
    RECORD_STATUS_FOLLOWING_UP,
    RECORD_STATUS_HANDLING,
    RECORD_STATUS_RECORDED,
    RECORD_STATUS_TEXT_MAP,
    RECORD_TYPE_ABNORMAL,
    RECORD_TYPE_FAULT,
    RECORD_TYPE_TEXT_MAP,
    RECORD_TYPE_VALUES,
    SystemOperationProgress,
    SystemOperationRecord,
)

# 有效影响优先级（高 → 低）：中断 > 故障 > 降级
HEALTH_IMPACT_PRIORITY = tuple(reversed(FAULT_IMPACT_VALUES))

# 仍会产生健康影响的记录状态：设备故障的处置中（唯一未结案的进行态）。
# closed/cancelled 不再产生影响；一般记录与异常事件不参与健康状态计算。
ACTIVE_FAULT_STATUSES = (RECORD_STATUS_HANDLING,)

HEALTH_RECALC_LOCK_NAME = 'el_health_recalc'
HEALTH_RECALC_LOCK_TIMEOUT = 10

RECORD_NO_LOCK_NAME = 'el_operation_record_no'
RECORD_NO_LOCK_TIMEOUT = 10

AUDIT_TARGET_HEALTH = 'el_system_health'
AUDIT_TARGET_RECORD = 'el_operation_record'


class HealthRecalcLockTimeout(Exception):
    """健康重算锁获取超时：并发过高时 fail-closed，避免状态覆盖。"""


class RecordNoExhausted(Exception):
    """当日记录编号序号已用尽（JLyyyymmdd-9999 已占用）：
    fail-closed，不产生五位序号或重复编号，由调用方转为业务错误。"""


class RecordLockTimeout(Exception):
    """记录编号锁获取超时：并发过高时 fail-closed，避免并发撞号。"""


class IdentifyFaultError(Exception):
    """异常事件认定设备故障业务校验失败（类型/状态不满足等）。"""


@contextmanager
def health_recalc_lock():
    """健康状态重算全局串行锁（GET_LOCK）。

    用法（视图层）：在事务外获取锁 → 事务内完成记录变更 + 重算 → 事务提交 → 释放锁。
    """
    cursor = connection.cursor()
    try:
        cursor.execute('SELECT GET_LOCK(%s, %s)',
                       [HEALTH_RECALC_LOCK_NAME, HEALTH_RECALC_LOCK_TIMEOUT])
        got = cursor.fetchone()[0]
        if got != 1:
            raise HealthRecalcLockTimeout('系统繁忙，请稍后重试')
        try:
            yield
        finally:
            try:
                cursor.execute('SELECT RELEASE_LOCK(%s)', [HEALTH_RECALC_LOCK_NAME])
            except Exception:
                connection.close()
    finally:
        try:
            cursor.close()
        except Exception:
            pass


@contextmanager
def record_no_lock():
    """统一记录编号生成全局串行锁（GET_LOCK，防并发创建撞号）。

    用法（视图层）：在事务外获取锁 → 事务内完成查重 + 创建 → 事务提交 → 释放锁。
    ATOMIC_REQUESTS 下持锁视图必须以 non_atomic_requests 装配并在视图内
    显式开启事务。
    """
    cursor = connection.cursor()
    try:
        cursor.execute('SELECT GET_LOCK(%s, %s)',
                       [RECORD_NO_LOCK_NAME, RECORD_NO_LOCK_TIMEOUT])
        got = cursor.fetchone()[0]
        if got != 1:
            raise RecordLockTimeout('系统繁忙，请稍后重试')
        try:
            yield
        finally:
            try:
                cursor.execute('SELECT RELEASE_LOCK(%s)', [RECORD_NO_LOCK_NAME])
            except Exception:
                connection.close()
    finally:
        try:
            cursor.close()
        except Exception:
            pass


def generate_record_no():
    """生成统一记录编号 JL{yyyymmdd}-{seq:04d}（须在编号锁 + 创建事务内调用）。

    - 日期取创建当天（不取用户填写的发生时间 occurred_at）；
    - 当日序号跨租户共享：按有效四位数字后缀的最大值加一（缺口不回填）；
    - 序号达 9999 后抛出 RecordNoExhausted，整个创建事务回滚。
    """
    # USE_TZ=False 下 timezone.localdate() 会对 naive now 抛错，直接取本地日期
    prefix = RECORD_NO_PREFIX + timezone.now().date().strftime('%Y%m%d') + '-'
    max_seq = 0
    existing = SystemOperationRecord.objects.filter(
        record_no__startswith=prefix).values_list('record_no', flat=True)
    for record_no in existing:
        suffix = record_no[len(prefix):]
        if len(suffix) == 4 and suffix.isdigit():
            max_seq = max(max_seq, int(suffix))
    if max_seq >= RECORD_NO_SEQ_LIMIT:
        raise RecordNoExhausted(f'当日记录编号已达上限（{prefix}9999），请明日再创建')
    return f'{prefix}{max_seq + 1:04d}'


def user_display_name(user):
    """用户展示姓名：昵称优先、回落用户名（登记人/记录人缺省姓名取此口径）。"""
    if not user:
        return ''
    return (getattr(user, 'nickname', '') or user.username)


def _clean(value, limit):
    """文本入参清洗：None 视为空串、去首尾空白、截断到 limit。"""
    return (value or '').strip()[:limit]


def collect_active_impacts(tenant_id, system_ids):
    """查询系统自身的有效故障影响，返回 {system_id: set(impact)}。

    只按本系统 ID 查询统一主表（record_type=fault、status=处置中、
    affects_system_operation=True）；一般记录/异常事件/已关闭/已撤销
    一律不参与系统健康状态计算。
    """
    system_ids = [int(i) for i in system_ids]
    impacts = {}
    if not system_ids or not tenant_id:
        return impacts
    fault_rows = SystemOperationRecord.objects.filter(
        tenant_id=tenant_id,
        record_type=RECORD_TYPE_FAULT,
        status__in=ACTIVE_FAULT_STATUSES,
        affects_system_operation=True,
        system_id__in=system_ids,
    ).values_list('system_id', 'impact_status')
    for system_id, impact in fault_rows:
        impacts.setdefault(system_id, set()).add(impact)
    return impacts


def derive_health_status(impacts):
    """由影响集合推导健康状态：中断 > 故障 > 降级，无影响 → 正常。"""
    for impact in HEALTH_IMPACT_PRIORITY:
        if impact in impacts:
            return impact
    return LifecycleSystem.HEALTH_NORMAL


def compute_health_status(tenant_id, system_ids):
    """只读计算（不落库）：返回 {system_id: health_status}。幂等、可重复调用。"""
    impacts = collect_active_impacts(tenant_id, system_ids)
    result = {}
    for system_id in system_ids:
        result[int(system_id)] = derive_health_status(impacts.get(int(system_id), set()))
    return result


def recompute_systems_health(tenant_id, system_ids, request=None, audit_reason=None):
    """重算并写入给定系统的健康状态（必须在持锁+事务内调用）。

    - 以"所有有效故障影响的重新查询结果"为准，重复调用结果一致（幂等）；
    - 仅写发生变化的系统，状态变化记录审计（request 为 None 时跳过审计）；
    - 原系统履历（SystemEvent）已退役，不再生成履历事件。
    """
    from apps.logs.audit import record_audit_event

    ids = sorted({int(i) for i in system_ids})
    if not ids or not tenant_id:
        return {}
    impacts = collect_active_impacts(tenant_id, ids)
    result = {}
    # 行锁锁定系统记录（按 id 排序加锁，避免并发事务交叉死锁）
    systems = LifecycleSystem.objects.select_for_update().order_by('id').filter(
        tenant_id=tenant_id, pk__in=ids)
    for system in systems:
        new_status = derive_health_status(impacts.get(system.id, set()))
        result[system.id] = new_status
        if system.health_status == new_status:
            continue
        old_status = system.health_status
        system.health_status = new_status
        system.save(update_fields=['health_status'])
        if request is not None:
            record_audit_event(
                request, 'other', AUDIT_TARGET_HEALTH, system.id, system.name,
                before_value={'health_status': old_status},
                after_value={'health_status': new_status, 'reason': audit_reason or ''})
    return result


def create_operation_record(*, tenant_id, created_by, system, record_type,
                            occurred_at, title, description,
                            registrar_name='', client_token='',
                            category='',
                            affects_system_operation=False, impact_status='',
                            fault_object_name='', fault_object_model='',
                            fault_object_sn='', reason_analysis='',
                            source_type=None, source_ledger_item=None,
                            request=None):
    """统一记录创建服务（唯一创建入口，直接登记故障/台账异常转故障复用）。

    - record_type 必须为三种合法类型之一；fault 默认"部件故障但不影响系统
      运行"（affects=False 时 impact_status 强制空串且不参与健康计算），
      创建即处置中；routine/abnormal 创建即已记录；
    - registrar_name 为登记人（手动输入姓名），缺省取创建人显示姓名
      （台账异常转故障等程序化入口由操作人登记）；
    - source_type/source_ledger_item 供台账异常转故障写入来源关系，
      转换与来源关系保存处于同一事务；
    - 编号由 generate_record_no 在编号锁 + 本事务内分配；
    - 设备故障创建后同事务重算系统健康状态（request 透传以记录审计）。

    返回 SystemOperationRecord 实例。
    """
    if record_type not in RECORD_TYPE_VALUES:
        raise ValueError('记录类型取值不合法')
    is_fault = record_type == RECORD_TYPE_FAULT
    if not is_fault:
        # 非故障记录不携带故障专用字段（后端独立兜底，视图层已拒绝）
        affects_system_operation = False
        impact_status = ''
        fault_object_name = fault_object_model = fault_object_sn = ''
        reason_analysis = ''
    elif not affects_system_operation:
        impact_status = ''
    record = SystemOperationRecord.objects.create(
        tenant_id=tenant_id,
        record_no=generate_record_no(),
        record_type=record_type,
        system=system,
        system_code=system.system_code,
        system_name=system.name,
        system_location=system.install_location,
        occurred_at=occurred_at,
        title=_clean(title, 200),
        description=_clean(description, 2000),
        registrar_name=_clean(registrar_name, 200) or user_display_name(created_by),
        status=RECORD_STATUS_HANDLING if is_fault else RECORD_STATUS_RECORDED,
        category='' if is_fault else _clean(category, 20),
        fault_object_name=_clean(fault_object_name, 100),
        fault_object_model=_clean(fault_object_model, 100),
        fault_object_sn=_clean(fault_object_sn, 100),
        affects_system_operation=bool(affects_system_operation),
        impact_status=_clean(impact_status, 20),
        reason_analysis=_clean(reason_analysis, 2000),
        source_type=source_type or 'manual',
        source_id=str(source_ledger_item.id) if source_ledger_item else '',
        source_ledger_item=source_ledger_item,
        client_token=_clean(client_token, 64),
        created_by=created_by,
    )
    if is_fault:
        recompute_systems_health(tenant_id, [system.id], request=request)
    return record


def identify_record_as_fault(*, record, created_by, occurred_at=None,
                             affects_system_operation=False, impact_status='',
                             reason_analysis='', fault_object_name='',
                             fault_object_model='', fault_object_sn='',
                             request=None):
    """异常事件在原记录上认定为设备故障（须在健康锁+事务内调用，行锁已加）。

    - 仅异常事件（abnormal）且处于 已记录/跟进中 可认定；
    - 更新同一条主记录：record_type→fault、status→处置中，补充故障所需字段
      （影响系统运行/影响程度/原因分析/对象快照）；原描述、登记人、附件与
      全部进展原样保留（进展外键挂主记录，类型变更不受影响）；
    - 发生时间缺省沿用原值（可显式覆盖，不得晚于当前时间——视图层校验）；
    - 同事务重算系统健康状态（认定可能新增有效影响）。

    返回更新后的 record。
    """
    if record.record_type != RECORD_TYPE_ABNORMAL:
        raise IdentifyFaultError('只有异常事件记录可以认定为设备故障')
    if record.status not in (RECORD_STATUS_RECORDED, RECORD_STATUS_FOLLOWING_UP):
        raise IdentifyFaultError('已完成或其他状态的异常事件不能认定为设备故障')
    if affects_system_operation:
        if impact_status not in FAULT_IMPACT_VALUES:
            raise IdentifyFaultError('影响系统运行时必须选择影响程度（降级/故障/中断）')
    else:
        impact_status = ''
    record.record_type = RECORD_TYPE_FAULT
    record.status = RECORD_STATUS_HANDLING
    if occurred_at is not None:
        record.occurred_at = occurred_at
    record.affects_system_operation = bool(affects_system_operation)
    record.impact_status = impact_status or ''
    record.reason_analysis = _clean(reason_analysis, 2000)
    record.fault_object_name = _clean(fault_object_name, 100)
    record.fault_object_model = _clean(fault_object_model, 100)
    record.fault_object_sn = _clean(fault_object_sn, 100)
    record.category = ''
    record.updated_at = timezone.now()
    record.updated_by = created_by
    record.save()
    recompute_systems_health(record.tenant_id, [record.system_id], request=request)
    return record


def attachment_count_map(records, user, module, object_type):
    """Count visible, active attachments for the records on one list page."""
    if not records:
        return {}
    record_keys = {(r.tenant_id, str(r.id)) for r in records}
    attachments = apply_tenant_filter(
        EvidenceAttachment.objects.filter(
            is_deleted=False, module=module, object_type=object_type,
            object_id__in=[str(r.id) for r in records],
            tenant_id__in={r.tenant_id for r in records},
        ), user)
    counts = attachments.values('tenant_id', 'object_id').annotate(
        total=Count('id')).order_by()
    return {(row['tenant_id'], row['object_id']): row['total']
            for row in counts if (row['tenant_id'], row['object_id']) in record_keys}


def progress_overview_map(record_ids):
    """批量统计统一进展概览（列表聚合，单次查询，无 N+1）。

    - progress_count：SystemOperationProgress 行数（追加式时间线）；
    - latest_progress_at / latest_progress_content：按（处理时间, id）最新
      一条进展的时间与内容；无进展的记录返回零值，调用方无需再补默认值。
    返回 {record_id: {...}}。
    """
    record_ids = [int(i) for i in record_ids if i]
    if not record_ids:
        return {}
    latest_qs = SystemOperationProgress.objects.filter(
        record_id=OuterRef('record_id')).order_by('-handled_at', '-id')
    rows = (
        SystemOperationProgress.objects
        .filter(record_id__in=record_ids)
        .values('record_id')
        .annotate(
            progress_count=Count('id'),
            latest_progress_at=Subquery(latest_qs.values('handled_at')[:1]),
            latest_progress_content=Subquery(latest_qs.values('content')[:1]),
        )
    )
    overview = {
        record_id: {
            'progress_count': 0,
            'latest_progress_at': None,
            'latest_progress_content': '',
        }
        for record_id in record_ids
    }
    for row in rows:
        overview[row['record_id']] = row
    return overview


def progress_overview(record):
    """单条记录的进展概览（详情出口用），字段口径与列表聚合一致。"""
    return progress_overview_map([record.id]).get(record.id) or {
        'progress_count': 0,
        'latest_progress_at': None,
        'latest_progress_content': '',
    }


def duration_seconds(occurred_at, latest_progress_at):
    """持续时间（秒）= 最新进展时间 - 主记录发生时间，三种记录类型同一口径。

    - 最新进展时间为全部进展处理时间的最大值（progress_overview_map 聚合），
      不使用主记录 updated_at、附件更新时间或当前时间，也不随当前时间递增；
    - 无进展（latest_progress_at 为空）返回 None（前端显示"-"）；
    - 最新进展早于发生时间（时间更正等）返回 None：不取绝对值、不静默归零；
    - 差值为 0 返回 0（前端显示"0分钟"）。
    """
    if not occurred_at or not latest_progress_at:
        return None
    seconds = (latest_progress_at - occurred_at).total_seconds()
    if seconds < 0:
        return None
    return int(seconds)


# "长期未关闭"判定天数：创建超过该天数仍未关闭的跟进中记录
STALE_OPEN_DAYS = 7

# 待办/积压口径：只统计"跟进中"（异常事件唯一需要用户持续处理的状态）。
# 已记录（留痕）同为合法稳定状态；设备故障的处置中由健康状态与故障列表暴露，
# 不混入跟进待办。
STAT_OPEN_STATUSES = [RECORD_STATUS_FOLLOWING_UP]


def compute_operation_statistics(tenant_id, occurred_from, occurred_to,
                                 now=None, with_trend=True, trend_limit_days=62,
                                 record_type=None):
    """运行处置记录只读统计（租户内，时间范围用 __gte/__lt）。

    record_type 统计口径筛选：None（默认，全部类型）/ routine / abnormal /
    fault。返回：
    - totals：总数/跟进中/处置中故障/长期未关闭（跟进中超 7 天）/
      平均处理时长（小时，仅异常事件有 resolved_at 参与）；
    - by_status / by_type：状态与类型维度计数；
    - by_system：按系统计数（数量倒序，最多 50 个系统）；
    - trend：按自然日计数（total/fault），范围超过 trend_limit_days 天时为空。
    """
    now = now or timezone.now()
    query = SystemOperationRecord.objects.filter(
        tenant_id=tenant_id, occurred_at__gte=occurred_from, occurred_at__lt=occurred_to)
    if record_type:
        query = query.filter(record_type=record_type)
    totals = query.aggregate(
        total=Count('id'),
        following_up=Count('id', filter=Q(status=RECORD_STATUS_FOLLOWING_UP)),
        handling_fault=Count(
            'id', filter=Q(record_type=RECORD_TYPE_FAULT, status=RECORD_STATUS_HANDLING)),
    )
    # 长期未关闭：不限统计时间范围（问题积压跨范围仍需暴露），只看未终态 + 创建时长
    stale_boundary = now - timedelta(days=STALE_OPEN_DAYS)
    stale_query = SystemOperationRecord.objects.filter(
        tenant_id=tenant_id,
        status__in=STAT_OPEN_STATUSES,
        created_at__lt=stale_boundary,
    )
    if record_type:
        stale_query = stale_query.filter(record_type=record_type)
    totals['stale_open'] = stale_query.count()

    by_status = dict(
        query.values_list('status').annotate(total=Count('id')).order_by())
    by_type = dict(
        query.values_list('record_type').annotate(total=Count('id')).order_by())

    by_system = [
        {
            'system': row['system_id'],
            'system_name': row['system_name'],
            'total': row['total'],
            'following_up': row['following_up'],
            'handling_fault': row['handling_fault'],
        }
        for row in query.values(
            'system_id', 'system_name',
        ).annotate(
            total=Count('id'),
            following_up=Count('id', filter=Q(status=RECORD_STATUS_FOLLOWING_UP)),
            handling_fault=Count(
                'id', filter=Q(record_type=RECORD_TYPE_FAULT, status=RECORD_STATUS_HANDLING)),
        ).order_by('-total')[:50]
    ]

    # 平均处理时长在 Python 侧计算（DurationField 聚合在 MariaDB 上兼容性差）
    resolve_pairs = query.filter(
        resolved_at__isnull=False).values_list('resolved_at', 'created_at')
    durations = [
        (resolved - created).total_seconds()
        for resolved, created in resolve_pairs if resolved and created
    ]
    totals['avg_resolve_hours'] = (
        round(sum(durations) / len(durations) / 3600.0, 1) if durations else None)
    totals['resolved_count'] = len(durations)

    trend = []
    if with_trend:
        range_days = (occurred_to - occurred_from).days
        if 0 <= range_days <= trend_limit_days:
            buckets = {}
            for occurred_at, rec_type in query.values_list('occurred_at', 'record_type'):
                day = occurred_at.date().isoformat()
                bucket = buckets.setdefault(day, {'total': 0, 'fault': 0})
                bucket['total'] += 1
                if rec_type == RECORD_TYPE_FAULT:
                    bucket['fault'] += 1
            day = occurred_from.date()
            while day < occurred_to.date():
                key = day.isoformat()
                trend.append({
                    'date': key,
                    'total': buckets.get(key, {}).get('total', 0),
                    'fault': buckets.get(key, {}).get('fault', 0),
                })
                day += timedelta(days=1)
    return {
        'totals': totals,
        'by_status': by_status,
        'by_type': by_type,
        'by_system': by_system,
        'trend': trend,
    }


# ==================== 统一记录导出（xlsx） ====================
#
# 单工作表排版：记录基础列（前 EXPORT_RECORD_COLUMN_COUNT 列）纵向合并为一格，
# 进展列（处理时间/处理人/处理内容）一条进展一行展平在其右侧，无进展记录占
# 一行（处理内容显示"—"）；相邻记录组交替底色、组间中粗分隔线、表头冻结。
# 合并单元格 Excel 不做自动行高，行高按文本显示宽度预估后显式固定。
# 工时（work_hours）已退役，不再导出。

RECORD_EXPORT_CONTENT_TYPE = (
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
RECORD_EXPORT_SHEET_NAME = '运行处置记录'

# 列规格：(表头, 列宽, 是否自动换行)；前 EXPORT_RECORD_COLUMN_COUNT 列随记录
# 合并，其余为进展列（逐行）。
EXPORT_RECORD_COLUMN_COUNT = 15
RECORD_EXPORT_COLUMN_SPECS = [
    ('记录编号', 16, False), ('记录类型', 10, False), ('所属系统', 14, False),
    ('系统编号', 12, False), ('状态', 9, False), ('影响系统运行', 12, False),
    ('系统影响程度', 12, False), ('发生时间', 17, False), ('标题', 22, True),
    ('情况描述', 40, True), ('登记人', 10, False), ('恢复时间', 17, False),
    ('结案说明', 24, True), ('恢复确认人', 12, False), ('来源类型', 11, False),
    ('处理时间', 17, False), ('处理人', 10, False), ('处理内容', 40, True),
]
# 记录基础列中需要自动换行的列（0 基下标：标题/情况描述/结案说明）
RECORD_EXPORT_WRAP_COLUMNS = (8, 9, 12)
# 记录基础列中的时间列（0 基下标：发生时间/恢复时间）
RECORD_EXPORT_DATETIME_COLUMNS = (7, 11)
RECORD_EXPORT_DATETIME_FORMAT = 'YYYY-MM-DD HH:MM'
# 每行文本行数折算行高（pt）与最小行高
RECORD_EXPORT_LINE_HEIGHT = 14
RECORD_EXPORT_MIN_ROW_HEIGHT = 20


def _export_text_lines(text, col_width):
    """按显示宽度估算自动换行行数（全角字符按 2 计）。"""
    if not text:
        return 1
    display_width = sum(2 if ord(ch) > 127 else 1 for ch in str(text))
    return max(1, math.ceil(display_width / max(col_width - 2, 1)))


def _export_base_values(item):
    """记录基础列取值（时间列写 datetime 对象，显示由数字格式控制）。"""
    return [
        item.record_no,
        RECORD_TYPE_TEXT_MAP.get(item.record_type, item.record_type),
        item.system_name,
        item.system_code,
        RECORD_STATUS_TEXT_MAP.get(item.status, item.status),
        '是' if item.affects_system_operation else '否',
        FAULT_IMPACT_TEXT_MAP.get(item.impact_status, ''),
        item.occurred_at,
        item.title,
        item.description,
        item.registrar_name or '',
        item.recovered_at,
        item.recovery_close_note,
        item.recovered_by_name or '',
        item.source_type,
    ]


def _export_progress_values(progress):
    """进展列取值：处理时间 / 处理人（手输名，空回落操作账号）/ 处理内容。"""
    return [
        progress.handled_at,
        progress.handler_name or (
            progress.handler.username if progress.handler_id else ''),
        progress.content,
    ]


def build_record_progress_map(records):
    """单次查询批量取进展并按记录分桶（桶内按处理时间正序，即时间线顺序）。"""
    record_ids = [item.id for item in records]
    if not record_ids:
        return {}
    rows = SystemOperationProgress.objects.filter(
        record_id__in=record_ids).order_by('record_id', 'handled_at', 'id')
    progress_map = {}
    for progress in rows:
        progress_map.setdefault(progress.record_id, []).append(progress)
    return progress_map


def _export_write_header(ws):
    """表头行：加粗底色、列宽预设、冻结首行。"""
    header_fill = PatternFill('solid', fgColor='D9E2F3')
    for col, (name, width, _) in enumerate(RECORD_EXPORT_COLUMN_SPECS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = 'A2'


def _export_style_group(ws, first_row, last_row, group_idx):
    """记录组样式：基础列纵向合并、交替底色、组框线（合并区边框需逐格设置）。"""
    total_cols = len(RECORD_EXPORT_COLUMN_SPECS)
    fill = PatternFill('solid', fgColor='F2F6FA') if group_idx % 2 else None
    thin = Side(style='thin', color='BFBFBF')
    medium = Side(style='medium', color='595959')
    for row in range(first_row, last_row + 1):
        for col in range(1, total_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = Border(
                left=medium if col == 1 else thin,
                right=medium if col == total_cols else thin,
                top=medium if row == first_row else thin,
                bottom=medium if row == last_row else thin)
            cell.alignment = Alignment(
                horizontal='left', vertical='center',
                wrap_text=RECORD_EXPORT_COLUMN_SPECS[col - 1][2])
            if fill is not None:
                cell.fill = fill
    if last_row > first_row:
        for col in range(1, EXPORT_RECORD_COLUMN_COUNT + 1):
            ws.merge_cells(
                start_row=first_row, start_column=col,
                end_row=last_row, end_column=col)


def _export_set_row_heights(ws, first_row, base_values, progress_values_list):
    """按文本预估行高并固定（合并跨行文本所需高度平摊到组内各行）。"""
    row_count = max(len(progress_values_list), 1)
    merged_lines = 1
    for idx in RECORD_EXPORT_WRAP_COLUMNS:
        width = RECORD_EXPORT_COLUMN_SPECS[idx][1]
        merged_lines = max(
            merged_lines, _export_text_lines(base_values[idx], width))
    content_width = RECORD_EXPORT_COLUMN_SPECS[-1][1]
    progress_height = RECORD_EXPORT_MIN_ROW_HEIGHT
    for values in progress_values_list:
        lines = _export_text_lines(values[-1], content_width)
        progress_height = max(
            progress_height, lines * RECORD_EXPORT_LINE_HEIGHT + 6)
    merged_height = math.ceil(
        (merged_lines * RECORD_EXPORT_LINE_HEIGHT + 6) / row_count)
    height = max(RECORD_EXPORT_MIN_ROW_HEIGHT, progress_height, merged_height)
    for offset in range(row_count):
        ws.row_dimensions[first_row + offset].height = height


def build_record_export_workbook(records):
    """构建统一记录导出 xlsx（单工作表展平排版），返回 (文件字节, 进展行数)。"""
    progress_map = build_record_progress_map(records)
    wb = Workbook()
    ws = wb.active
    ws.title = RECORD_EXPORT_SHEET_NAME
    _export_write_header(ws)
    progress_rows = 0
    row = 2
    for group_idx, item in enumerate(records):
        values_list = [
            _export_progress_values(progress)
            for progress in progress_map.get(item.id, [])]
        if not values_list:
            values_list = [['', '', '—']]
        base_values = _export_base_values(item)
        first_row = row
        for values in values_list:
            for offset, value in enumerate(values):
                cell = ws.cell(
                    row=row, column=EXPORT_RECORD_COLUMN_COUNT + 1 + offset,
                    value=value)
                if offset == 0 and value:
                    cell.number_format = RECORD_EXPORT_DATETIME_FORMAT
            row += 1
        for col, value in enumerate(base_values, start=1):
            cell = ws.cell(row=first_row, column=col, value=value)
            if (col - 1) in RECORD_EXPORT_DATETIME_COLUMNS and value:
                cell.number_format = RECORD_EXPORT_DATETIME_FORMAT
        _export_style_group(ws, first_row, row - 1, group_idx)
        _export_set_row_heights(ws, first_row, base_values, values_list)
        progress_rows += len(values_list)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), progress_rows
