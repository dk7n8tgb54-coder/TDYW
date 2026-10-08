# Copyright: (c) OpenSpug Organization. https://github.com/openspug
# Released under the AGPL-3.0 License.
"""值班巡检单导出模块（列表 Excel / 详情 Excel / 详情 PDF / 批量详情导出）。

导出口径（与巡检单详情页一致）：
- 四个固定轮次（11:00-11:20 / 16:00-16:20 / 22:00-22:20 / 次日 6:30-6:50），第四轮归属前一值班日；
- 明细以巡检单已保存数据为准（全量快照），不根据当前时间推断轮次状态；
- 各轮次完成情况读取 recorded_rounds 最终状态；
- 不导出明细实际填写时间（created_at/updated_at 仅为系统审计字段，业务上
  不记录填写时间），也不存在迟填/逾期标记。

实现边界：
- 列表 Excel 复用 libs.export_utils.build_excel_response（样式/文件名/响应头）；
- 详情 Excel 为四轮矩阵双工作表，业务专用构建，但沿用公共样式语言与
  RFC 5987 中文文件名约定；
- PDF 为巡检单专用模块（横向四轮矩阵版式与其他业务不同），复用
  libs.font_manager 的中文字体注册，不做与业务无关的万能 PDF 生成器；
- 批量详情导出（Excel 汇总 + 每单双工作表 / PDF 合并文档流）复用上述
  单条详情的表格与版式构建函数，PDF 由同一 ReportLab 文档流连续生成
  （每单从新页开始），不做"多文件临时拼接"。
"""
import io
import re
from datetime import datetime
from urllib.parse import quote
from xml.sax.saxutils import escape

from django.conf import settings
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    Flowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from libs.font_manager import FONT_NAME, register_chinese_font

from apps.equipment_lifecycle.models import (
    INSPECTION_RESULT_ABNORMAL,
    INSPECTION_ROUND_TEXT_MAP,
    INSPECTION_ROUND_VALUES,
)

# 未记录轮次与异常结果的视觉区分色（与公共 Excel 样式语言一致）
_COLOR_ABNORMAL = 'FF' + 'FF0000'
_COLOR_UNRECORDED = 'FF' + '808080'

# Excel 头部样式常量（与 libs/export_utils.build_excel_response 保持一致）
_HEADER_FILL_HEX = '4472C4'
_BORDER_HEX = 'BFBFBF'

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
PDF_CONTENT_TYPE = 'application/pdf'

LIST_SHEET_NAME = '巡检单记录'

# 列表导出列定义：(字段, 表头)。轮次完成情况四列由 recorded_rounds 派生。
LIST_COLUMNS = [
    ('duty_date', '值班日'),
    ('sheet_no', '单号'),
    ('template_name', '巡检模板'),
    ('persons_text', '值班人员'),
    ('status_text', '状态'),
    ('round_r1', '11:00-11:20'),
    ('round_r2', '16:00-16:20'),
    ('round_r3', '22:00-22:20'),
    ('round_r4', '6:30-6:50'),
    ('abnormal_count', '异常数量'),
    ('submitted_at', '提交时间'),
    ('submitted_by_name', '提交人'),
    ('created_by_name', '创建人'),
]

# 详情 Excel 工作表名
INFO_SHEET_NAME = '巡查单信息'
MATRIX_SHEET_NAME = '四轮巡检明细'

# 批量详情 Excel：汇总工作表名 + 每单两个详情工作表名前缀（"前缀_单号"）
BATCH_SUMMARY_SHEET_NAME = '巡检单汇总'
BATCH_INFO_TITLE_PREFIX = '信息_'
BATCH_MATRIX_TITLE_PREFIX = '明细_'

# 批量详情导出数量上限（Excel/PDF 分别校验，可经 Django settings 覆盖）
BATCH_EXPORT_LIMIT_MESSAGES = {
    'xlsx': '批量导出数据超过 {limit} 条，请缩小值班日期范围或增加筛选条件后重试',
    'pdf': '批量 PDF 导出数据超过 {limit} 条，请缩小值班日期范围或增加筛选条件后重试',
}
_DEFAULT_BATCH_LIMITS = {'xlsx': 500, 'pdf': 100}
_LIMIT_SETTING_NAMES = {
    'xlsx': 'INSPECTION_BATCH_EXPORT_XLSX_LIMIT',
    'pdf': 'INSPECTION_BATCH_EXPORT_PDF_LIMIT',
}

_MATRIX_HEADERS = ['检查项目', '检查内容'] + [
    INSPECTION_ROUND_TEXT_MAP[r] for r in INSPECTION_ROUND_VALUES]


def export_filename(base, ext):
    """导出文件名：{业务名}_{时间戳}.{扩展名}。"""
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    return f'{base}_{stamp}.{ext}'


def content_disposition(filename):
    """中文文件名响应头（与 libs.export_utils 同一 RFC 5987 约定）。"""
    return "attachment; filename*=UTF-8''%s" % quote(filename)


def batch_export_limit(export_format):
    """批量详情导出数量上限（Excel 默认 500 / PDF 默认 100，settings 可覆盖）。"""
    default = _DEFAULT_BATCH_LIMITS.get(export_format, _DEFAULT_BATCH_LIMITS['xlsx'])
    return getattr(settings, _LIMIT_SETTING_NAMES.get(export_format,
                                                      _LIMIT_SETTING_NAMES['xlsx']), default)


def batch_limit_message(export_format, limit):
    """批量导出超限错误消息（按导出格式区分文案）。"""
    template = BATCH_EXPORT_LIMIT_MESSAGES.get(
        export_format, BATCH_EXPORT_LIMIT_MESSAGES['xlsx'])
    return template.format(limit=limit)


def check_batch_export_limit(count, export_format):
    """批量导出数量上限校验：超限返回错误消息，未超限返回 None。"""
    limit = batch_export_limit(export_format)
    if count > limit:
        return batch_limit_message(export_format, limit)
    return None


# ----------------------------------------------------------------------
# 列表 Excel
# ----------------------------------------------------------------------

def build_list_rows(sheets):
    """把巡检单查询集构建为列表导出行（轮次完成情况取 recorded_rounds）。

    调用方需 prefetch_related('persons', 'records')，避免逐单查库。
    """
    rows = []
    for sheet in sheets:
        records = list(sheet.records.all())
        abnormal_count = sum(
            1 for r in records if r.result == INSPECTION_RESULT_ABNORMAL)
        recorded = set(sheet.get_recorded_rounds())
        row = {
            'duty_date': str(sheet.duty_date),
            'sheet_no': sheet.sheet_no,
            'template_name': sheet.template.name,
            'persons_text': '、'.join(
                p.person_name for p in sheet.persons.all()) or '（未登记）',
            'status_text': sheet.get_status_display(),
            'abnormal_count': abnormal_count,
            'submitted_at': sheet.submitted_at,
            'submitted_by_name': sheet.submitted_by.username if sheet.submitted_by_id else '',
            'created_by_name': sheet.created_by.username if sheet.created_by_id else '',
        }
        for round_no in INSPECTION_ROUND_VALUES:
            row[f'round_{round_no}'] = '已记录' if round_no in recorded else '未记录'
        rows.append(row)
    return rows


# ----------------------------------------------------------------------
# 详情 Excel（业务专用：巡查单信息 + 四轮巡检明细双工作表）
# ----------------------------------------------------------------------

def _group_detail_rows(records):
    """一行 = 检查项目（设备名称快照）× 检查内容，四轮取值横向展开。

    与详情页 SheetDetailDrawer 的分组规则一致：模板行按 template_item、
    临时加项按 client_key（无则记录 ID）分组；临时加项设备名归"临时检查项"。
    """
    item_map = {}
    for record in records:
        if record.get('template_item'):
            key = f"t_{record['template_item']}"
        else:
            key = f"m_{record.get('client_key') or record.get('id')}"
        item = item_map.get(key)
        if item is None:
            item = item_map[key] = {
                'key': key,
                'project': record.get('device_name') or '临时检查项',
                'content': record.get('item_name'),
                'display_order': record.get('display_order') or 0,
                'rounds': {},
            }
        item['rounds'][record.get('round_no')] = record
    return sorted(
        item_map.values(),
        key=lambda i: (i['project'], i['display_order'], i['key']))


def round_cell_lines(record):
    """单轮单元格文本行：填写值 + 异常说明 + 现场处理 + 转单去向。

    返回 (lines, state)；state ∈ normal / abnormal / none（未记录）。
    不包含实际填写时间，也没有迟填/逾期概念。
    """
    if record is None:
        return ['未记录'], 'none'
    lines = [record.get('value_text') or '—']
    state = 'abnormal' if record.get('result') == INSPECTION_RESULT_ABNORMAL else 'normal'
    if record.get('abnormal_desc'):
        lines.append(f"异常说明：{record['abnormal_desc']}")
    if record.get('onsite_action'):
        lines.append(f"现场处理：{record['onsite_action']}")
    disposition = record.get('disposition') or {}
    if disposition.get('status') == 'converted_fault' and disposition.get('fault'):
        lines.append(f"已转故障：{disposition['fault'].get('fault_no')}")
    elif disposition.get('status') == 'target_cancelled':
        lines.append('转单目标已作废')
    return lines, state


def _info_rows(detail):
    """巡查单信息工作表的 (项目, 内容) 行（含作废留痕）。"""
    rounds_text = '、'.join(
        INSPECTION_ROUND_TEXT_MAP[r] for r in detail.get('recorded_rounds') or [])
    persons_text = '、'.join(
        p.get('person_name') for p in detail.get('persons') or []) or '（未登记）'
    rows = [
        ('单号', detail.get('sheet_no')),
        ('值班日', detail.get('duty_date')),
        ('巡检模板', detail.get('template_name')),
        ('状态', detail.get('status_text')),
        ('值班人员', persons_text),
        ('异常数量', detail.get('abnormal_count')),
        ('已记录轮次', rounds_text or '无'),
        ('提交时间', detail.get('submitted_at') or '-'),
        ('提交人', detail.get('submitted_by_name') or '-'),
        ('创建人', detail.get('created_by_name') or '-'),
    ]
    if detail.get('status') == 'void':
        rows.append(('作废原因', detail.get('void_reason') or '-'))
        voided = detail.get('voided_at')
        rows.append(('作废信息', f"{voided} {detail.get('voided_by_name') or ''}".strip() or '-'))
    return rows


def _styled_header_cell(cell, text):
    cell.value = text
    cell.font = Font(name='宋体', bold=True, color='FFFFFF', size=11)
    cell.fill = PatternFill(start_color=_HEADER_FILL_HEX, end_color=_HEADER_FILL_HEX,
                            fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = _thin_border()
    return cell


def _thin_border():
    thin = Side(border_style='thin', color=_BORDER_HEX)
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _write_info_sheet(ws, detail):
    """填写巡查单信息工作表（项目/内容两列，含作废留痕）。"""
    _styled_header_cell(ws.cell(row=1, column=1), '项目')
    _styled_header_cell(ws.cell(row=1, column=2), '内容')
    body_align = Alignment(vertical='center', wrap_text=True)
    for row_idx, (label, value) in enumerate(_info_rows(detail), start=2):
        label_cell = ws.cell(row=row_idx, column=1, value=label)
        label_cell.font = Font(name='宋体', size=10, bold=True)
        label_cell.border = _thin_border()
        value_cell = ws.cell(row=row_idx, column=2,
                             value=value.strftime('%Y-%m-%d %H:%M:%S')
                             if hasattr(value, 'strftime') else value)
        value_cell.font = Font(name='宋体', size=10)
        value_cell.alignment = body_align
        value_cell.border = _thin_border()
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 50
    ws.freeze_panes = 'A2'


def _write_matrix_sheet(ws, detail):
    """填写四轮巡检明细工作表（一行 = 检查项目 × 检查内容，四轮取值横向展开）。"""
    for col_idx, title in enumerate(_MATRIX_HEADERS, start=1):
        _styled_header_cell(ws.cell(row=1, column=col_idx), title)
    body_font = Font(name='宋体', size=10)
    unrecorded_font = Font(name='宋体', size=10, color=_COLOR_UNRECORDED)
    abnormal_font = Font(name='宋体', size=10, bold=True, color=_COLOR_ABNORMAL)
    top_align = Alignment(vertical='top', wrap_text=True)
    for row_idx, item in enumerate(_group_detail_rows(detail.get('records') or []), start=2):
        ws.cell(row=row_idx, column=1, value=item['project'])
        ws.cell(row=row_idx, column=2, value=item['content'])
        for col_idx, round_no in enumerate(INSPECTION_ROUND_VALUES, start=3):
            lines, state = round_cell_lines(item['rounds'].get(round_no))
            cell = ws.cell(row=row_idx, column=col_idx, value='\n'.join(lines))
            if state == 'abnormal':
                cell.font = abnormal_font
            elif state == 'none':
                cell.font = unrecorded_font
            else:
                cell.font = body_font
            cell.alignment = top_align
            cell.border = _thin_border()
        for col_idx in (1, 2):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = body_font
            cell.alignment = top_align
            cell.border = _thin_border()
    widths = [16, 26, 30, 30, 30, 30]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    # 冻结首行 + 左侧两列（检查项目/检查内容）
    ws.freeze_panes = 'C2'


def build_detail_workbook(detail):
    """构建巡检单详情工作簿：巡查单信息 + 四轮巡检明细。"""
    wb = Workbook()
    ws = wb.active
    ws.title = INFO_SHEET_NAME
    _write_info_sheet(ws, detail)
    _write_matrix_sheet(wb.create_sheet(MATRIX_SHEET_NAME), detail)
    return wb


def build_detail_response(detail, export_format='xlsx'):
    """构建巡检单详情导出响应：xlsx（双工作表矩阵）或 pdf（A4 纵向）。"""
    if export_format == 'pdf':
        response = HttpResponse(generate_sheet_pdf(detail), content_type=PDF_CONTENT_TYPE)
    else:
        output = io.BytesIO()
        build_detail_workbook(detail).save(output)
        response = HttpResponse(output.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response['Content-Disposition'] = content_disposition(
        export_filename(f"巡检单_{detail.get('sheet_no')}", export_format))
    return response


# ----------------------------------------------------------------------
# 批量详情 Excel（一个工作簿：巡检单汇总 + 每单两个详情工作表）
# ----------------------------------------------------------------------

# Excel 工作表名限制：≤31 字符，不含 \ / ? * [ ] :
_INVALID_TITLE_CHARS = re.compile(r'[\\/?*\[\]:]')
# 预留"_序号"后缀空间（27 + '_' + 最多 3 位序号 ≤ 31 字符）
_MAX_TITLE_BASE_LEN = 27

# 汇总表列宽（与 LIST_COLUMNS 一一对应）
_SUMMARY_COLUMN_WIDTHS = [12, 20, 18, 14, 10, 12, 12, 12, 12, 10, 20, 12, 12]


def _summary_cell_value(value):
    """汇总表单元格取值：datetime 转文本，None 转空串（与列表导出口径一致）。"""
    if value is None:
        return ''
    if hasattr(value, 'strftime'):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    return value


def _write_summary_sheet(ws, rows):
    """巡检单汇总工作表：复用列表导出列定义与样式语言，每单一行。"""
    for col_idx, (_key, title) in enumerate(LIST_COLUMNS, start=1):
        _styled_header_cell(ws.cell(row=1, column=col_idx), title)
    body_font = Font(name='宋体', size=10)
    body_align = Alignment(vertical='center', wrap_text=True)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (key, _title) in enumerate(LIST_COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx,
                           value=_summary_cell_value(row.get(key)))
            cell.font = body_font
            cell.alignment = body_align
            cell.border = _thin_border()
    for col_idx, width in enumerate(_SUMMARY_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = 'A2'


def _unique_sheet_title(wb, base):
    """合法且互不重复的工作表名：剔除 Excel 非法字符、限长、冲突追加序号。"""
    base = _INVALID_TITLE_CHARS.sub('_', base)[:_MAX_TITLE_BASE_LEN]
    candidate, seq = base, 2
    while candidate in wb.sheetnames:
        candidate = f'{base}_{seq}'
        seq += 1
    return candidate


def build_batch_detail_workbook(sheets, details):
    """构建批量详情工作簿：巡检单汇总 + 每张巡检单两个详情工作表。

    sheets：已 prefetch 的巡检单序列（供复用 build_list_rows 汇总口径）；
    details：与 sheets 一一对应的 to_detail_view() 载荷（含处置状态）。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = BATCH_SUMMARY_SHEET_NAME
    _write_summary_sheet(ws, build_list_rows(sheets))
    for detail in details:
        sheet_no = str(detail.get('sheet_no') or '')
        info_title = _unique_sheet_title(wb, BATCH_INFO_TITLE_PREFIX + sheet_no)
        matrix_title = _unique_sheet_title(wb, BATCH_MATRIX_TITLE_PREFIX + sheet_no)
        _write_info_sheet(wb.create_sheet(info_title), detail)
        _write_matrix_sheet(wb.create_sheet(matrix_title), detail)
    return wb


def build_batch_detail_response(sheets, details, export_format='xlsx'):
    """构建批量详情导出响应：xlsx（汇总 + 每单双工作表）或 pdf（合并文档流）。"""
    if export_format == 'pdf':
        response = HttpResponse(
            generate_batch_pdf(details), content_type=PDF_CONTENT_TYPE)
    else:
        output = io.BytesIO()
        build_batch_detail_workbook(sheets, details).save(output)
        response = HttpResponse(output.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response['Content-Disposition'] = content_disposition(
        export_filename('巡检单详情批量', export_format))
    return response


# ----------------------------------------------------------------------
# 详情 PDF（巡检单专用版式：横向四轮矩阵，A4 纵向打印归档）
# ----------------------------------------------------------------------

# PDF 主题色（蓝系表头 + 红/灰状态区分，与 Excel 视觉语言一致）
THEME_BORDER = colors.HexColor('#BFBFBF')
THEME_HEADER_BLUE = colors.HexColor('#4472C4')
THEME_BG_HEADER = colors.HexColor('#EAF0FA')
THEME_TEXT = colors.HexColor('#333333')
THEME_TEXT_SECONDARY = colors.HexColor('#666666')
THEME_TEXT_ABNORMAL = colors.HexColor('#FF0000')
THEME_TEXT_NONE = colors.HexColor('#808080')


def _title_style():
    return ParagraphStyle(
        'InsTitle', fontName=FONT_NAME, fontSize=16, alignment=1,
        spaceAfter=4, textColor=colors.HexColor('#096DD9'), leading=22)


def _subtitle_style():
    return ParagraphStyle(
        'InsSubtitle', fontName=FONT_NAME, fontSize=10, alignment=1,
        spaceAfter=10, textColor=THEME_TEXT_SECONDARY, leading=14)


def _section_style():
    return ParagraphStyle(
        'InsSection', fontName=FONT_NAME, fontSize=12, alignment=0,
        spaceAfter=6, spaceBefore=2, textColor=colors.HexColor('#096DD9'), leading=16)


def _cell_style():
    return ParagraphStyle(
        'InsCell', fontName=FONT_NAME, fontSize=8.5, alignment=0,
        leading=12, wordWrap='CJK', textColor=THEME_TEXT)


def _label_style():
    return ParagraphStyle(
        'InsLabel', fontName=FONT_NAME, fontSize=8.5, alignment=2,
        leading=12, wordWrap='CJK', textColor=THEME_TEXT_SECONDARY)


def _pdf_header_style():
    return ParagraphStyle(
        'InsHeader', fontName=FONT_NAME, fontSize=9, alignment=1,
        leading=12, wordWrap='CJK', textColor=colors.white)


def _abnormal_cell_style():
    return ParagraphStyle(
        'InsAbnormal', fontName=FONT_NAME, fontSize=8.5, alignment=0,
        leading=12, wordWrap='CJK', textColor=THEME_TEXT_ABNORMAL)


def _none_cell_style():
    return ParagraphStyle(
        'InsNone', fontName=FONT_NAME, fontSize=8.5, alignment=0,
        leading=12, wordWrap='CJK', textColor=THEME_TEXT_NONE)


def _label(text, style):
    return Paragraph(escape(str(text)), style)


def _value(text, style):
    return Paragraph(escape('' if text is None else str(text)), style)


def _pdf_footer(detail):
    """页面底部：巡检单编号 + 页码 + 导出时间（供打印归档溯源）。"""
    sheet_no = detail.get('sheet_no')
    exported_at = datetime.now().strftime('%Y-%m-%d %H:%M')

    def _draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(FONT_NAME, 8)
        canvas.setFillColor(THEME_TEXT_SECONDARY)
        canvas.drawString(doc.leftMargin, 1.0 * cm, f'巡检单编号：{sheet_no}')
        canvas.drawCentredString(A4[0] / 2, 1.0 * cm, f'第 {canvas.getPageNumber()} 页')
        canvas.drawRightString(A4[0] - doc.rightMargin, 1.0 * cm, f'导出时间：{exported_at}')
        canvas.restoreState()

    return _draw


def _multiline_para(lines, style):
    """多行文本 Paragraph（逐行 XML 转义后以 <br/> 连接）。"""
    return Paragraph('<br/>'.join(escape(str(line)) for line in lines), style)


def _build_pdf_info_table(detail, cell_style, label_style):
    """基本信息表：两列键值一组，值班人员等长值跨列。"""
    rounds_text = '、'.join(
        INSPECTION_ROUND_TEXT_MAP[r] for r in detail.get('recorded_rounds') or []) or '无'
    persons_text = '、'.join(
        p.get('person_name') for p in detail.get('persons') or []) or '（未登记）'
    rows = [
        [_label('单号', label_style), _value(detail.get('sheet_no'), cell_style),
         _label('状态', label_style), _value(detail.get('status_text'), cell_style)],
        [_label('值班日', label_style), _value(detail.get('duty_date'), cell_style),
         _label('巡检模板', label_style), _value(detail.get('template_name'), cell_style)],
        [_label('值班人员', label_style), _value(persons_text, cell_style), '', ''],
        [_label('异常数量', label_style), _value(str(detail.get('abnormal_count')), cell_style),
         _label('已记录轮次', label_style), _value(rounds_text, cell_style)],
        [_label('提交时间', label_style),
         _value(f"{detail.get('submitted_at') or '-'}"
                f"{'（%s）' % detail['submitted_by_name'] if detail.get('submitted_by_name') else ''}",
                cell_style),
         _label('创建人', label_style), _value(detail.get('created_by_name') or '-', cell_style)],
    ]
    if detail.get('status') == 'void':
        rows.append([_label('作废原因', label_style),
                     _value(f"{detail.get('void_reason') or '-'}"
                            f"{'（%s %s）' % (detail.get('voided_at') or '', detail.get('voided_by_name') or '')}".strip(),
                            cell_style), '', ''])
    table = Table(rows, colWidths=[2.6 * cm, 6.4 * cm, 2.6 * cm, 6.4 * cm])
    table.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, THEME_BORDER),
        ('BACKGROUND', (0, 0), (0, -1), THEME_BG_HEADER),
        ('BACKGROUND', (2, 0), (2, -1), THEME_BG_HEADER),
        ('SPAN', (1, 2), (3, 2)),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ] + ([('SPAN', (1, len(rows) - 1), (3, len(rows) - 1))] if detail.get('status') == 'void' else [])))
    return table


def _build_pdf_matrix_table(detail, header_style, normal_style,
                            abnormal_style, none_style):
    """四轮巡检明细矩阵表：一行 = 检查项目 × 检查内容，四轮取值横向展开。"""
    header = [Paragraph(escape(t), header_style) for t in _MATRIX_HEADERS]
    data = [header]
    for item in _group_detail_rows(detail.get('records') or []):
        row = [Paragraph(escape(item['project']), normal_style),
               Paragraph(escape(item['content']), normal_style)]
        for round_no in INSPECTION_ROUND_VALUES:
            lines, state = round_cell_lines(item['rounds'].get(round_no))
            style = abnormal_style if state == 'abnormal' \
                else none_style if state == 'none' else normal_style
            row.append(_multiline_para(lines, style))
        data.append(row)
    table = Table(data, colWidths=[3.0 * cm, 4.0 * cm, 2.75 * cm, 2.75 * cm, 2.75 * cm, 2.75 * cm],
                  repeatRows=1)
    table.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, THEME_BORDER),
        ('BACKGROUND', (0, 0), (-1, 0), THEME_HEADER_BLUE),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    return table


def build_detail_story(detail):
    """单个巡检单的 PDF 内容块（标题 + 副标题 + 信息表 + 四轮矩阵）。

    单条导出与批量合并导出共用同一版式；每次调用返回全新流对象，
    不可跨文档复用。
    """
    return [
        Paragraph('值班巡检单', _title_style()),
        Paragraph(
            escape(f"单号 {detail.get('sheet_no')}　值班日 {detail.get('duty_date')}"),
            _subtitle_style()),
        _build_pdf_info_table(detail, _cell_style(), _label_style()),
        Spacer(1, 10),
        Paragraph('四轮巡检明细', _section_style()),
        _build_pdf_matrix_table(
            detail, _pdf_header_style(), _cell_style(),
            _abnormal_cell_style(), _none_cell_style()),
    ]


def generate_sheet_pdf(detail):
    """生成巡检单详情 PDF（A4 纵向、中文、表头重复、四轮矩阵），返回字节。"""
    register_chinese_font()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.6 * cm, bottomMargin=1.8 * cm,
        title=f"值班巡检单 {detail.get('sheet_no')}", author='设备运行维护管理')
    doc.build(build_detail_story(detail),
              onFirstPage=_pdf_footer(detail), onLaterPages=_pdf_footer(detail))
    return buf.getvalue()


class _SheetFooterMarker(Flowable):
    """零尺寸标记流：绘制时登记当前页所属巡检单编号（批量 PDF 页脚归属用）。"""

    def __init__(self, holder, sheet_no):
        Flowable.__init__(self)
        self._holder = holder
        self._sheet_no = sheet_no
        self.width = 0
        self.height = 0

    def draw(self):
        self._holder['sheet_no'] = self._sheet_no


class _BatchPdfDocTemplate(SimpleDocTemplate):
    """批量合并 PDF 文档模板：afterPage 时绘制页脚。

    onPage 回调在页面开始时触发，读不到本页的巡检单标记；afterPage 在
    本页全部内容绘制完、翻页前触发，此时本页首个巡检单标记已绘制，
    页脚能正确归属当前巡检单（ReportLab 文档流不区分页脚回调对象）。
    """

    footer_drawer = None

    def afterPage(self):
        if self.footer_drawer:
            self.footer_drawer(self.canv, self.canv.getPageNumber())


def generate_batch_pdf(details):
    """批量合并 PDF：一个 ReportLab 文档流连续生成，每个巡检单从新页开始。

    不做多文件临时拼接；单张巡检单内容较长时允许跨页（矩阵表头重复），
    下一张巡检单通过 PageBreak 从新页开始。页脚含巡检单编号、页码与导出时间。
    """
    register_chinese_font()
    holder = {}
    buf = io.BytesIO()
    doc = _BatchPdfDocTemplate(
        buf, pagesize=A4,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.6 * cm, bottomMargin=1.8 * cm,
        title='值班巡检单批量导出', author='设备运行维护管理')
    left_margin, right_margin = doc.leftMargin, doc.rightMargin
    exported_at = datetime.now().strftime('%Y-%m-%d %H:%M')

    def _draw_footer(canvas, page_number):
        canvas.saveState()
        canvas.setFont(FONT_NAME, 8)
        canvas.setFillColor(THEME_TEXT_SECONDARY)
        canvas.drawString(left_margin, 1.0 * cm,
                          f'巡检单编号：{holder.get("sheet_no") or "-"}')
        canvas.drawCentredString(A4[0] / 2, 1.0 * cm, f'第 {page_number} 页')
        canvas.drawRightString(A4[0] - right_margin, 1.0 * cm,
                               f'导出时间：{exported_at}')
        canvas.restoreState()

    doc.footer_drawer = _draw_footer
    story = []
    for index, detail in enumerate(details):
        if index:
            story.append(PageBreak())
        story.append(_SheetFooterMarker(holder, detail.get('sheet_no')))
        story.extend(build_detail_story(detail))
    doc.build(story)
    return buf.getvalue()
