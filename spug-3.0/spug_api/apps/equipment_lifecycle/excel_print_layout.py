# Copyright: (c) OpenSpug Organization. https://github.com/openspug
# Released under the AGPL-3.0 License.
"""设备全生命周期——Excel 台账统一打印布局与分页计划（2026-10-21）。

网页填写视图与浏览器引擎 PDF（excel_print_html + excel_browser_pdf_service）
共用同一份打印布局规则：本模块从版本快照（布局 meta + pdf_params）计算
权威的「打印计划」，纸张/方向/边距/缩放/分页只在此处定义一次。

- 纸张：真实物理尺寸（A4 = 210×297mm，横向交换宽高），整份 PDF 统一使用
  主（绑定）工作表解析出的纸张与方向（网页与 PDF 一致的前提）；
- 边距：复用模板打印参数的页边距（cm），缺省下边距 1.91cm（≥15mm 建议
  值）；页高计算始终包含下边距，正文（表格/维护人/日期）只落在可打印
  区域内，不把维护人日期改成固定页脚；
- 缩放：actual=100%、percent=指定比例、fit_width=宽度收缩至一页
  （只缩不放）、fit_page=宽高同时收缩至一页；缩放同时作用于行高列宽与
  字号（打印 HTML 用 transform 语义的等比数值，网页用 transform scale）；
- 分页：行方向按缩放后行高贪婪填充（合并区域整块移动，块高于整页时才
  跨页拆分并告警），手动分页行在指定行后断页，重复表头行在后续页顶部
  重排；列方向超出可打印宽度时按列分块（band），page_order 决定遍历顺序；
- 兼容：preset=original 从布局 meta 的 print 设置（上传时解析自源文件，
  历史快照由读取端修正补齐）取纸张/方向/缩放/打印区域；分页参数
  （重复表头/手动分页/显式打印区域）与 LibreOffice 路径一致只作用于
  主工作表；不支持的情形以 warnings 明确提示，绝不静默裁切。

计划同时下发给网页（store.version.print_plan）：网页据此画分页参考线，
不再自行推测分页；打印区域参考线仍属网页辅助，不进入 PDF。
"""
import copy

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from apps.equipment_lifecycle.excel_pdf_params import (
    DEFAULT_PDF_PARAMS,
    PAPER_SIZE_CODES,
    PAPER_SIZE_MM,
    _resolve_meta_orientation,
)

PRINT_PLAN_VERSION = 1
# 96dpi 网页像素换算（与前端 PX_PER_MM 同源）：1mm = 96/25.4 px
PX_PER_MM = 96 / 25.4
# Excel 缺省行高（pt）/列宽（字符，Excel 2007+ 缺省 8.43）与缺省字号
DEFAULT_ROW_HEIGHT_PT = 14.25
DEFAULT_COLUMN_CHARS = 8.43
DEFAULT_FONT_SIZE_PT = 11.0
# 列宽字符单位 → 像素（与网页 LayoutTable 同一换算：chars*7+5）
COL_CHARS_TO_PX_BASE = 7
COL_CHARS_TO_PX_EXTRA = 5
# 行高 pt → 像素（96dpi：pt*4/3）
ROW_PT_TO_PX = 4 / 3
# 浮点比较容差（px），避免边界值抖动
_EPSILON_PX = 0.5
# fit 缩放低于该比例提示字号过小（与 estimate_pdf_warnings 同口径）
_MIN_READABLE_SCALE = 0.5
# 单份 PDF 最大页数（防呆上限，超出明确告警并截断）
MAX_PRINT_PAGES = 500


def column_width_px(chars):
    """Excel 字符列宽 → 像素（网页与打印共用同一换算）。"""
    return ((chars if chars is not None else DEFAULT_COLUMN_CHARS)
            * COL_CHARS_TO_PX_BASE + COL_CHARS_TO_PX_EXTRA)


def row_height_px(pt):
    """Excel 行高 pt → 像素（96dpi）。"""
    return (pt if pt is not None else DEFAULT_ROW_HEIGHT_PT) * ROW_PT_TO_PX


# ==================== 参数归并 ====================

def _effective_params(snapshot_params, override_params):
    """归并出生效打印参数：override 整体优先（与导出覆盖语义一致），
    缺失键回落内置缺省（历史快照字段不全时兜底）。"""
    source = override_params if override_params is not None else snapshot_params
    params = copy.deepcopy(DEFAULT_PDF_PARAMS)
    if isinstance(source, dict):
        for key, value in source.items():
            if key == 'margins':
                continue
            if value is not None and value != '':
                params[key] = value
        raw_margins = source.get('margins')
        if isinstance(raw_margins, dict):
            params['margins'].update({
                k: v for k, v in raw_margins.items() if v is not None})
    return params


def _paper_name(params, main_sheet):
    """纸张名称：preset=original 优先用主工作表 print 设置里的纸张编号。"""
    if params.get('preset') == 'original':
        code = (main_sheet.get('print') or {}).get('paper_size')
        for name, value in PAPER_SIZE_CODES.items():
            if code is not None and int(code) == value:
                return name
    name = params.get('paper_size')
    return name if name in PAPER_SIZE_MM else 'A4'


def _orientation(params, main_sheet, paper_name):
    """整份 PDF 的方向：preset=original 优先用主工作表 print 设置的方向；
    显式配置次之；auto 按主工作表内容宽度解析（纵向放不下转横向）。"""
    if params.get('preset') == 'original':
        from_print = (main_sheet.get('print') or {}).get('orientation')
        if from_print in ('landscape', 'portrait'):
            return from_print
    orientation = params.get('orientation')
    if orientation in ('landscape', 'portrait'):
        return orientation
    return _resolve_meta_orientation(main_sheet, {
        'paper_size': paper_name,
        'orientation': 'auto',
        'margins': params.get('margins') or {},
    })


def _margins_mm(params, main_sheet):
    """页边距 mm：preset=original 且 print 设置携带物理边距时用源文件值。"""
    margins = params.get('margins') or {}
    if params.get('preset') == 'original':
        from_print = (main_sheet.get('print') or {}).get('page_margins') or {}
        margins = {**margins, **{k: v for k, v in from_print.items()
                                 if v is not None}}
    result = {}
    for key in ('top', 'bottom', 'left', 'right'):
        try:
            result[key] = round(float(margins.get(key)) * 10, 2)
        except (TypeError, ValueError):
            result[key] = round(
                float(DEFAULT_PDF_PARAMS['margins'][key]) * 10, 2)
    return result


def _print_range(sheet, area_text, warnings):
    """打印区域文本 → 行列范围（多段取包围盒并提示）；空/非法回落全部内容。"""
    text = str(area_text or '').strip()
    if not text:
        return None
    min_row = min_col = None
    max_row = max_col = None
    for part in text.split(','):
        plain = part.strip().upper().replace('$', '')
        if '!' in plain:
            plain = plain.rsplit('!', 1)[-1]
        try:
            c1, r1, c2, r2 = range_boundaries(plain)
        except Exception:
            warnings.append('打印区域「%s」无法识别，已按全部内容处理' % part.strip())
            return None
        min_row = r1 if min_row is None else min(min_row, r1)
        min_col = c1 if min_col is None else min(min_col, c1)
        max_row = r2 if max_row is None else max(max_row, r2)
        max_col = c2 if max_col is None else max(max_col, c2)
    if min_row is None:
        return None
    return (min_row, max_row, min_col, max_col)


# ==================== 行分页 ====================

def _row_blocks(sheet, row_start, row_end):
    """跨行合并块：行号 → 块末行（块内所有行映射到同一末行）。"""
    blocks = {}
    for merged in sheet.get('merged_ranges') or []:
        try:
            _c1, r1, _c2, r2 = range_boundaries(str(merged))
        except Exception:
            continue
        if r2 <= r1:
            continue
        lo, hi = max(r1, row_start), min(r2, row_end)
        for row in range(lo, hi + 1):
            blocks[row] = hi
    return blocks


def _row_heights_px(sheet, row_start, row_end, scale):
    hidden = set(sheet.get('hidden_rows') or [])
    heights = sheet.get('row_heights') or {}
    result = {}
    for row in range(row_start, row_end + 1):
        if str(row) in hidden:
            result[row] = 0.0
        else:
            result[row] = round(row_height_px(heights.get(str(row))) * scale, 2)
    return result


def _repeat_rows_and_height(heights, repeat_range, row_start, row_end):
    """重复表头行（与打印范围求交）及其总高度。"""
    if not repeat_range:
        return [], 0.0
    a, b = repeat_range
    a, b = max(a, row_start), min(b, row_end)
    if a > b:
        return [], 0.0
    rows = list(range(a, b + 1))
    return rows, round(sum(heights[r] for r in rows), 2)


def _split_take_end(heights, row, block_end, allowed):
    """块高于整页时的跨页拆分：从 row 起在 allowed 高度内取行，返回末行。"""
    take_end = row
    acc = 0.0
    for r in range(row, block_end + 1):
        if acc + heights[r] > allowed + _EPSILON_PX and acc > 0:
            break
        acc += heights[r]
        take_end = r
    return take_end


def _fill_one_page(heights, blocks, break_after, row, row_end,
                   printable_h_px, repeat_h, with_repeat, warnings, name,
                   split_warned):
    """填充一页：贪婪装箱 + 合并块整体移动 + 块高于整页时拆分告警。

    返回 (页首行, 该页末行, 下一页起始行, 拆分告警是否已触发)。
    """
    page_start = row
    page_h = repeat_h if with_repeat else 0.0
    content_rows = 0
    while row <= row_end:
        block_end = blocks.get(row, row)
        block_h = round(sum(heights[r] for r in range(row, block_end + 1)), 2)
        if content_rows and block_h > printable_h_px - page_h + _EPSILON_PX:
            break  # 当前页已满：在块前断页
        if content_rows == 0 \
                and repeat_h + block_h > printable_h_px + _EPSILON_PX:
            # 块（含重排表头）高于整页：拆分跨页并明确告警
            if not split_warned:
                warnings.append(
                    '工作表「%s」第 %d 行起的合并区域高于单页可打印高度，'
                    '将跨页拆分显示（建议调整行高或分页）' % (name, row))
                split_warned = True
            allowed = max(printable_h_px - page_h, heights[row])
            row = _split_take_end(heights, row, block_end, allowed) + 1
            break
        page_h = round(page_h + block_h, 2)
        content_rows += block_end - row + 1
        last_row = block_end
        row = block_end + 1
        if last_row in break_after:
            break
    return page_start, max(page_start, row - 1), row, split_warned


def _paginate_rows(sheet, row_start, row_end, printable_h_px, scale,
                   repeat_range, manual_breaks, warnings):
    """行方向分页：贪婪填充 + 合并块整体移动 + 手动分页 + 重复表头。

    返回 (页列表, 分页行号列表)；页 = {'row_start','row_end','repeat_rows'}
    （repeat_rows 为该页顶部重排的表头行，首页自然包含时不重排）。
    """
    name = sheet.get('name') or ''
    heights = _row_heights_px(sheet, row_start, row_end, scale)
    blocks = _row_blocks(sheet, row_start, row_end)
    break_after = {int(row) for row in (manual_breaks or [])
                   if row_start <= int(row) <= row_end}
    repeat_rows, repeat_h = _repeat_rows_and_height(
        heights, repeat_range, row_start, row_end)

    pages = []
    row = row_start
    split_warned = False
    while row <= row_end and len(pages) < MAX_PRINT_PAGES:
        with_repeat = bool(repeat_rows) and row > repeat_rows[-1]
        page_start, end_row, row, split_warned = _fill_one_page(
            heights, blocks, break_after, row, row_end, printable_h_px,
            repeat_h, with_repeat, warnings, name, split_warned)
        pages.append({'row_start': page_start, 'row_end': end_row,
                      'repeat_rows': repeat_rows if with_repeat else None})
    if row <= row_end:
        warnings.append('工作表「%s」分页数超过上限（%d 页），超出部分未生成'
                        % (name, MAX_PRINT_PAGES))
    page_breaks = [page['row_start'] for page in pages[1:]]
    return pages, page_breaks


# ==================== 列分块 ====================

def _column_bands(sheet, col_start, col_end, printable_w_px, scale, warnings):
    """列方向分块：可见列按缩放宽贪婪装箱，返回 [(起始列, 结束列)]。"""
    name = sheet.get('name') or ''
    hidden = set(sheet.get('hidden_columns') or [])
    widths = sheet.get('column_widths') or {}
    bands = []
    current = []
    current_w = 0.0
    for col in range(col_start, col_end + 1):
        letter = get_column_letter(col)
        if letter in hidden:
            continue
        width = round(column_width_px(widths.get(letter)) * scale, 2)
        if current and current_w + width > printable_w_px + _EPSILON_PX:
            bands.append((current[0], current[-1]))
            current, current_w = [], 0.0
        if width > printable_w_px + _EPSILON_PX:
            warnings.append('工作表「%s」第 %d 列宽度超出单页可打印宽度，'
                            '该列内容可能被裁切' % (name, col))
        current.append(col)
        current_w = round(current_w + width, 2)
    if current:
        bands.append((current[0], current[-1]))
    return bands or [(col_start, col_end)]


# ==================== 单表配置与缩放 ====================

def _content_size_px(sheet, row_start, row_end, col_start, col_end, scale):
    hidden_cols = set(sheet.get('hidden_columns') or [])
    widths = sheet.get('column_widths') or {}
    width = 0.0
    for col in range(col_start, col_end + 1):
        letter = get_column_letter(col)
        if letter in hidden_cols:
            continue
        width += column_width_px(widths.get(letter))
    width = round(width * scale, 2)
    height = round(sum(_row_heights_px(sheet, row_start, row_end, scale).values()), 2)
    return width, height


def _sheet_print_config(sheet, params, is_main):
    """单表打印设置：缩放模式/比例 + 生效打印区域文本。

    分页参数（重复表头/手动分页/显式打印区域）与 LibreOffice 路径一致
    只作用于主工作表；其余表用源文件自带打印区域（上传时解析）。
    """
    preset = params.get('preset')
    print_cfg = sheet.get('print') or {}
    mode = params.get('scale_mode') or 'fit_width'
    percent = params.get('scale_percent') or 100
    if preset == 'original':
        fit_w, fit_h = print_cfg.get('fit_to_width'), print_cfg.get('fit_to_height')
        scale = print_cfg.get('scale')
        if fit_w is not None or fit_h is not None:
            mode = 'fit_page' if (fit_h or 0) > 0 else 'fit_width'
            percent = 100
        elif scale is not None:
            mode, percent = 'percent', max(10, min(400, int(scale)))
        else:
            mode, percent = 'fit_width', 100
    area = print_cfg.get('print_area') or ''
    if is_main and params.get('print_area'):
        area = params.get('print_area')
    return mode, percent, area


def _resolve_scale(mode, percent, content_w_px, content_h_px,
                   printable_w_px, printable_h_px):
    """缩放倍率：fit 只缩不放；actual=1；percent=指定比例。"""
    if mode == 'percent':
        return round(max(0.1, min(4.0, float(percent) / 100.0)), 4)
    if mode == 'actual':
        return 1.0
    scale = 1.0
    if content_w_px > 0:
        scale = min(scale, printable_w_px / content_w_px)
    if mode == 'fit_page' and content_h_px > 0:
        scale = min(scale, printable_h_px / content_h_px)
    return round(max(0.05, scale), 4)


def _sheet_is_blank(sheet, bindings):
    """工作表没有任何单元格内容且没有绑定字段（打印会产生空白页）。"""
    if any((cell or {}).get('v') not in (None, '')
           for cell in (sheet.get('cells') or {}).values()):
        return False
    if any(binding.get('sheet_name') == sheet.get('name')
           for binding in bindings or []):
        return False
    return True


def _assemble_pages(row_pages, bands, page_order):
    """行页 × 列块组装（down_then_over：先下后右；over_then_down：先右后下）。"""
    pages = []
    if page_order == 'over_then_down':
        for row_page in row_pages:
            for band in bands:
                pages.append(dict(row_page, col_start=band[0], col_end=band[1]))
    else:
        for band in bands:
            for row_page in row_pages:
                pages.append(dict(row_page, col_start=band[0], col_end=band[1]))
    return pages


def _empty_plan(warnings=None):
    return {
        'plan_version': PRINT_PLAN_VERSION,
        'paper': {'name': 'A4', 'width_mm': 210, 'height_mm': 297,
                  'orientation': 'portrait'},
        'margins_mm': {'top': 19.1, 'bottom': 19.1, 'left': 17.8, 'right': 17.8},
        'scale': 1.0,
        'page_count': 0,
        'warnings': warnings or ['布局数据缺失，无法计算统一打印布局'],
        'sheets': {},
    }


def _sheet_plan(sheet, params, main_name, bindings, warnings,
                printable_w_px, printable_h_px):
    """单工作表计划：打印范围/缩放/告警/行列分页与分块（行页×列块组装）。

    跳过空表返回 None（告警由本函数追加）；分页参数（重复表头/手动
    分页/显式打印区域）与 LibreOffice 路径一致只作用于主工作表。
    """
    name = sheet.get('name') or ''
    if _sheet_is_blank(sheet, bindings):
        warnings.append('工作表「%s」没有可打印内容，已跳过' % name)
        return None
    is_main = name == main_name
    mode, percent, area_text = _sheet_print_config(sheet, params, is_main)
    area_range = _print_range(sheet, area_text, warnings)
    if area_range:
        row_start, row_end, col_start, col_end = area_range
    else:
        row_start, row_end = 1, int(sheet.get('max_row') or 1)
        col_start, col_end = 1, int(sheet.get('max_column') or 1)
    row_start = max(1, min(row_start, int(sheet.get('max_row') or 1)))
    col_start = max(1, min(col_start, int(sheet.get('max_column') or 1)))
    content_w, content_h = _content_size_px(
        sheet, row_start, row_end, col_start, col_end, 1.0)
    scale = _resolve_scale(mode, percent, content_w, content_h,
                           printable_w_px, printable_h_px)
    _append_scale_warnings(name, mode, scale, content_w,
                           printable_w_px, warnings)
    repeat_range = None
    manual_breaks = []
    if is_main:
        repeat_range = _repeat_range(params, row_end, warnings, name)
        manual_breaks = [int(r) for r in params.get('row_breaks') or []]
    row_pages, row_breaks = _paginate_rows(
        sheet, row_start, row_end, printable_h_px, scale,
        repeat_range, manual_breaks, warnings)
    bands = _column_bands(sheet, col_start, col_end,
                          printable_w_px, scale, warnings)
    pages = _assemble_pages(row_pages, bands, params.get('page_order'))
    return name, {
        'scale': scale,
        'page_count': len(pages),
        'print_row_start': row_start, 'print_row_end': row_end,
        'print_col_start': col_start, 'print_col_end': col_end,
        'row_breaks': row_breaks,
        'col_breaks': [band[0] for band in bands[1:]],
        'pages': pages,
    }


def _append_scale_warnings(name, mode, scale, content_w, printable_w_px,
                           warnings):
    """缩放相关告警：fit 过小提示字号过小；固定缩放横向溢出提示分页。"""
    if mode in ('fit_width', 'fit_page') and scale < _MIN_READABLE_SCALE:
        warnings.append('工作表「%s」按纸张适配后缩放约 %d%%，打印字号可能过小'
                        % (name, round(scale * 100)))
    elif mode in ('actual', 'percent'):
        scaled_w = round(content_w * scale, 2)
        if scaled_w > printable_w_px + 28.3:  # 约 0.75cm 容差
            warnings.append('工作表「%s」内容宽度超过一页可打印宽度，将横向分页'
                            % name)


def _plan_paper_geometry(params, main):
    """整份 PDF 的纸张几何：名称/方向/页宽高（横向交换）/边距/可打印区 px。

    纸张与方向全 PDF 统一（取主工作表解析结果）；auto 方向按主工作表
    内容宽度解析。
    """
    paper_name = _paper_name(params, main)
    orientation = _orientation(params, main, paper_name)
    short_mm, long_mm = PAPER_SIZE_MM[paper_name]
    page_w_mm, page_h_mm = (
        (long_mm, short_mm) if orientation == 'landscape' else (short_mm, long_mm))
    margins_mm = _margins_mm(params, main)
    return {
        'paper_name': paper_name, 'orientation': orientation,
        'page_w_mm': page_w_mm, 'page_h_mm': page_h_mm,
        'margins_mm': margins_mm,
        'printable_w_px': max(
            2.0, (page_w_mm - margins_mm['left'] - margins_mm['right'])
            * PX_PER_MM),
        'printable_h_px': max(
            2.0, (page_h_mm - margins_mm['top'] - margins_mm['bottom'])
            * PX_PER_MM),
    }


def _collect_sheet_plans(sheets, params, main_name, bindings, warnings,
                         geometry):
    """逐表构建计划（空表跳过；总页数超上限明确告警并截断）。"""
    sheet_plans = {}
    total_pages = 0
    for sheet in sheets:
        item = _sheet_plan(sheet, params, main_name, bindings, warnings,
                           geometry['printable_w_px'],
                           geometry['printable_h_px'])
        if item is None:
            continue
        name, plan = item
        sheet_plans[name] = plan
        total_pages += plan['page_count']
        if total_pages > MAX_PRINT_PAGES:
            warnings.append('分页总数超过上限（%d 页），超出部分未生成'
                            % MAX_PRINT_PAGES)
            break
    return sheet_plans


def build_print_plan(meta, pdf_params=None):
    """计算统一打印计划（权威分页规则；网页与浏览器 PDF 共用）。

    meta：版本快照形态 {'sheet_name', 'layout': {'sheets': [...]},
    'pdf_params': {...}}；pdf_params 显式传入时整体覆盖（导出高级设置）。
    返回计划 dict（结构见模块 docstring），不抛业务异常——异常数据回落
    缺省计划并以 warnings 提示，绝不静默失真。
    """
    layout = (meta or {}).get('layout') or {}
    sheets = layout.get('sheets') or []
    warnings = []
    if not sheets:
        return _empty_plan()
    params = _effective_params((meta or {}).get('pdf_params'), pdf_params)
    main = next((sheet for sheet in sheets
                 if sheet.get('name') == (meta or {}).get('sheet_name')),
                sheets[0])
    main_name = main.get('name') or ''
    geometry = _plan_paper_geometry(params, main)
    bindings = (meta or {}).get('field_bindings') or []

    sheet_plans = _collect_sheet_plans(sheets, params, main_name, bindings,
                                       warnings, geometry)
    if not sheet_plans:
        warnings.append('没有可打印的工作表内容')
    main_plan = sheet_plans.get(main_name) or {}
    return {
        'plan_version': PRINT_PLAN_VERSION,
        'paper': {'name': geometry['paper_name'],
                  'width_mm': geometry['page_w_mm'],
                  'height_mm': geometry['page_h_mm'],
                  'orientation': geometry['orientation']},
        'margins_mm': geometry['margins_mm'],
        'scale': main_plan.get('scale', 1.0),
        'horizontal_centered': bool(params.get('horizontal_centered')),
        'vertical_centered': bool(params.get('vertical_centered')),
        'page_count': sum(item['page_count'] for item in sheet_plans.values()),
        'warnings': warnings,
        'sheets': sheet_plans,
    }


def _repeat_range(params, row_end, warnings, name):
    """重复表头行 'a:b' → (a, b)；越界回落 None 并告警（历史快照防御）。"""
    text = str(params.get('repeat_header_rows') or '').strip()
    if not text:
        return None
    try:
        a, b = (int(part) for part in text.split(':', 1))
    except (ValueError, TypeError):
        return None
    if a < 1 or a > b:
        return None
    if b > row_end:
        warnings.append('重复表头行 %s 超出工作表「%s」打印范围，已忽略'
                        % (text, name))
        return None
    return (a, b)
