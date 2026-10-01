# -*- coding: utf-8 -*-
"""ICS 200/60 内话系统周月维护记录表 Excel 路线真实端到端验收（2026-10-01）。

真实 HTTP 全栈：Windows 宿主机 → nginx(8080, /api 剥离) → gunicorn → MariaDB，
账号 ics_e2e（真实 Role 权限，setup_ics_user.py 在容器内创建）。

验收项对齐任务书：
 1 上传真实 ICS 文件           2 解析预览（布局/合并/打印设置/PDF参数播种）
 3 空格保真（布局 JSON 逐字符） 4 绑定（单字段+批量+负向）
 5 草稿保存并重新打开（空格仍在）6 PDF 打印参数（preset a4_auto/居中/打印区域+负向）
 7 草稿 PDF 预览               8 发布（版本快照携带新参数）
 9 创建台账+填写（含空格值）+重新打开 10 导出预检 + Excel 导出（空格逐字符）
11 PDF 导出（字形级空格断言 + fit_width 无横向分页 + 页数合理）
12 历史版本隔离（发布 v2 后 v1 参数与导出不变）

宿主机运行：python e2e_ics.py（需 requests/openpyxl/pdfminer.six）。
"""
import hashlib
import io
import json
import os
import sys
import uuid

import requests

BASE = 'http://localhost:8080/api'
TOKEN = ('ics-e2e-' * 8)[:32]
OUT = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(OUT, '00_ICS200-60内话系统周月维护记录表.xlsx')

# 模板内关键空格场景（与原始文件逐字符一致）
TITLE = 'ICS 200/60内话系统周月维护记录表'
MAINT_RECORD = '维护记录：                              温度：       湿度：'
DATE_LINE = '     年   月   日'
FILLED_TEMP = '  温 度 25.5  '
FILLED_HUMIDITY = '湿 度：60 %RH，连续  空格'

S = requests.Session()
S.headers.update({'x-token': TOKEN})
RESULTS = []


def step(no, name, ok, detail=''):
    RESULTS.append((no, name, 'PASS' if ok else 'FAIL', detail))
    print('[%s] %s %s %s' % ('PASS' if ok else 'FAIL', no, name, detail), flush=True)
    return ok


def body_of(resp):
    try:
        payload = resp.json()
    except ValueError:
        return None, 'HTTP %s 非JSON: %s' % (resp.status_code, resp.text[:200])
    if isinstance(payload, dict):
        err = payload.get('error') or ''
        data = payload.get('data')
        if isinstance(data, dict) and 'data' in data and set(data.keys()) <= {'data', 'total', 'error'}:
            data = data['data']
        return data, err
    return payload, ''


def api(method, path, expect_error=None, **kw):
    resp = S.request(method, BASE + path, timeout=180, **kw)
    data, err = body_of(resp)
    if expect_error is not None:
        return resp, data, err
    if err or resp.status_code != 200:
        raise AssertionError('%s %s -> HTTP %s error=%r data=%r'
                             % (method, path, resp.status_code, err, str(data)[:300]))
    return resp, data, err


with open(SRC, 'rb') as f:
    ICS = f.read()
SRC_SHA = hashlib.sha256(ICS).hexdigest()
print('ICS xlsx sha256=%s size=%d' % (SRC_SHA, len(ICS)), flush=True)

# ---------- 0 认证连通 ----------
resp, data, err = api('GET', '/equipment-lifecycle/excel-templates/')
step('0', '认证与列表连通（GET /excel-templates/）', isinstance(data, list),
     '现有模板数=%s' % len(data) if isinstance(data, list) else type(data).__name__)

# ---------- 1 上传 ----------
tpl_name = 'ICS200-60内话系统周月维护记录表（E2E验收）'
resp = S.post(BASE + '/equipment-lifecycle/excel-templates/',
              files={'file': ('00_ICS200-60内话系统周月维护记录表.xlsx', io.BytesIO(ICS),
                              'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
              data={'name': tpl_name, 'ledger_type': 'week_month',
                    'description': 'ICS 200/60 真实 WPS/Excel 模板端到端验收（2026-10-01）'},
              timeout=120)
tpl, err = body_of(resp)
ok = step('1', '上传真实 ICS 200/60 模板', not err and resp.status_code == 200 and bool(tpl and tpl.get('id')),
          err or 'template_id=%s' % (tpl and tpl.get('id')))
if not ok:
    print('上传失败，终止')
    json.dump(RESULTS, open(os.path.join(OUT, 'results.json'), 'w'), ensure_ascii=False)
    sys.exit(1)
TID = tpl['id']
step('1a', '上传登记哈希 == 原始文件 SHA-256', tpl.get('draft_source_hash') == SRC_SHA,
     '%s' % SRC_SHA[:16])

# ---------- 2 解析预览 ----------
draft = tpl.get('draft') or {}
layout = draft.get('layout') or {}
sheets = layout.get('sheets') or []
s1 = next((s for s in sheets if s.get('name') == 'Sheet1'), None)
pdf_seeded = draft.get('pdf_params') or {}
checks = {
    '识别主工作表 Sheet1': draft.get('sheet_name') == 'Sheet1' and s1 is not None,
    '合并单元格 23 处': s1 and len(s1.get('merged_ranges') or []) == 23,
    '单元格值+样式落 meta': s1 and len(s1.get('cells') or {}) >= 15,
    '行高捕获（13-25 行 x 30）': s1 and all(
        (s1.get('row_heights') or {}).get(str(r)) == 30.0 for r in range(13, 26)),
    '打印设置捕获(A4纵向)': s1 and s1.get('print', {}).get('paper_size') == 9,
    '播种 preset=a4_auto（样板纸张方向不触发 original）': pdf_seeded.get('preset') == 'a4_auto',
    '播种 方向=auto': pdf_seeded.get('orientation') == 'auto',
    '播种 缩放=fit_width': pdf_seeded.get('scale_mode') == 'fit_width',
    '兼容性检查 ok': (tpl.get('compatibility') or {}).get('status') == 'ok',
}
step('2', '解析预览：布局 meta 完整 + PDF 参数播种 a4_auto', all(checks.values()),
     '; '.join('%s=%s' % (k, v) for k, v in checks.items() if not v) or '全部通过')

# ---------- 3 空格保真（布局 JSON 逐字符） ----------
cells = (s1 or {}).get('cells') or {}
step('3', '空格保真：A1 标题中间空格', cells.get('A1', {}).get('v') == TITLE,
     'A1=%r' % cells.get('A1', {}).get('v'))
step('3a', '空格保真：A13 维护记录留空线（30+7 连续空格）',
     cells.get('A13', {}).get('v') == MAINT_RECORD,
     'A13=%r' % cells.get('A13', {}).get('v'))
step('3b', '空格保真：E25 年月日（首部5空格+字间3空格）',
     cells.get('E25', {}).get('v') == DATE_LINE,
     'E25=%r' % cells.get('E25', {}).get('v'))


def b(coord, name, itype, required=True, options=None):
    return {'field_id': str(uuid.uuid4()), 'sheet_name': 'Sheet1', 'cell_coordinate': coord,
            'field_name': name, 'input_type': itype, 'is_required': required,
            'unit': '', 'options': options or []}


# ---------- 4 绑定 ----------
single = [b('C3', '服务器A/B状态显示', 'status')]
resp, data, err = api('PUT', '/equipment-lifecycle/excel-templates/%s/bindings/' % TID,
                      json={'field_bindings': single})
step('4', '单字段绑定（C3 合并锚点 状态显示）',
     not err and len((data.get('draft') or {}).get('field_bindings') or []) == 1,
     err or 'draft.field_bindings=1')
resp, data, err = api('PUT', '/equipment-lifecycle/excel-templates/%s/bindings/' % TID,
                      expect_error=True,
                      json={'field_bindings': [b('A13', '维护记录格', 'text')]})
step('4a', '负向：有内容的合并锚点格拒绝绑定', '已有内容' in err, err)

batch = [
    b('C4', '无线板卡状态', 'status'), b('E4', '无线板卡备注', 'text', False),
    b('C5', '有线板卡状态', 'status'), b('C6', '电源电压状态', 'status'),
    b('E6', '电源电压备注', 'text', False), b('C7', '系统管理终端状态', 'status'),
    b('C8', '操作席位状态', 'status'), b('C9', '环路连接状态', 'status'),
    b('C10', '运控环状态', 'status'), b('C12', '线路连接状态', 'status'),
    b('A25', '维护人签字', 'text', True),
    b('D25', '维护日期', 'date', True),
]
resp, data, err = api('PUT', '/equipment-lifecycle/excel-templates/%s/bindings/' % TID,
                      json={'field_bindings': single + batch})
got = (data.get('draft') or {}).get('field_bindings') or []
step('4b', '批量绑定合计 13 字段整体落库', not err and len(got) == 13,
     err or '落库=%d' % len(got))

# ---------- 5 草稿保存 & 重新打开 ----------
resp, again, err = api('GET', '/equipment-lifecycle/excel-templates/%s/' % TID)
d2 = again.get('draft') or {}
names = {f['field_name']: f['field_id'] for f in d2.get('field_bindings') or []}
step('5', '草稿保存并重新打开（绑定+PDF参数持久化）',
     len(names) == 13 and bool(d2.get('pdf_params')),
     '字段=%d pdf_params.preset=%s' % (len(names), (d2.get('pdf_params') or {}).get('preset')))
# 重新打开后空格仍在（与步骤 3 同一断言走独立 GET 路径）
cells2 = next((s for s in (d2.get('layout') or {}).get('sheets') or []
               if s.get('name') == 'Sheet1'), {}).get('cells') or {}
step('5a', '重新打开后空格仍在（A13/E25 逐字符）',
     cells2.get('A13', {}).get('v') == MAINT_RECORD
     and cells2.get('E25', {}).get('v') == DATE_LINE, '')

# ---------- 6 PDF 打印参数（新协议：preset/auto/居中/打印区域） ----------
params = {
    'preset': 'a4_auto', 'paper_size': 'A4', 'orientation': 'auto',
    'margins': {'top': 1.91, 'bottom': 1.91, 'left': 1.78, 'right': 1.78,
                'header': 0.76, 'footer': 0.76},
    'scale_mode': 'fit_width', 'scale_percent': 100,
    'repeat_header_rows': '', 'row_breaks': [], 'page_order': 'down_then_over',
    'print_area': 'A1:E25', 'horizontal_centered': True, 'vertical_centered': False,
}
resp, data, err = api('PUT', '/equipment-lifecycle/excel-templates/%s/pdf-params/' % TID,
                      json={'pdf_params': params})
saved = (data.get('draft') or {}).get('pdf_params') or {}
ok = not err and saved.get('preset') == 'a4_auto' and saved.get('print_area') == 'A1:E25' \
     and saved.get('horizontal_centered') is True
step('6', 'PDF 参数保存（a4_auto/方向auto/打印区域A1:E25/水平居中）', ok,
     err or json.dumps({k: saved.get(k) for k in ('preset', 'orientation', 'print_area', 'horizontal_centered')},
                       ensure_ascii=False))
bad = dict(params)
bad['print_area'] = 'A1:ZZ99'
resp, data, err = api('PUT', '/equipment-lifecycle/excel-templates/%s/pdf-params/' % TID,
                      expect_error=True, json={'pdf_params': bad})
step('6a', '负向：打印区域越界（ZZ 列）拒绝', bool(err), err)
bad2 = dict(params)
bad2['preset'] = 'magic'
resp, data, err = api('PUT', '/equipment-lifecycle/excel-templates/%s/pdf-params/' % TID,
                      expect_error=True, json={'pdf_params': bad2})
step('6b', '负向：未知打印模式拒绝', bool(err), err)

# ---------- 7 草稿 PDF 预览 ----------
resp = S.post(BASE + '/equipment-lifecycle/excel-templates/%s/preview/pdf/' % TID, timeout=300)
if resp.status_code == 200 and resp.content[:5] == b'%PDF-':
    open(os.path.join(OUT, '03_草稿PDF预览.pdf'), 'wb').write(resp.content)
    step('7', '草稿 PDF 预览（未填写留空，应用 a4_auto 参数）', True,
         'bytes=%d' % len(resp.content))
else:
    step('7', '草稿 PDF 预览（未填写留空，应用 a4_auto 参数）', False,
         'HTTP %s %s' % (resp.status_code, resp.text[:200]))

# ---------- 8 发布 ----------
resp, ver, err = api('POST', '/equipment-lifecycle/excel-templates/%s/publish/' % TID)
ok = not err and ver and ver.get('version_no') == 1 and ver.get('field_count') == 13
step('8', '发布不可变版本 v1', ok, err or 'version_id=%s field_count=%s' % (
    ver and ver.get('id'), ver and ver.get('field_count')))
VID = ver['id']
resp, vd, err = api('GET', '/equipment-lifecycle/excel-templates/%s/versions/1/' % TID)
vp = vd.get('pdf_params') or {}
step('8a', '版本快照携带新 PDF 参数（preset/print_area/居中）',
     vp.get('preset') == 'a4_auto' and vp.get('print_area') == 'A1:E25'
     and vp.get('horizontal_centered') is True, '')
step('8b', '版本快照空格仍在（layout.A13 逐字符）',
     next((s for s in (vd.get('layout') or {}).get('sheets') or []
           if s.get('name') == 'Sheet1'), {}).get('cells', {}).get('A13', {}).get('v')
     == MAINT_RECORD, '')

# ---------- 9 创建台账 + 填写 ----------
RUN_TOKEN = 'ics-e2e-%s' % uuid.uuid4().hex[:12]
resp, ledger, err = api('POST', '/equipment-lifecycle/excel-ledgers/',
                        json={'version_id': VID, 'executed_date': '2026-10-01',
                              'executor_names': 'ICS验收员', 'ledger_type': 'week_month',
                              'client_token': RUN_TOKEN})
ok = not err and ledger and ledger.get('created_from') == 'excel_template'
step('9', '创建台账（created_from=excel_template）', ok,
     err or 'id=%s no=%s' % (ledger and ledger.get('id'), ledger and ledger.get('ledger_no')))
LID = ledger['id']
values = {}
for field in (single + batch):
    name, itype = field['field_name'], field['input_type']
    if itype == 'status':
        values[field['field_id']] = '正常'
    elif itype == 'date':
        values[field['field_id']] = '2026-10-01'
    elif name == '无线板卡备注':
        values[field['field_id']] = FILLED_TEMP
    elif name == '电源电压备注':
        values[field['field_id']] = FILLED_HUMIDITY
    elif name == '维护人签字':
        values[field['field_id']] = '  ICS 验收员  '
    else:
        values[field['field_id']] = '记录正常'
resp, data, err = api('PUT', '/equipment-lifecycle/excel-ledgers/%s/' % LID,
                      json={'values': values})
got_values = ((data.get('excel') or {}).get('values') or {})
note_id = next(f['field_id'] for f in single + batch if f['field_name'] == '无线板卡备注')
ok = not err and got_values.get(note_id) == FILLED_TEMP
step('9a', '填写值保存（含空格值逐字符保留）', ok, err or '备注=%r' % got_values.get(note_id))
resp, detail, err = api('GET', '/equipment-lifecycle/excel-ledgers/%s/' % LID)
values2 = (detail.get('excel') or {}).get('values') or {}
step('9b', '重新打开台账：空格填写值仍在', values2.get(note_id) == FILLED_TEMP,
     '%r' % values2.get(note_id))
resp, dup_ledger, err = api('POST', '/equipment-lifecycle/excel-ledgers/',
                            json={'version_id': VID, 'executed_date': '2026-10-01',
                                  'executor_names': 'ICS验收员', 'ledger_type': 'week_month',
                                  'client_token': RUN_TOKEN})
step('9c', '同 client_token 防重复创建', dup_ledger and dup_ledger.get('id') == LID, '')

# ---------- 10 导出预检 + Excel 导出 ----------
resp, warnings, err = api('POST', '/equipment-lifecycle/excel-ledgers/%s/export/check/' % LID)
warning_items = warnings.get('warnings', []) if isinstance(warnings, dict) else (warnings or [])
step('10', '导出预检（必填已填、fit_width 宽度适配→无横向溢出告警）',
     not err and not any('横向分页' in w for w in warning_items),
     'warnings=%s' % (warning_items,))
resp = S.post(BASE + '/equipment-lifecycle/excel-ledgers/%s/export/excel/' % LID, timeout=180)
if resp.status_code == 200 and resp.content[:2] == b'PK':
    open(os.path.join(OUT, '06_已提交台账导出.xlsx'), 'wb').write(resp.content)
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb['Sheet1']
    ok = (ws['A1'].value == TITLE and ws['A13'].value == MAINT_RECORD
          and ws['E25'].value == DATE_LINE and ws['E4'].value == FILLED_TEMP)
    step('10a', 'Excel 导出空格保真（模板格+填写值逐字符）', ok,
         'A13=%r E25=%r' % (ws['A13'].value, ws['E25'].value))
else:
    step('10a', 'Excel 导出空格保真（模板格+填写值逐字符）', False,
         'HTTP %s' % resp.status_code)

# 提交台账（导出 PDF 前置状态按任务书无强制，但提交走全流程）
resp, data, err = api('POST', '/equipment-lifecycle/excel-ledgers/%s/submit/' % LID)
step('10b', '提交台账（草稿→已提交）', not err, err or '')

# ---------- 11 PDF 导出（字形级空格断言） ----------
resp = S.post(BASE + '/equipment-lifecycle/excel-ledgers/%s/export/pdf/' % LID, timeout=300)
if resp.status_code == 200 and resp.content[:5] == b'%PDF-':
    open(os.path.join(OUT, '07_已提交台账导出.pdf'), 'wb').write(resp.content)
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LTChar, LTTextContainer, LTTextLine
    page_count = 0
    pages_lines = []
    for page in extract_pages(io.BytesIO(resp.content)):
        page_count += 1
        page_lines = []
        for element in page:
            if isinstance(element, LTTextContainer):
                for line in element:
                    if isinstance(line, LTTextLine):
                        page_lines.append(''.join(c.get_text() for c in line
                                                  if isinstance(c, LTChar)))
        pages_lines.append(page_lines)
    joined = '\n'.join('\n'.join(page) for page in pages_lines)
    import re
    match = re.search(r'维护记录：( +)温度：( +)湿度：', joined)
    ok30 = bool(match and len(match.group(1)) == 30 and len(match.group(2)) == 7)
    gap_detail = ('未匹配' if not match else
                  '(%d, %d)' % (len(match.group(1)), len(match.group(2))))
    step('11', 'PDF 导出：A13 留空线 30+7 空格逐字形存在', ok30, '实际=%s' % gap_detail)
    date_match = re.search(r'( +)年( +)月( +)日', joined)
    ok_date = bool(date_match and len(date_match.group(1)) == 5
                   and len(date_match.group(2)) == 3 and len(date_match.group(3)) == 3)
    date_detail = ('未匹配' if not date_match else
                   '(%d, %d, %d)' % (len(date_match.group(1)), len(date_match.group(2)),
                                     len(date_match.group(3))))
    step('11a', 'PDF 导出：E25 年月日（5+3+3 空格）逐字形存在', ok_date, '实际=%s' % date_detail)
    step('11b', 'PDF 导出：填写值空格保留（温 度 25.5）', '  温 度 25.5' in joined, '')
    # fit_width：首页包含最右列「备注」表头（无横向分页）
    first_page = pages_lines[0] if pages_lines else []
    step('11c', 'PDF 导出：宽度适配一页（首页含全部 5 列表头，无横向分页）',
         any('维护内容' in t for t in first_page) and any('备注' in t for t in first_page),
         'page_count=%d' % page_count)
    step('11d', 'PDF 导出：页数合理（1-3 页，无 4 页式横向溢出）', 1 <= page_count <= 3,
         'page_count=%d' % page_count)
else:
    step('11', 'PDF 导出（真实 LibreOffice）', False, 'HTTP %s %s' % (resp.status_code, resp.text[:200]))

# ---------- 12 历史版本隔离 ----------
params_v2 = {'preset': 'single_page', 'paper_size': 'A4', 'orientation': 'auto',
             'scale_mode': 'fit_page', 'margins': params['margins'],
             'repeat_header_rows': '', 'row_breaks': [], 'page_order': 'down_then_over',
             'print_area': '', 'horizontal_centered': False, 'vertical_centered': False}
resp, data, err = api('PUT', '/equipment-lifecycle/excel-templates/%s/pdf-params/' % TID,
                      json={'pdf_params': params_v2})
step('12', '草稿改为 single_page 并发布 v2', not err, err)
resp, ver2, err = api('POST', '/equipment-lifecycle/excel-templates/%s/publish/' % TID)
step('12a', '发布 v2（single_page 快照）', not err and ver2.get('version_no') == 2, err)
resp, v1, err = api('GET', '/equipment-lifecycle/excel-templates/%s/versions/1/' % TID)
step('12b', 'v1 快照参数不变（历史台账不受新版本影响）',
     (v1.get('pdf_params') or {}).get('preset') == 'a4_auto'
     and (v1.get('pdf_params') or {}).get('scale_mode') == 'fit_width', '')
resp = S.post(BASE + '/equipment-lifecycle/excel-ledgers/%s/export/pdf/' % LID, timeout=300)
same = False
detail_12c = 'HTTP %s' % resp.status_code
if resp.status_code == 200 and resp.content[:5] == b'%PDF-':
    from pdfminer.high_level import extract_pages as _ep
    from pdfminer.layout import LTChar as _C, LTTextContainer as _TC, LTTextLine as _TL
    def _pdf_lines(data):
        out = []
        for page in _ep(io.BytesIO(data)):
            for element in page:
                if isinstance(element, _TC):
                    for line in element:
                        if isinstance(line, _TL):
                            out.append(''.join(c.get_text() for c in line if isinstance(c, _C)))
        return out
    before = _pdf_lines(open(os.path.join(OUT, '07_已提交台账导出.pdf'), 'rb').read())
    after = _pdf_lines(resp.content)
    same = before == after
    detail_12c = 'lines_before=%d lines_after=%d equal=%s（LibreOffice 内嵌时间戳导致字节不恒等，按字形内容比对）' % (
        len(before), len(after), same)
step('12c', 'v1 台账再次导出 PDF 与 v2 发布前内容一致（按字形文本）', same, detail_12c)

# ---------- 汇总 ----------
passed = sum(1 for r in RESULTS if r[2] == 'PASS')
print('=' * 60)
print('TOTAL %d  PASS %d  FAIL %d' % (len(RESULTS), passed, len(RESULTS) - passed), flush=True)
json.dump(RESULTS, open(os.path.join(OUT, 'results.json'), 'w'), ensure_ascii=False, indent=1)
sys.exit(0 if passed == len(RESULTS) else 1)
