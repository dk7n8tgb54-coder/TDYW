# -*- coding: utf-8 -*-
"""NB500.xlsx 结构体检：为 E2E 验收建立原始基准。只读。"""
import json
import sys

import openpyxl

PATH = '/data/spug/spug_api/NB500_tmp_e2e.xlsx'
wb = openpyxl.load_workbook(PATH, data_only=False)

report = {}
report['sheetnames'] = wb.sheetnames
try:
    dn = getattr(wb.defined_names, 'definedName', None) or {}
    report['defined_names'] = {d.name: str(d.value) for d in dn} if not isinstance(dn, dict) else {k: str(v.value) for k, v in dn.items()}
except Exception as e:
    report['defined_names'] = 'inspect-error: %s' % e

for ws in wb.worksheets:
    info = {
        'dims': ws.dimensions,
        'max_row': ws.max_row, 'max_col': ws.max_column,
        'merged': sorted(str(r) for r in ws.merged_cells.ranges),
        'print_area': str(ws.print_area) if ws.print_area else None,
        'page_setup': {
            'paperSize': ws.page_setup.paperSize,
            'orientation': ws.page_setup.orientation,
            'scale': ws.page_setup.scale,
            'fitToWidth': ws.page_setup.fitToWidth,
            'fitToHeight': ws.page_setup.fitToHeight,
            'fitToPage': bool(ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage),
        },
        'page_margins': {k: getattr(ws.page_margins, k) for k in
                         ('left', 'right', 'top', 'bottom', 'header', 'footer')},
        'header': ws.oddHeader.center.text if ws.oddHeader and ws.oddHeader.center else None,
        'footer': ws.oddFooter.center.text if ws.oddFooter and ws.oddFooter.center else None,
        'col_widths': {c: round(d.width, 2) if d.width else None
                       for c, d in ws.column_dimensions.items()},
        'row_heights': {r: d.height for r, d in ws.row_dimensions.items() if d.height},
        'row_breaks': [b.id for b in ws.row_breaks.brk] if ws.row_breaks else [],
        'col_breaks': [b.id for b in ws.col_breaks.brk] if ws.col_breaks else [],
        'images': len(ws._images) if hasattr(ws, '_images') else 0,
        'charts': len(ws._charts) if hasattr(ws, '_charts') else 0,
        'data_validations': len(ws.data_validations.dataValidation),
        'cells': [],
    }
    for row in ws.iter_rows():
        for c in row:
            if c.value is None and not (c.border and (c.border.left.style or c.border.top.style
                                                      or c.border.right.style or c.border.bottom.style)):
                continue
            f = c.font
            cell = {
                'coord': c.coordinate,
                'value': c.value,
                'font': {'name': f.name, 'size': f.size, 'bold': f.bold,
                         'underline': f.underline, 'strike': f.strike, 'color': str(f.color.rgb) if f.color and f.color.rgb else None},
                'align': {'h': c.alignment.horizontal, 'v': c.alignment.vertical, 'wrap': c.alignment.wrap_text},
                'fill': str(c.fill.fgColor.rgb) if c.fill and c.fill.patternType and c.fill.fgColor and c.fill.fgColor.rgb else None,
                'border': {s: getattr(c.border, s).style for s in ('left', 'right', 'top', 'bottom')
                           if getattr(c.border, s) and getattr(c.border, s).style},
                'numfmt': c.number_format if c.number_format != 'General' else None,
            }
            info['cells'].append(cell)
    report[ws.title] = info

print(json.dumps(report, ensure_ascii=False, indent=1, default=str))
