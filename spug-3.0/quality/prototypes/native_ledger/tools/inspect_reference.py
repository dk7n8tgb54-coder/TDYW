#!/usr/bin/env python
"""Read-only inspection of the existing Excel ledger sample."""
from openpyxl import load_workbook
from pathlib import Path

SOURCE = Path(
    r"E:\TDYW\spug-3.0\quality\reports\ledger_excel_e2e"
    r"\01_信道维护台账验收样例.xlsx"
)

wb = load_workbook(SOURCE, data_only=False)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"\nSHEET {ws.title} dim={ws.max_row}x{ws.max_column}")
    print("merged:", [str(r) for r in ws.merged_cells.ranges])
    print(
        "page:",
        "area=", ws.print_area,
        "orientation=", ws.page_setup.orientation,
        "paper=", ws.page_setup.paperSize,
        "fitToWidth=", ws.page_setup.fitToWidth,
        "fitToHeight=", ws.page_setup.fitToHeight,
    )
    print(
        "margins:",
        ws.page_margins.left,
        ws.page_margins.right,
        ws.page_margins.top,
        ws.page_margins.bottom,
    )
    print(
        "row heights:",
        {k: round(v.height, 2) for k, v in ws.row_dimensions.items() if v.height},
    )
    print(
        "col widths:",
        {k: round(v.width, 2) for k, v in ws.column_dimensions.items() if v.width},
    )
    for row in ws.iter_rows(
        min_row=1, max_row=min(ws.max_row, 10), values_only=False
    ):
        cells = [
            (c.coordinate, c.value, c.style_id) for c in row if c.value is not None
        ]
        if cells:
            print("row", row[0].row, cells)
